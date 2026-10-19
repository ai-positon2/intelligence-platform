"""Paid Marketing Agents: Unknown Domains Alert, the URL Checker (is the address an ad sends people to broken, or does it redirect).

The logic of it, with no database and no clock of its own (so it is tested without a network):

  * which addresses to check: the final URLs of the sightings the scan already read (never a tracking template; ValueTrack placeholders taken out),
  * how to check one: ask for it through the platform's guarded fetch (public addresses only, standard ports, every redirect checked again), never
    read the page, ask again after a moment when it fails, and judge the answer: ok, broken (404/410, a server error, the domain does not resolve, a bad
    certificate, no connection, no answer, a redirect loop), blocked (the site refuses an automated check: never reported as broken) or unchecked,
  * a redirect: the address ends up somewhere else. That is a finding of its own (an ad whose page was moved, an expired offer sent to the home page, a domain
    that now forwards somewhere it should not), and one to another domain is marked as such. A change that only normalises the same page (http to https,
    adding or dropping "www.", a trailing slash, capital letters, a query string) is not a redirect worth an alert,
  * checking many politely: a few at a time, at most two to one host, inside a time budget,
  * what a finding is and when it is new, fixed or back (one finding per address in one place and of one kind),
  * the sheet of broken and redirecting links and the words of the alert.
"""
from __future__ import annotations

import hashlib
import html
import io
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from urllib.parse import urlsplit

from . import paid_domains as dom

TIMEOUT_S = 12
WORKERS = 8
PER_HOST = 2
RETRY_WAIT_S = 2
RECHECK_AFTER_S = 20 * 3600       # an address that worked is asked again after this long; a broken one at every scan
BUDGET_S = 12 * 60                # the most time one scan spends checking (the rest are checked at the next scan)
MAX_URLS = 4000
UA = "Mozilla/5.0 (compatible; Position2-URLChecker/1.0; +https://position2.com)"
SHOWN = 15
SHEET_NAME = "broken-links.xlsx"
SHEET_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
KIND_TEXT = {"broken": "Broken", "redirect": "Redirect", "redirect_domain": "Redirect to another domain"}

_PLACEHOLDER = re.compile(r"\{[^{}]*\}")


def default_fetch():
    from .event_intel_http import public_get
    return public_get


# ── Which addresses ─────────────────────────────────────────────────────────

def clean_url(url):
    """The address to check, or None: it has no host of its own (it begins with a ValueTrack placeholder), is not http(s), or is a place that is
    never a landing page (Facebook, the app stores). ValueTrack and custom parameters are taken out ("?c={campaignid}" becomes "?c=")."""
    u = str(url or "").strip()
    if not u or u.startswith("{") or not re.match(r"^https?://", u, re.I):
        return None
    u = _PLACEHOLDER.sub("", u)
    host = dom.host_of(u)
    return u if host and not dom.is_ignored(host) else None


def fingerprint(o, kind="broken"):
    key = "|".join([kind, o["platform"], o["accountId"], o["where"], str(o.get("objectId") or ""), o["url"]])
    return hashlib.sha1(key.encode()).hexdigest()[:24]


def occurrences(sightings):
    """One entry per address in one place, from the scan's sightings (a tracking template is not a final URL)."""
    out = {}
    for s in sightings:
        if s.get("redirect"):
            continue
        u = clean_url(s.get("url"))
        if not u:
            continue
        o = {"platform": s["platform"], "accountId": s["accountId"], "accountName": s["accountName"], "where": s["where"], "objectId": str(s.get("objectId") or ""),
             "objectName": s.get("objectName") or "", "campaign": s.get("campaign") or "", "adGroup": s.get("adGroup") or "", "url": u, "host": dom.host_of(u)}
        out.setdefault(fingerprint(o, "place"), o)
    return list(out.values())


# ── Checking one address ────────────────────────────────────────────────────

def judge_code(code):
    """(verdict, problem) for an HTTP status."""
    if code < 400:
        return "ok", ""
    if code in (404, 410):
        return "broken", "Page not found (%d)" % code
    if 500 <= code < 600:
        return "broken", "Server error (%d)" % code
    return "blocked", "The site answered %d to an automated check, so the link could not be verified" % code


def judge_error(e):
    """(verdict, problem) for a failed attempt."""
    text = ("%s %s" % (type(e).__name__, e)).lower()
    if "redirect" in text:
        return "broken", "Redirect loop or too many redirects"
    if any(k in text for k in ("getaddrinfo", "name or service not known", "nodename nor servname", "name resolution", "no address associated", "gaierror")):
        return "broken", "The domain does not resolve"
    if "ssl" in text or "certificate" in text or "tls" in text:
        return "broken", "The secure connection failed (certificate or SSL error)"
    if "timed out" in text or "timeout" in text:
        return "broken", "No answer within %d seconds" % TIMEOUT_S
    if isinstance(e, ValueError):
        return "unchecked", "Not checked: %s" % str(e)[:100]
    if isinstance(e, (ConnectionError, OSError)) or "connect" in text or "reset" in text or "refused" in text:
        return "broken", "Could not connect to the site"
    return "unchecked", "Not checked: %s" % type(e).__name__


def _once(url, fetch):
    try:
        r = fetch(url, timeout=TIMEOUT_S, stream=True, headers={"User-Agent": UA, "Accept": "text/html,*/*;q=0.8"})
    except Exception as e:                                                      # noqa: BLE001 (DNS, TLS, timeout, a private address, a loop)
        verdict, problem = judge_error(e)
        return {"verdict": verdict, "problem": problem, "code": None, "finalUrl": ""}
    try:
        code, final = int(r.status_code), str(getattr(r, "url", "") or "")
    finally:
        try:
            r.close()
        except Exception:                                                       # noqa: BLE001
            pass
    verdict, problem = judge_code(code)
    return {"verdict": verdict, "problem": problem, "code": code, "finalUrl": final.split("?")[0].split("#")[0][:300]}


def check_url(url, fetch, sleep=time.sleep):
    """{verdict: ok|broken|blocked|unchecked, problem, code, finalUrl}. A broken answer is asked for again after a moment, and the second answer counts."""
    r = _once(url, fetch)
    if r["verdict"] == "broken":
        sleep(RETRY_WAIT_S)
        r = _once(url, fetch)
    return r


# ── Redirects ───────────────────────────────────────────────────────────────

def _place(url):
    """A web address reduced to where it is: the host without "www.", and the path without a trailing slash, in lower case. The scheme and the
    query string do not move a page."""
    p = urlsplit(str(url or ""))
    host = (p.hostname or "").lower().rstrip(".")
    host = host[4:] if host.startswith("www.") else host
    return host, (p.path or "/").rstrip("/").lower() or "/"


def redirect_of(url, final_url):
    """None when the address ends where it started (or only normalised), else {to, crossDomain}: where it ends up, and whether that is another domain."""
    if not final_url:
        return None
    a, b = _place(url), _place(final_url)
    if a == b:
        return None
    return {"to": final_url, "crossDomain": dom.registrable(a[0]) != dom.registrable(b[0])}


def redirect_problem(red):
    return "Redirects to %s%s" % (red["to"], " (another domain)" if red["crossDomain"] else "")


# ── Checking many ───────────────────────────────────────────────────────────

def due(urls, cache, now_s):
    """The addresses to check this scan: never checked, broken last time, or last checked more than RECHECK_AFTER_S ago (one that refuses automated
    checks is not asked again every scan). `cache` is {url: {verdict, checkedAt (epoch seconds)}}."""
    out = []
    for u in urls:
        c = cache.get(u)
        if not c or c["verdict"] == "broken" or now_s - c["checkedAt"] >= RECHECK_AFTER_S:
            out.append(u)
    return out


def check_many(urls, fetch, *, sleep=time.sleep, budget_s=BUDGET_S, workers=WORKERS, per_host=PER_HOST, on_done=None, clock=time.monotonic):
    """{url: result} for the addresses reached inside the budget, a few at a time and at most per_host to one host."""
    deadline = clock() + budget_s
    sems, lock, done = {}, threading.Lock(), [0]

    def one(u):
        if clock() > deadline:
            return u, None
        host = dom.host_of(u)
        with lock:
            sem = sems.setdefault(host, threading.BoundedSemaphore(per_host))
        with sem:
            r = check_url(u, fetch, sleep)
        if on_done:
            with lock:
                done[0] += 1
                on_done(done[0])
        return u, r
    with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
        return {u: r for u, r in pool.map(one, urls) if r is not None}


def confirm(results, fetch, *, sleep=time.sleep, **kw):
    """Ask again, once the whole pass is over, for every address that came back broken: it is broken only if it fails both times."""
    again = [u for u, r in results.items() if r["verdict"] == "broken"]
    second = check_many(again, fetch, sleep=sleep, **kw)
    out = dict(results)
    for u in again:
        if u in second and second[u]["verdict"] != "broken":
            out[u] = second[u]
    return out


# ── Findings ────────────────────────────────────────────────────────────────

def _problem_of(url, c):
    """(kind, problem, target) of the latest answer for an address, or None when it is fine."""
    if c["verdict"] == "broken":
        return "broken", c["problem"], ""
    if c["verdict"] == "ok":
        red = redirect_of(url, c.get("finalUrl"))
        if red:
            return ("redirect_domain" if red["crossDomain"] else "redirect"), redirect_problem(red), red["to"]
    return None


def apply(occs, cache, existing, ok_ids, stamp):
    """The findings after a scan. `occs` are this scan's occurrences, `cache` {url: {verdict, problem, code, finalUrl}} the latest answer for each address,
    `existing` {fingerprint: finding}. Returns (rows to save, {new, reopened, fixed})."""
    writes, seen, new, reopened, fixed = [], set(), 0, 0, 0
    for o in occs:
        c = cache.get(o["url"])
        found = _problem_of(o["url"], c) if c else None
        if not found:
            continue
        kind, problem, target = found
        fp = fingerprint(o, kind)
        seen.add(fp)
        base = dict(o, fingerprint=fp, kind=kind, problem=problem, target=target, code=c.get("code"), lastSeen=stamp)
        old = existing.get(fp)
        if old is None:
            writes.append(dict(base, status="open", firstSeen=stamp, alertedAt=None, alert={}, decidedBy=None, decidedAt=None))
            new += 1
        elif old["status"] == "dismissed":
            writes.append(dict(old, problem=problem, target=target, code=c.get("code"), lastSeen=stamp))
        elif old["status"] == "fixed":
            writes.append(dict(base, status="open", firstSeen=stamp, alertedAt=None, alert={}, decidedBy=None, decidedAt=None))
            reopened += 1
        else:
            writes.append(dict(base, status="open", firstSeen=old["firstSeen"], alertedAt=old.get("alertedAt"), alert=old.get("alert") or {},
                               decidedBy=old.get("decidedBy"), decidedAt=old.get("decidedAt")))
    for fp, old in existing.items():
        if fp not in seen and old["status"] == "open" and old["accountId"] in ok_ids:
            writes.append(dict(old, status="fixed", decidedBy=None, decidedAt=None))
            fixed += 1
    return writes, {"new": new, "reopened": reopened, "fixed": fixed}


def order(rows):
    return sorted(rows, key=lambda f: (f["accountName"].lower(), f["campaign"].lower(), f["adGroup"].lower(), f["objectName"], f["url"]))


# ── The sheet and the alert ─────────────────────────────────────────────────

SHEET_HEAD = ["Account", "Platform", "Campaign", "Ad group", "Where", "Ad or asset", "Link in the ad", "Type", "Problem", "Redirects to", "Status code",
              "Since", "New this check", "Fixed or paused? (fill in)"]


def sheet(rows, new_fingerprints=()):
    """The broken and redirecting links as an Excel file (bytes), one row per link in one place, broken ones first."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Links to fix"
    ws.append(SHEET_HEAD)
    marked = set(new_fingerprints)
    for f in sorted(order(rows), key=lambda f: f["kind"] != "broken"):
        ws.append([f["accountName"], dom.PLATFORM.get(f["platform"], f["platform"]), f["campaign"], f["adGroup"], f["where"], f["objectName"], f["url"],
                   KIND_TEXT.get(f["kind"], f["kind"]), f["problem"], f.get("target") or "", f.get("code") or "", str(f.get("firstSeen") or "")[:10],
                   "Yes" if f["fingerprint"] in marked else "", ""])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for col, w in zip("ABCDEFGHIJKLMN", (28, 11, 30, 26, 22, 28, 60, 24, 44, 60, 11, 12, 14, 26)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _what(f):
    return " / ".join(x for x in (f["campaign"], f["adGroup"], f["objectName"]) if x) or f["where"]


def _line(f):
    return "%s (%s) · %s · %s · %s" % (f["accountName"], dom.PLATFORM.get(f["platform"], f["platform"]), _what(f), f["url"], f["problem"])


def _count(rows):
    nb = sum(1 for f in rows if f["kind"] == "broken")
    nr = len(rows) - nb
    parts = []
    if nb:
        parts.append("%d broken link%s" % (nb, "" if nb == 1 else "s"))
    if nr:
        parts.append("%d redirect%s" % (nr, "" if nr == 1 else "s"))
    return " and ".join(parts)


def digest(open_rows, new_rows, link):
    """{subject, text, html, slack_text, slack_blocks, attachments} for a check that found new broken or redirecting links. The sheet holds every open
    one, the words list the new ones."""
    accounts = len({f["accountId"] for f in new_rows})
    n, total = len(new_rows), len(open_rows)
    subject = "URL Checker: %s in %d ad account%s: fix or pause the affected ads" % (_count(new_rows), accounts, "" if accounts == 1 else "s")
    shown = sorted(order(new_rows), key=lambda f: f["kind"] != "broken")[:SHOWN]
    more = n - len(shown)
    head = "%s found, %d to fix in all. The attached sheet lists every one with its campaign and ad, so the affected campaigns can be paused and fixed." % (_count(new_rows), total)
    text = "\n".join([head, ""] + ["- " + _line(f) for f in shown] + (["...and %d more in the sheet." % more] if more else []) + ["", "The page: " + link])
    e = html.escape
    rows = "".join("<li><b>%s</b> (%s)<br>%s<br><code>%s</code><br>%s</li>" % (e(f["accountName"]), e(dom.PLATFORM.get(f["platform"], f["platform"])), e(_what(f)),
                                                                              e(f["url"]), e(f["problem"])) for f in shown)
    html_body = ("<div style=\"font-family:Arial,sans-serif;font-size:14px;line-height:1.5\"><p>%s found, %d to fix in all. The attached sheet (<b>%s</b>) lists every one with its "
                 "campaign and ad, so the affected campaigns can be paused and fixed.</p><ul>%s</ul>%s<p><a href=\"%s\">Open the URL Checker</a></p></div>"
                 % (e(_count(new_rows)), total, SHEET_NAME, rows, ("<p>...and %d more in the sheet.</p>" % more) if more else "", e(link)))
    blocks = [{"type": "header", "text": {"type": "plain_text", "text": subject[:150]}}]
    for f in shown[:8]:
        blocks.append({"type": "section", "text": {"type": "mrkdwn", "text": "*%s* (%s)\n%s\n`%s`\n%s" % (
            f["accountName"], dom.PLATFORM.get(f["platform"], f["platform"]), _what(f), f["url"], f["problem"])}})
    blocks.append({"type": "section", "text": {"type": "mrkdwn", "text": ("...and %d more. " % (n - 8) if n > 8 else "") + "The sheet was emailed. <%s|Open the URL Checker>" % link}})
    attachment = (SHEET_NAME, sheet(open_rows, [f["fingerprint"] for f in new_rows]), SHEET_TYPE)
    return {"subject": subject, "text": text, "html": html_body, "slack_text": "*%s*" % subject, "slack_blocks": blocks, "attachments": [attachment]}
