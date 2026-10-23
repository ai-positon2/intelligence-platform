"""Paid Marketing Agents: Strategy & Campaign Book, the downloads.

One plan, three files:

  Excel   every table of the plan as a sheet, and one Campaign Book sheet per campaign in the QA Agent's own layout
          (Field | Value | Other things to check)
  CSV     the Campaign Book of every campaign as one flat table; one campaign's own book in the QA layout; or every table as a zip
  PDF     the whole strategy and campaign book as a document to read and to send

`tables(rec)` is the one place that says what is in each table; the three formats are drawn from it, so they cannot drift apart.
A value that a spreadsheet would run as a formula (it starts with = + - @) is written as plain text in every format: marked as text in a
workbook, with a leading quote mark in a CSV.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile

from . import paid_strategy as ps, paid_strategy_book as book
from .paid_strategy_check import LIMITS, text as clean

SHEET_NAME_BAD = re.compile(r"[\[\]\*\?/\\:]")
CELL_CHARS = 1500           # the most a PDF table cell shows: a longer one would not fit a page
FORMULA = re.compile(r"^[=+\-@\t\r]")


def plain(v):
    """A cell value for a workbook: text stays text (the writer marks it as text, so a spreadsheet never runs it), numbers stay numbers."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "Yes" if v else "No"
    return v if isinstance(v, (int, float)) else ps.scrub(str(v))                 # a control character in one cell would stop the whole workbook being written


def safe(v):
    """A CSV value that no spreadsheet will run as a formula: text that starts with = + - @ gets a leading quote mark (a plain number is left alone)."""
    v = plain(v)
    if isinstance(v, str) and FORMULA.match(v) and not re.match(r"^[+-]?\d[\d,.]*%?$", v):
        return "'" + v
    return v


def _join(v, sep="; "):
    return sep.join(str(x) for x in v if x) if isinstance(v, list) else (str(v) if v is not None else "")


def _money(v, cur=""):
    return ("%s %s" % ("{:,.2f}".format(v), cur)).strip() if isinstance(v, (int, float)) else ""


def file_stem(rec):
    base = re.sub(r"[^A-Za-z0-9]+", "-", rec.get("name") or "client").strip("-")[:50] or "client"
    return "%s-strategy-and-campaign-book" % base


# ── The tables ──────────────────────────────────────────────────────────────

def tables(rec):
    """{sheet name: (header, rows)} in the order they are shown. Text only; empty tables are left out."""
    plan = rec.get("plan") or {}
    inputs = rec.get("inputs") or {}
    research = rec.get("research") or {}
    prof, comp, st = plan.get("profile") or {}, plan.get("competitive") or {}, plan.get("strategy") or {}
    camps = plan.get("campaigns") or []
    cur = inputs.get("currency") or "USD"
    out = {}

    ov = [["Client", inputs.get("name")], ["Website", inputs.get("website")], ["Industry", inputs.get("industry")], ["Goal", ps.GOALS.get(inputs.get("goal"), inputs.get("goal"))],
          ["Monthly budget", _money(inputs.get("budget"), cur)], ["Launch date", inputs.get("launch")], ["Places", inputs.get("locations")],
          ["Google Ads account", (rec.get("account_name") or rec.get("account_id")) if inputs.get("account") else ""],
          ["Meta ad account", (((research.get("metahistory") or {}).get("data") or {}).get("account") or {}).get("name") or inputs.get("metaAccount")],
          ["Built by", rec.get("created_by")], ["Built on", (plan.get("generatedAt") or "")[:10]]]
    ov += [["Summary", s] for s in st.get("summary") or []]
    out["Overview"] = (["Item", "Detail"], [r for r in ov if r[1]])

    fc = plan.get("forecast")
    if fc:
        out["Forecast"] = (["Scenario", "Cost per click", "Conversion rate %", "Clicks", "Conversions", "Cost per conversion"],
                           [[s["scenario"], s["cpc"], s["cvr"], s["clicks"], s["conversions"], s["cpa"]] for s in fc["scenarios"]] +
                           [["Cost per click from", fc["cpcFrom"], "", "", "", ""], ["Conversion rate from", fc["cvrFrom"], "", "", "", ""], ["Note", fc["note"], "", "", "", ""]])

    rows = []
    for o in st.get("objectives") or []:
        rows.append(["Objective", o["objective"], "KPI: %s. Target: %s." % (o["kpi"], o["target"]), o["basis"]])
    for f in st.get("funnel") or []:
        rows.append(["Funnel", f["stage"], f["role"], "%s Channels: %s" % (f["message"], f["channels"])])
    for c in st.get("channels") or []:
        rows.append(["Channel", c["channel"], "%s%% of budget, %s, %s" % (c["share"] if c["share"] is not None else "?", c["priority"], c["phase"]), c["why"]])
    b = st.get("budget") or {}
    for a in b.get("allocation") or []:
        rows.append(["Budget", a["campaign"], "%s%% of budget, %s a day, %s a month" % (a["share"], _money(a["daily"], cur) or "n/a", _money(a["monthly"], cur) or "n/a"), a["note"]])
    for label, key in (("Pacing", "pacing"), ("When to scale", "scaleRule")):
        if b.get(key):
            rows.append(["Budget", label, b[key], ""])
    for x in st.get("bidding") or []:
        rows.append(["Bidding", x["phase"], x["strategy"], "Move on when: %s. %s" % (x["trigger"], x["note"])])
    t = st.get("targeting") or {}
    for loc in t.get("locations") or []:
        rows.append(["Targeting", "Place: " + loc["name"], loc["why"], ""])
    for label, key in (("Devices", "devices"), ("Schedule", "schedule"), ("Demographics", "demographics")):
        if t.get(key):
            rows.append(["Targeting", label, t[key], ""])
    for a in t.get("audiences") or []:
        rows.append(["Targeting", "Audience", a, ""])
    m = st.get("messaging") or {}
    for p in m.get("pillars") or []:
        rows.append(["Messaging", p["pillar"], p["example"], "Proof: " + p["proof"]])
    for label, key in (("Tone", "tone"),):
        if m.get(key):
            rows.append(["Messaging", label, m[key], ""])
    for o in m.get("offers") or []:
        rows.append(["Messaging", "Offer to use", o, ""])
    for o in m.get("avoid") or []:
        rows.append(["Messaging", "Avoid", o, ""])
    for s in st.get("social") or []:
        rows.append(["Social insight", s["insight"], s["use"], ""])
    for s in st.get("seasonality") or []:
        rows.append(["Seasonality", s["window"], s["action"], ""])
    meas = st.get("measurement") or {}
    for c in meas.get("conversions") or []:
        rows.append(["Measurement", "Conversion: " + c["name"], "%s, %s" % (c["type"], "primary" if c["primary"] else "secondary"), c["how"]])
    for x in meas.get("tracking") or []:
        rows.append(["Measurement", x["item"], x["detail"], ""])
    if meas.get("reporting"):
        rows.append(["Measurement", "Reporting", meas["reporting"], ""])
    for x in st.get("tests") or []:
        rows.append(["Test", x["test"], x["hypothesis"], "Metric: %s. When: %s." % (x["metric"], x["when"])])
    for r in st.get("roadmap") or []:
        for a in r["actions"]:
            rows.append(["Roadmap", r["phase"], a, ""])
    for x in st.get("risks") or []:
        rows.append(["Risk", x["risk"], x["mitigation"], ""])
    for a in st.get("assumptions") or []:
        rows.append(["Assumption", a, "", ""])
    for a in st.get("clientNeeds") or []:
        rows.append(["Needed from the client", a, "", ""])
    if rows:
        out["Strategy"] = (["Section", "Item", "Detail", "Basis or note"], rows)

    rows = []
    for o in prof.get("offerings") or []:
        rows.append(["Offering", o["name"], o["description"], o["priority"]])
    for a in prof.get("audiences") or []:
        rows.append(["Audience", a["name"], "%s Need: %s" % (a["who"], a["need"]), a["intent"]])
    pos = prof.get("positioning") or {}
    if pos.get("promise"):
        rows.append(["Positioning", "Promise", pos["promise"], ""])
    for d in pos.get("differentiators") or []:
        rows.append(["Positioning", "What sets it apart", d, ""])
    if pos.get("tone"):
        rows.append(["Positioning", "Voice", pos["tone"], ""])
    for p in prof.get("proof") or []:
        rows.append(["Proof", p["point"], p["source"], ""])
    for p in prof.get("conversionPaths") or []:
        rows.append(["Conversion path", p["action"], p["where"], p["note"]])
    for c in prof.get("compliance") or []:
        rows.append(["Compliance", c["topic"], c["rule"], c["impact"]])
    for g in prof.get("siteGaps") or []:
        rows.append(["Site gap", g["issue"], g["why"], "Fix: " + g["fix"]])
    for f in prof.get("facts") or []:
        rows.append(["Fact", f["fact"], f["quote"], ("On the page: " if f["verified"] else "NOT FOUND ON A PAGE, confirm with the client: ") + f["source"]])
    for a in prof.get("assumptions") or []:
        rows.append(["Assumption", a, "", ""])
    for q in prof.get("questions") or []:
        rows.append(["Question for the client", q, "", ""])
    if prof.get("summary"):
        rows.insert(0, ["Summary", prof["summary"], "", ""])
    if rows:
        out["Business"] = (["Section", "Item", "Detail", "Note"], rows)

    if comp.get("competitors"):
        rows = [[c["name"], c["url"], c["positioning"], _join(c["offers"]), _join(c["messages"]), _join(c["ctas"]), _join(c["strengths"]), _join(c["weaknesses"]), _join(c["themes"])]
                for c in comp["competitors"]]
        out["Competitors"] = (["Competitor", "Website", "Positioning", "Offers", "Messages", "Calls to action", "Strengths", "Weaknesses", "Search themes"], rows)
        rows = [["Landscape", comp.get("landscape"), ""]] if comp.get("landscape") else []
        rows += [["Open space", w["opportunity"], w["why"]] for w in comp.get("whiteSpace") or []]
        rows += [["How to stand apart", d, ""] for d in comp.get("differentiation") or []]
        rows += [["Watch out", d, ""] for d in comp.get("watchOuts") or []]
        if rows:
            out["Competitive read"] = (["Section", "Finding", "Why"], rows)

    soc = (research.get("social") or {}).get("data") or {}
    if soc.get("platforms") or soc.get("crossPlatform"):
        rows = [["Platform", "%s %s, %s posts" % (p["platform"], p.get("handle") or "", p.get("posts")), _join(p.get("summary"), " | "), _join(p.get("messaging"), " | ")]
                for p in soc.get("platforms") or [] if p.get("posts")]
        cp = soc.get("crossPlatform") or {}
        rows += [["Across platforms", "", _join(cp.get("summary"), " | "), _join(cp.get("messaging"), " | ")]] if cp.get("summary") or cp.get("messaging") else []
        rows += [["Top post", "%s %s" % (p["platform"], p.get("type") or ""), p.get("caption"), "Engagement %s" % p.get("engagement")] for p in soc.get("topPosts") or []]
        out["Social"] = (["Section", "Platform", "What it posts", "Messaging and strategy"], rows)

    hist = (research.get("history") or {}).get("data") or {}
    if hist.get("campaigns"):
        out["Account history"] = (["Campaign", "Type", "Status", "Bid strategy", "Daily budget", "Clicks", "Cost", "Conversions", "Cost per conversion"],
                                  [[c["name"], c["type"], c["status"], c["bid"], c["daily"], c["clicks"], c["cost"], c["conv"], c["cpa"]] for c in hist["campaigns"]])
        rows = [["Best keyword", k["text"], k["match"], k["campaign"], k["clicks"], k["cost"], k["conv"]] for k in hist.get("keywordsBest") or []]
        rows += [["Wasted keyword", k["text"], k["match"], k["campaign"], k["clicks"], k["cost"], k["conv"]] for k in hist.get("keywordsWasted") or []]
        t = hist.get("terms") or {}
        rows += [["Converting search term", x["term"], "", x["campaign"], x["clicks"], x["cost"], x["conv"]] for x in t.get("converting") or []]
        rows += [["Wasted search term", x["term"], "", x["campaign"], x["clicks"], x["cost"], 0] for x in t.get("wasteful") or []]
        rows += [["Device", d["device"], "", "", d["clicks"], d["cost"], d["conv"]] for d in hist.get("devices") or []]
        rows += [["Place", r["name"], "", "", r["clicks"], r["cost"], r["conv"]] for r in hist.get("regions") or []]
        out["Account detail"] = (["What", "Item", "Match", "Campaign", "Clicks", "Cost", "Conversions"], rows)

    mh = (research.get("metahistory") or {}).get("data") or {}
    if mh.get("campaigns"):
        rows = [["Campaign", c["name"], c["type"], c["cost"], c["conv"], c.get("cpa"), c["impr"], c["clicks"]] for c in mh["campaigns"]]
        rows += [["Placement", p["name"], "", p["cost"], p["conv"], p.get("cpa"), p["impr"], p["clicks"]] for p in mh.get("placements") or []]
        rows += [["Device", d["name"], "", d["cost"], d["conv"], d.get("cpa"), d["impr"], d["clicks"]] for d in mh.get("devices") or []]
        rows += [["Live ad", a["name"], a.get("format") or "", "", "", "", "", ("%s | %s | %s" % (a["title"], a["body"], a["cta"])).strip(" |")] for a in mh.get("ads") or []]
        out["Meta account history"] = (["What", "Item", "Objective or type", "Cost", "Results", "Cost per result", "Impressions", "Clicks or text"], rows)
    ma = (research.get("metaaudience") or {}).get("data") or {}
    if ma.get("candidates"):
        out["Meta interests"] = (["Interest", "Audience size", "Where Meta lists it", "Found from", "Aimed at by ad sets", "Spent", "Results"],
                                 [[c["name"], c.get("sizeText") or "", " > ".join(c.get("path") or []), c.get("from") or "", (c.get("inUse") or {}).get("adsets", ""), (c.get("inUse") or {}).get("cost", ""),
                                   (c.get("inUse") or {}).get("conv", "")] for c in ma["candidates"]])

    kw = (research.get("keywords") or {}).get("data") or {}
    if kw.get("ideas"):
        out["Keyword research"] = (["Keyword", "Searches a month", "Competition", "Low bid", "High bid", "Peak month", "Lowest month", "Found from"],
                                   [[i["text"], i["avg"], i["competition"], i["lowBid"], i["highBid"], i["peak"], i["low"], i["from"]] for i in kw["ideas"]])

    mk = (research.get("market") or {}).get("data") or {}
    rows = []
    for label, key in (("Demand", "demand"), ("How people choose", "decisionFactors"), ("Trend", "trends"), ("Cost signal", "costSignals")):
        rows += [[label, x, ""] for x in mk.get(key) or []]
    rows += [["Advertising rule", p["rule"], p["appliesTo"]] for p in mk.get("adPolicy") or []]
    rows += [["Source", s, ""] for s in mk.get("sources") or []]
    if rows:
        out["Market"] = (["Section", "Finding", "Applies to"], rows)

    if camps:
        out["Campaigns"] = (["Campaign", "Type", "Share of budget %", "Daily budget", "Monthly budget", "Purpose", "Landing page", "Bid strategy", "Locations", "Audience", "Notes"],
                            [[c["name"], c["kindLabel"], c["share"], c.get("daily"), c.get("monthly"), c["purpose"], c["landingPage"], c["settings"]["bidStrategy"]["value"],
                              c["settings"]["locations"]["value"], c["audience"], c["notes"]] for c in camps])
    ag, ng, ad, asx, theme, mset, mads = [], [], [], [], [], [], []
    for c in camps:
        bld = c.get("build") or {}
        for g in bld.get("adSets") or []:
            checked = {x["name"]: x for x in g.get("interestsChecked") or []}
            mset.append([c["name"], g["name"], g.get("audienceType") or "", ", ".join(("%s (%s)" % (n, checked[n]["sizeText"]) if checked.get(n, {}).get("sizeText") else n + ("" if checked.get(n, {}).get("verified", True) else " (not checked)")) for n in g.get("interests") or []),
                         _join(g.get("behaviors")), _join(g.get("customAudiences")), _join(g.get("lookalikes")), _join(g.get("excludes")), g.get("budgetShare"),
                         round(c["daily"] * g["budgetShare"] / 100.0, 2) if c.get("daily") else "", g.get("placements") or "", g.get("landingPage") or ""])
            for i, a in enumerate(g.get("ads") or [], 1):
                mads.append([c["name"], g["name"], a.get("name") or "Ad %d" % i, a["format"], a["primaryText"], len(a["primaryText"]), a["headline"], len(a["headline"]), a.get("description") or "",
                             len(a.get("description") or ""), a.get("cta") or "", a.get("finalUrl") or "", a.get("concept") or "", a.get("script") or ""])
        for g in bld.get("adGroups") or []:
            for k in g.get("keywords") or []:
                ag.append([c["name"], g["name"], g.get("theme") or "", g.get("intent") or "", g.get("landingPage") or "", k["text"], k["match"].title(), k.get("intent") or "", k.get("note") or ""])
            for i, a in enumerate(g.get("ads") or [], 1):
                for kind, items, lim in (("Headline", a.get("headlines") or [], LIMITS["headline"]), ("Description", a.get("descriptions") or [], LIMITS["description"])):
                    for j, t in enumerate(items, 1):
                        pin = next((p["position"] for p in a.get("pins") or [] if p["text"] == t), "")
                        ad.append([c["name"], g["name"], "Ad %d" % i, kind, j, t, len(t), lim, "Pinned to position %s" % pin if pin else ""])
                for kind, t in (("Path 1", a.get("path1")), ("Path 2", a.get("path2"))):
                    if t:
                        ad.append([c["name"], g["name"], "Ad %d" % i, kind, "", t, len(t), LIMITS["path"], ""])
        for g in bld.get("assetGroups") or []:
            for kind, items, lim in (("Headline", g.get("headlines") or [], LIMITS["headline"]), ("Long headline", g.get("longHeadlines") or [], LIMITS["long_headline"]),
                                     ("Description", g.get("descriptions") or [], LIMITS["description"])):
                for j, t in enumerate(items, 1):
                    ad.append([c["name"], g["name"], "Asset group", kind, j, t, len(t), lim, ""])
            if g.get("businessName"):
                ad.append([c["name"], g["name"], "Asset group", "Business name", "", g["businessName"], len(g["businessName"]), LIMITS["business_name"], ""])
            for im in g.get("images") or []:
                asx.append([c["name"], g["name"], "Image brief", im["ratio"], im["concept"], ""])
            for v in g.get("videos") or []:
                asx.append([c["name"], g["name"], "Video brief", "", v["concept"], ""])
            s = g.get("signals") or {}
            for label, key in (("Audience signal: custom segment", "customSegments"), ("Audience signal: interest", "interests"), ("Audience signal: your data", "yourData")):
                for x in s.get(key) or []:
                    asx.append([c["name"], g["name"], label, "", x, ""])
            if s.get("demographics"):
                asx.append([c["name"], g["name"], "Audience signal: demographics", "", s["demographics"], ""])
            for x in g.get("searchThemes") or []:
                theme.append([c["name"], g["name"], x, len(x)])
        if c["kind"] not in ("search", "pmax"):
            for g in bld.get("adGroups") or []:
                asx.append([c["name"], g["name"], "Audiences", "", _join(g.get("audiences")), _join(g.get("placements"))])
                for i, a in enumerate(g.get("ads") or [], 1):
                    for kind, items, lim in (("Headline", a.get("headlines") or [], LIMITS["dg_headline"] if c["kind"] == "demandgen" else LIMITS["headline"]),
                                             ("Description", a.get("descriptions") or [], LIMITS["description"])):
                        for j, t in enumerate(items, 1):
                            ad.append([c["name"], g["name"], "Ad %d (%s)" % (i, a.get("format") or "ad"), kind, j, t, len(t), lim, ""])
                    if a.get("longHeadline"):
                        ad.append([c["name"], g["name"], "Ad %d (%s)" % (i, a.get("format") or "ad"), "Long headline", "", a["longHeadline"], len(a["longHeadline"]), LIMITS["long_headline"], ""])
                    asx.append([c["name"], g["name"], "Creative concept, ad %d" % i, a.get("cta") or "", a.get("concept") or "", a.get("script") or ""])
        a = bld.get("assets") or {}
        for s in a.get("sitelinks") or []:
            asx.append([c["name"], "", "Sitelink", s["text"], "%s | %s" % (s["desc1"], s["desc2"]), s["url"] + ("" if s["verified"] else "  (page not seen on the site)")])
        for x in a.get("callouts") or []:
            asx.append([c["name"], "", "Callout", x, "", ""])
        for s in a.get("snippets") or []:
            asx.append([c["name"], "", "Structured snippet", s["header"], _join(s["values"], ", "), ""])
        if (a.get("call") or {}).get("use"):
            asx.append([c["name"], "", "Call asset", "", a["call"].get("note") or "", ""])
        for p in a.get("promotions") or []:
            asx.append([c["name"], "", "Promotion", p["text"], p["detail"], ""])
        n = bld.get("negatives") or {}
        for x in n.get("campaign") or []:
            ng.append([c["name"], "Campaign", x["text"], x["match"].title(), x["reason"]])
        for lst in n.get("lists") or []:
            for x in lst["keywords"]:
                ng.append([c["name"], "List: " + lst["name"], x, "Phrase", lst["reason"]])
    if ag:
        out["Ad groups and keywords"] = (["Campaign", "Ad group", "Theme", "Intent", "Landing page", "Keyword", "Match type", "Keyword intent", "Note"], ag)
    if ng:
        out["Negatives"] = (["Campaign", "Level", "Negative keyword", "Match type", "Reason"], ng)
    if ad:
        out["Ad copy"] = (["Campaign", "Ad group", "Ad", "Part", "No.", "Text", "Characters", "Limit", "Note"], ad)
    if mset:
        out["Meta ad sets"] = (["Campaign", "Ad set", "Audience type", "Interests", "Behaviors", "Custom audiences", "Lookalikes", "Excluded", "Share of the campaign %", "Daily budget", "Placements", "Landing page"], mset)
    if mads:
        out["Meta ads"] = (["Campaign", "Ad set", "Ad", "Format", "Primary text", "Characters (limit %d)" % 125, "Headline", "Characters (limit %d)" % 40, "Description", "Characters (limit %d)" % 30, "Button",
                            "Final URL", "Creative concept", "Script"], mads)
    if asx:
        out["Assets"] = (["Campaign", "Group", "Asset", "Detail", "Text", "Address or note"], asx)
    if theme:
        out["Search themes"] = (["Campaign", "Asset group", "Search theme", "Characters"], theme)
    if plan.get("claims"):
        out["Claims to confirm"] = (["Campaign", "Where", "Text", "Claim not found on the client's pages"], [[c.get("campaign"), c["where"], c["text"], c["claim"]] for c in plan["claims"]])
    chk = plan.get("checks") or {}
    rows = [["Dropped for being over a limit", "%s: %d" % (k, v)] for k, v in (chk.get("dropped") or {}).items()]
    rows += [["Negative removed", "%s [%s] would block the keyword '%s' (%s)" % (r["negative"], r["match"], r["blocked"], r["campaign"])] for r in chk.get("removedNegatives") or []]
    rows += [["Short", s] for s in chk.get("shortfalls") or []]
    rows += [["Meta interests", s] for s in chk.get("interestNotes") or []]
    if chk.get("learningNotChecked"):
        rows.append(["Meta learning bar not checked", "%d Meta campaign%s: the account has no cost per result for the objective, or the objective has no result to count." % (chk["learningNotChecked"], "" if chk["learningNotChecked"] == 1 else "s")])
    rows += [["Not available", s] for s in chk.get("missing") or []]
    if chk.get("unconfirmedFacts"):
        rows.append(["Facts not found on a page", str(chk["unconfirmedFacts"])])
    if rows:
        out["Checks"] = (["Check", "Detail"], rows)
    return out


def book_sheet_name(camp, used):
    base = SHEET_NAME_BAD.sub("", "Book %s" % camp["name"].replace("P2_", ""))[:31] or "Book"
    name, n = base, 2
    while name in used:
        name = base[:28] + "-%d" % n
        n += 1
    return name


# ── Excel ───────────────────────────────────────────────────────────────────

def xlsx(rec):
    """The workbook as bytes: the plan's tables, then one Campaign Book sheet per campaign."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head_fill, head_font = PatternFill("solid", fgColor="1F2A44"), Font(bold=True, color="FFFFFF")
    sec_fill = PatternFill("solid", fgColor="E7ECF5")
    wrap = Alignment(wrap_text=True, vertical="top")
    used = set()

    def sheet(name, header, rows, widths=None, section_col=None):
        name = SHEET_NAME_BAD.sub("", name)[:31] or "Sheet"
        base, n = name, 2
        while name in used:
            name, n = base[:28] + "-%d" % n, n + 1
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append([plain(h) for h in header])
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, wrap
        for r in rows:
            ws.append([plain(v) for v in r])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"
                c.alignment = wrap
        for i in range(1, len(header) + 1):
            longest = max([len(str(header[i - 1]))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None] or [10])
            ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, min(60, max(12, longest + 2)))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
        return ws

    for name, (header, rows) in tables(rec).items():
        sheet(name, header, rows)
    for camp in (rec.get("plan") or {}).get("campaigns") or []:
        if not camp.get("book"):
            continue
        ws = sheet(book_sheet_name(camp, used), book.HEADER, book.sheet_rows(camp["book"])[1:], widths={1: 38, 2: 60, 3: 60})
        for row in ws.iter_rows(min_row=2):
            if row[0].value and not row[1].value and not row[2].value:
                for c in row:
                    c.fill, c.font = sec_fill, Font(bold=True)
        ws.auto_filter.ref = None
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ── CSV ─────────────────────────────────────────────────────────────────────

def _csv_text(header, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow([safe(h) for h in header])
    for r in rows:
        w.writerow([safe(v) for v in r])
    return buf.getvalue()


def csv_book(rec):
    """The Campaign Book of every campaign in one table: Campaign, Section, Field, Value, Other things to check."""
    rows = book.all_rows((rec.get("plan") or {}).get("campaigns") or [])
    return _csv_text(rows[0], rows[1:])


def csv_campaign(rec, key):
    """One campaign's book in the QA Agent's layout (Field | Value | Other things to check), or None when there is no such campaign."""
    for c in (rec.get("plan") or {}).get("campaigns") or []:
        if c["key"] == key and c.get("book"):
            rows = book.sheet_rows(c["book"])
            return _csv_text(rows[0], rows[1:])
    return None


def xlsx_campaign(rec, key):
    """One campaign's book alone as a workbook the QA Agent can read (one sheet, Field | Value | Other things to check)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    for c in (rec.get("plan") or {}).get("campaigns") or []:
        if c["key"] == key and c.get("book"):
            wb = Workbook()
            ws = wb.active
            ws.title = "Campaign Book"
            for r in book.sheet_rows(c["book"]):
                ws.append([plain(v) for v in r])
            for cell in ws[1]:
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F2A44")
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                if row[0].value and not row[1].value and not row[2].value:
                    for cell in row:
                        cell.font, cell.fill = Font(bold=True), PatternFill("solid", fgColor="E7ECF5")
            for col, w in (("A", 38), ("B", 60), ("C", 60)):
                ws.column_dimensions[col].width = w
            ws.freeze_panes = "A2"
            out = io.BytesIO()
            wb.save(out)
            return out.getvalue()
    return None


def csv_zip(rec):
    """Every table of the plan as its own CSV, and every campaign's book, in one zip."""
    out = io.BytesIO()
    used = set()

    def name(base):
        stem = re.sub(r"[^A-Za-z0-9]+", "-", base).strip("-").lower() or "table"
        n, cand = 2, stem
        while cand in used:
            cand, n = "%s-%d" % (stem, n), n + 1
        used.add(cand)
        return cand + ".csv"

    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for sheet, (header, rows) in tables(rec).items():
            z.writestr(name(sheet), _csv_text(header, rows))
        z.writestr(name("campaign book all campaigns"), csv_book(rec))
        for c in (rec.get("plan") or {}).get("campaigns") or []:
            body = csv_campaign(rec, c["key"])
            if body:
                z.writestr(name("campaign book " + c["name"]), body)
    return out.getvalue()


# ── PDF ─────────────────────────────────────────────────────────────────────

def pdf(rec):
    """The strategy and the campaign book as a document."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from .event_intel_pdf import BODY, BOLD, _fonts, text_markup
    _fonts()
    plan = rec.get("plan") or {}
    inputs = rec.get("inputs") or {}
    ink, mute, rule = colors.HexColor("#1B2333"), colors.HexColor("#5B6577"), colors.HexColor("#D7DDE8")
    base = ParagraphStyle("b", fontName=BODY, fontSize=8.6, leading=11.8, textColor=ink)
    small = ParagraphStyle("s", parent=base, fontSize=7.6, leading=10, textColor=mute)
    h1 = ParagraphStyle("h1", parent=base, fontName=BOLD, fontSize=17, leading=21, spaceBefore=4, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=base, fontName=BOLD, fontSize=12, leading=15, spaceBefore=10, spaceAfter=4)
    h3 = ParagraphStyle("h3", parent=base, fontName=BOLD, fontSize=9.6, leading=12.5, spaceBefore=6, spaceAfter=2)
    head = ParagraphStyle("th", parent=base, fontName=BOLD, textColor=colors.white, fontSize=8)
    bullet = ParagraphStyle("bl", parent=base, leftIndent=10, bulletIndent=1, spaceAfter=1.5)
    page = landscape(A4)
    width = page[0] - 28 * mm

    def P(t, st=base):
        return Paragraph(text_markup(str(t if t is not None else "")), st)

    def bullets(items):
        return [Paragraph(text_markup(str(x)), bullet, bulletText="•") for x in items if x]

    sechead = ParagraphStyle("sh", parent=base, fontName=BOLD, textColor=colors.HexColor("#1F2A44"))

    def table(header, rows, widths, sections=()):
        """`sections` are the numbers (from 0) of rows that are a heading for the rows below: one cell across the whole width."""
        total = sum(widths)
        cw = [width * w / total for w in widths]
        data = [[Paragraph(text_markup(h), head) for h in header]]
        for i, r in enumerate(rows):
            data.append([Paragraph(text_markup(str(r[0])), sechead)] + [""] * (len(header) - 1) if i in sections else [P(str(c)[:CELL_CHARS]) for c in r])
        t = Table(data, colWidths=cw, repeatRows=1, splitInRow=1)                    # a row taller than a page carries on over the next one
        style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A44")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                 ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FA")]), ("LINEBELOW", (0, 0), (-1, -1), 0.25, rule),
                 ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        for i in sections:
            style += [("SPAN", (0, i + 1), (-1, i + 1)), ("BACKGROUND", (0, i + 1), (-1, i + 1), colors.HexColor("#E7ECF5"))]
        t.setStyle(TableStyle(style))
        return t

    tb = tables(rec)
    story = [P("Strategy and Campaign Book", h1), P(inputs.get("name"), h2),
             P("%s | %s | goal: %s | launch %s%s" % (inputs.get("website") or "no website", inputs.get("industry") or "", ps.GOALS.get(inputs.get("goal"), ""), inputs.get("launch"),
                                                 (" | budget %s a month" % _money(inputs.get("budget"), inputs.get("currency"))) if inputs.get("budget") else ""), small), Spacer(1, 6)]
    st = plan.get("strategy") or {}
    story += bullets(st.get("summary") or [])
    contents = ["Strategy", "Business", "Competitors", "Competitive read", "Social", "Account history", "Meta account history", "Keyword research", "Market", "Campaigns"]
    story += [Spacer(1, 6), P("Contents", h3), P(", ".join(n for n in contents if n in tb) + ", the Campaign Book of each campaign, and what needs confirming.", small)]
    fc = plan.get("forecast")
    if fc:
        scenes = [r for r in tb["Forecast"][1] if r[0] in ("Cautious", "Expected", "Strong")]
        story += [P("What the budget buys", h2), table(tb["Forecast"][0], scenes, [1.2, 1, 1.2, 1, 1, 1.4]),
                  P("Cost per click: %s." % fc["cpcFrom"], small), P("Conversion rate: %s." % fc["cvrFrom"], small), P(fc["note"], small)]
    for name in ("Strategy", "Business", "Competitors", "Competitive read", "Social", "Account history", "Account detail", "Meta account history", "Keyword research", "Market"):
        if name not in tb:
            continue
        header, rows = tb[name]
        widths = {"Strategy": [1.1, 2, 3.4, 3.4], "Business": [1.3, 2.4, 3, 3], "Competitors": [1.2, 1.3, 2.2, 1.7, 1.7, 1.2, 1.7, 1.7, 1.5], "Competitive read": [1.3, 4, 3],
                  "Social": [1.2, 1.8, 3.4, 3.4], "Account history": [3.2, 1.2, 1, 1.5, 1, 1, 1, 1, 1], "Account detail": [1.8, 3, 1, 2.6, 1, 1, 1],
                  "Keyword research": [4, 1.5, 1.4, 1, 1, 1.2, 1.2, 1.3], "Market": [1.5, 5, 2.4], "Meta account history": [1.2, 2.4, 1.4, 1, 1, 1.2, 1.2, 3]}[name]
        story += [PageBreak() if name in ("Strategy", "Business") else Spacer(1, 4), P(name, h2), table(header, rows[:400], widths)]
    camps = plan.get("campaigns") or []
    if camps:
        story += [PageBreak(), P("Campaigns", h1), table(*_cols(tb["Campaigns"], [3, 1.2, 1, 1, 1, 3, 2.4, 2.4, 2.4, 2, 2]))]
    for c in camps:
        story += [PageBreak(), P(c["name"], h1), P("%s | %s%% of the budget | %s a day | landing page: %s" % (c["kindLabel"], c["share"], _money(c.get("daily"), inputs.get("currency")) or "budget not set",
                                                                                                      c.get("landingPage") or "to be confirmed"), small), P(c["purpose"]), Spacer(1, 4)]
        bld = c.get("build") or {}
        if not bld:
            story.append(P("This campaign could not be written: %s" % (c.get("buildError") or "no reason was given"), base))
        for g in bld.get("adGroups") or []:
            story += [P("%s%s" % (g["name"], " (%s)" % g["intent"] if g.get("intent") else ""), h3)]
            if g.get("keywords"):
                story.append(P("Keywords: " + ", ".join("%s [%s]" % (k["text"], k["match"][0].upper()) for k in g["keywords"]), small))
            for i, a in enumerate(g.get("ads") or [], 1):
                story += [P("Ad %d" % i, small)] + bullets(["Headlines: " + " | ".join(a.get("headlines") or []), "Descriptions: " + " | ".join(a.get("descriptions") or [])])
            if g.get("audiences"):
                story.append(P("Audiences: " + ", ".join(g["audiences"]), small))
            for a in g.get("ads") or []:
                if isinstance(a, dict) and a.get("concept"):
                    story += bullets(["Creative: " + a["concept"]])
        for g in bld.get("adSets") or []:
            story += [P("%s (%s, %s%% of the campaign)" % (g["name"], g.get("audienceType") or "interests", g.get("budgetShare")), h3)]
            aud = [x for x in (("Interests", g.get("interests")), ("Behaviors", g.get("behaviors")), ("Custom audiences", g.get("customAudiences")), ("Lookalikes", g.get("lookalikes")), ("Excluded", g.get("excludes"))) if x[1]]
            story += bullets(["%s: %s" % (k, ", ".join(v)) for k, v in aud])
            for i, a in enumerate(g.get("ads") or [], 1):
                story += [P("Ad %d (%s)" % (i, a["format"]), small)] + bullets(["Primary text: " + a["primaryText"], "Headline: " + a["headline"] + ((" | Description: " + a["description"]) if a.get("description") else "") +
                                                                           ((" | Button: " + a["cta"]) if a.get("cta") else ""), ("Creative: " + a["concept"]) if a.get("concept") else ""])
        for g in bld.get("assetGroups") or []:
            story += [P(g["name"], h3)] + bullets(["Headlines: " + " | ".join(g["headlines"]), "Long headlines: " + " | ".join(g["longHeadlines"]), "Descriptions: " + " | ".join(g["descriptions"]),
                                                  "Search themes: " + ", ".join(g["searchThemes"])])
        n = bld.get("negatives") or {}
        if n.get("campaign"):
            story += [P("Negative keywords", h3), P(", ".join('%s [%s]' % (x["text"], x["match"]) for x in n["campaign"]), small)]
        a = bld.get("assets") or {}
        if a.get("sitelinks") or a.get("callouts"):
            story += [P("Assets", h3)] + bullets(["Sitelinks: " + " | ".join("%s (%s)" % (s["text"], s["url"]) for s in a.get("sitelinks") or []),
                                                  "Callouts: " + " | ".join(a.get("callouts") or []),
                                                  "Snippets: " + " | ".join("%s: %s" % (s["header"], ", ".join(s["values"])) for s in a.get("snippets") or [])])
        if c.get("book"):
            rows, heads, last = [], [], None
            for r in c["book"]:
                if r["section"] != last:
                    heads.append(len(rows))
                    rows.append([r["section"], "", ""])
                    last = r["section"]
                rows.append([r["field"], r["value"], r["other"]])
            story += [PageBreak(), P("Campaign Book: " + c["name"], h2), table(book.HEADER, rows, [2.2, 4, 3.6], sections=heads)]
    if "Claims to confirm" in tb:
        story += [PageBreak(), P("Claims to confirm with the client before launch", h2), P("These are in the ad text but not on the client's pages or in the brief.", small),
                  table(*_cols(tb["Claims to confirm"], [2, 3, 4, 2]))]
    story += [PageBreak(), P("Open questions and what is needed from the client", h2)]
    story += bullets((plan.get("profile") or {}).get("questions") or []) + bullets(st.get("clientNeeds") or []) + bullets(st.get("assumptions") or [])
    out = io.BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(BODY, 7)
        canvas.setFillColor(mute)
        canvas.drawString(14 * mm, 8 * mm, "Position2: %s. Strategy and Campaign Book" % clean(inputs.get("name"), 60))
        canvas.drawRightString(page[0] - 14 * mm, 8 * mm, "Page %d" % doc.page)
        canvas.restoreState()

    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                            title="Strategy and Campaign Book: %s" % inputs.get("name"), author="Position2")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()


def _cols(tab, widths):
    header, rows = tab
    return header, rows, widths
