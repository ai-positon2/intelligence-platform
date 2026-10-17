"""Paid Marketing Agents: Landing Page & Relevancy Check, the downloads.

One report, three files, all drawn from `tables(report)` so they cannot drift apart:

  Excel    every table as a sheet
  CSV      every table as its own CSV in a zip
  PDF      the report as a document to read and to send: the summary, the checks to put right first, every page's score, then each page's problems

A value that a spreadsheet would run as a formula (it starts with = + - @) is written as plain text in every format (tracker/paid_strategy_export.py).
The report sends its findings small, with the wording once in `catalog`; the files join the two so each row reads in full.
"""
from __future__ import annotations

import io
import re
import zipfile

from .paid_strategy_check import text as clean
from .paid_strategy_export import CELL_CHARS, SHEET_NAME_BAD, _csv_text, plain

RESULT_TEXT = {"pass": "Pass", "warn": "Warning", "fail": "Fail", "na": "Not applicable", "unknown": "Could not check"}
SEVERITY_TEXT = {"critical": "Critical", "high": "High", "medium": "Medium", "low": "Low"}
ON_TEXT = {"headline": "In the headline", "page": "On the page", "partial": "Part of it", "none": "Not on the page", None: "Not checked"}
RATING_TEXT = {"below": "Below average", "average": "Average", "above": "Above average", None: ""}
PDF_AFFECTED = 5
PDF_PAGES = 30                                     # pages whose problems the PDF lists (the Excel file has every one)


def file_stem(report):
    acct = report.get("account") or {}
    base = re.sub(r"[^A-Za-z0-9]+", "-", acct.get("name") or acct.get("id") or "")[:60].strip("-") or "account"
    return "%s-landing-page-check" % base


def _full(f, catalog):
    """A finding joined with its catalog wording."""
    c = catalog.get(f["id"], {})
    return dict(f, title=c.get("title", f["id"]), why=c.get("why", ""), fix=c.get("fix", ""), section=c.get("section", ""), severity=f.get("severity") or c.get("severity", ""))


def _affected(f, limit=None):
    shown = f.get("items") or []
    if not shown:
        return ""
    if limit is not None:
        shown = shown[:limit]
    more = f.get("count", len(f.get("items") or [])) - len(shown)
    return "; ".join(str(i) for i in shown) + (" (and %d more)" % more if more > 0 else "")


def _pct(x):
    return None if x is None else round(float(x), 2)


def _devices(p):
    d = p.get("devices") or {}
    out = {}
    for k in ("mobile", "desktop"):
        v = d.get(k) or {}
        out[k] = (100.0 * v["conversions"] / v["clicks"]) if v.get("clicks") else None
    return out


def tables(report):
    """{sheet name: (header, rows)} in the order they are shown. Empty tables are left out. Numbers stay numbers."""
    catalog = report.get("catalog") or {}
    labels = {s["key"]: s["label"] for s in report.get("sections") or []}
    summary = report.get("summary") or {}
    acct = report.get("account") or {}
    pages = report.get("pages") or []
    out = {}
    t = summary.get("tally") or {}
    ov = [["Google Ads account", "%s (%s)" % (acct.get("name") or "", acct.get("id") or "")], ["Window", report.get("window") or ""],
          ["Read on", (report.get("readAt") or "")[:16].replace("T", " ")], ["Report made on", (report.get("generatedAt") or "")[:10]],
          ["Overall score", summary.get("score") if summary.get("score") is not None else "Nothing could be judged"], ["Overall grade", summary.get("grade") or ""],
          ["Grade held at C or below because", "; ".join(summary.get("capReasons") or [])],
          ["Landing pages that took money", summary.get("pagesOfSpend")], ["Pages checked", summary.get("checked")], ["Pages with a score", summary.get("rated")],
          ["Spend on those pages (%s)" % (acct.get("currency") or ""), summary.get("spend")], ["Spend on pages checked", summary.get("spendChecked")],
          ["Spend on pages that could not be verified", summary.get("spendUnverified")], ["Spend on pages with a critical problem or a D or F", summary.get("spendAtRisk")],
          ["Checks passed", t.get("pass")], ["Warnings", t.get("warn")], ["Failed", t.get("fail")], ["Not applicable", t.get("na")], ["Could not check", t.get("unknown")],
          ["Failed critical checks", summary.get("critical")], ["Failed high checks", summary.get("high")]]
    for s in report.get("sections") or []:
        ov.append(["Section: %s" % s["label"], "%s (%s)" % (s["score"], s["grade"]) if s["score"] is not None else "Nothing could be judged"])
    out["Summary"] = (["Item", "Detail"], [r for r in ov if r[1] not in (None, "")])

    rows = []
    for p in pages:
        s, seen, dev = p.get("summary") or {}, p.get("seen") or {}, _devices(p)
        tally = s.get("tally") or {}
        rows.append([p["url"], p["spend"], p["clicks"], p["conversions"], _pct(p.get("convRate")), p.get("cpa"), p["score"] if p.get("rated") else "Not verified", p["grade"] if p.get("rated") else "",
                     tally.get("fail", 0), tally.get("warn", 0), tally.get("unknown", 0), seen.get("status"), seen.get("responseMs"), seen.get("words"), p.get("speed"),
                     _pct(dev["mobile"]), _pct(dev["desktop"]), (seen.get("title") or ""), "; ".join(seen.get("h1") or []), p.get("groupCount", 0)])
    if rows:
        out["Pages"] = (["Page", "Spend", "Clicks", "Conversions", "Conv rate %", "Cost per conversion", "Score", "Grade", "Failed", "Warnings", "Could not check", "HTTP status",
                         "Server ms", "Words", "Google speed score", "Mobile conv rate %", "Desktop conv rate %", "Title", "Main headline", "Ad groups"], rows)

    fix = report.get("fixFirst") or []
    if fix:
        fr = []
        for i, f in enumerate(fix):
            f = _full(f, catalog)
            fr.append([i + 1, f.get("pageName") or f.get("page"), f.get("spend"), labels.get(f["section"], f["section"]), "%s %s" % (f["id"], f["title"]), SEVERITY_TEXT.get(f["severity"], f["severity"]),
                       RESULT_TEXT.get(f["result"], f["result"]), f["found"], _affected(f), f["fix"]])
        out["Fix first"] = (["Order", "Page", "Page spend", "Section", "Check", "Severity", "Result", "What was found", "Affected", "What to do"], fr)

    allrows, unknown = [], []
    for p in pages:
        for sec in p.get("sections") or []:
            for f in sec["findings"]:
                f = _full(f, catalog)
                allrows.append([p["url"], sec["label"], "%s %s" % (f["id"], f["title"]), SEVERITY_TEXT.get(f["severity"], f["severity"]), RESULT_TEXT.get(f["result"], f["result"]), f["found"],
                                _affected(f), f["why"], f["fix"]])
                if f["result"] == "unknown":
                    unknown.append([p["url"], "%s %s" % (f["id"], f["title"]), f["found"]])
    if allrows:
        out["All checks"] = (["Page", "Section", "Check", "Severity", "Result", "What was found", "Affected", "Why it matters", "What to do"], allrows)
    if unknown:
        out["Could not check"] = (["Page", "Check", "Why not"], unknown)

    kw, terms, offers = [], [], []
    for p in pages:
        for g in p.get("adGroups") or []:
            for k in g.get("keywords") or []:
                kw.append([p["url"], g["campaign"], g["name"], k["text"], k["match"], k["cost"], k["clicks"], k["qs"], RATING_TEXT.get(k["landingPage"], ""), RATING_TEXT.get(k["adRelevance"], ""),
                           RATING_TEXT.get(k["expectedCtr"], ""), ON_TEXT.get(k["on"], "")])
            for x in g.get("terms") or []:
                terms.append([p["url"], g["campaign"], g["name"], x["text"], x["clicks"], x["cost"], x["conversions"], ON_TEXT.get(x["on"], "")])
            for o in g.get("offers") or []:
                offers.append([p["url"], g["campaign"], g["name"], o["text"], {True: "On the page", False: "Not on the page", None: "Not checked"}[o["on"]]])
    if kw:
        out["Keywords"] = (["Page", "Campaign", "Ad group", "Keyword", "Match", "Spend", "Clicks", "Quality score", "Landing page experience", "Ad relevance", "Expected click rate", "On the page"], kw)
    if terms:
        out["Search terms"] = (["Page", "Campaign", "Ad group", "Search term", "Clicks", "Spend", "Conversions", "On the page"], terms)
    if offers:
        out["Ad offers"] = (["Page", "Campaign", "Ad group", "Promise in the ad", "On the page"], offers)

    notes = [["Note", n] for n in report.get("notes") or []] + [["Part that could not be read", w] for w in report.get("warnings") or []]
    for st in report.get("stale") or []:
        notes.append(["Saved read", "The read is the one saved on %s because Google Ads would not answer now." % str(st.get("savedAt") or "")[:16].replace("T", " ")])
    if notes:
        out["Notes"] = (["Kind", "Detail"], notes)
    return out


def xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head_fill, head_font, wrap = PatternFill("solid", fgColor="1F2A44"), Font(bold=True, color="FFFFFF"), Alignment(wrap_text=True, vertical="top")
    fills = {"Fail": "FFD9D6", "Warning": "FFF1CC", "Pass": "DDF3E2", "Could not check": "E4E8F0", "Not on the page": "FFD9D6", "In the headline": "DDF3E2", "On the page": "DDF3E2",
             "Part of it": "FFF1CC"}
    used = set()
    for name, (header, rows) in tables(report).items():
        name = SHEET_NAME_BAD.sub("", name)[:31] or "Sheet"
        base, n = name, 2
        while name in used:
            name, n = base[:28] + "-%d" % n, n + 1
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append([plain(h) for h in header])
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, wrap
        for r in rows:
            ws.append([plain(v) for v in r])
        tint_cols = [i for i, h in enumerate(header) if h in ("Result", "On the page")]
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"
                c.alignment = wrap
            for i in tint_cols:
                if fills.get(row[i].value):
                    row[i].fill = PatternFill("solid", fgColor=fills[row[i].value])
        for i in range(1, len(header) + 1):
            longest = max([len(str(header[i - 1]))] + [min(len(str(r[i - 1])), 70) for r in rows if i - 1 < len(r) and r[i - 1] is not None] or [10])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def csv_zip(report):
    out = io.BytesIO()
    used = set()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for sheet, (header, rows) in tables(report).items():
            stem = re.sub(r"[^A-Za-z0-9]+", "-", sheet).strip("-").lower() or "table"
            cand, n = stem, 2
            while cand in used:
                cand, n = "%s-%d" % (stem, n), n + 1
            used.add(cand)
            z.writestr(cand + ".csv", _csv_text(header, rows))
    return out.getvalue()


def pdf_pages(report):
    """[(page url, score text, spend, rows)] for the PDF: each page's failures and warnings (what was found, what to do), the busiest pages first."""
    catalog = report.get("catalog") or {}
    out = []
    for p in (report.get("pages") or [])[:PDF_PAGES]:
        rows = []
        for sec in p.get("sections") or []:
            for f in sec["findings"]:
                if f["result"] not in ("fail", "warn"):
                    continue
                f = _full(f, catalog)
                rows.append([RESULT_TEXT.get(f["result"], f["result"]), SEVERITY_TEXT.get(f["severity"], f["severity"]), f["title"],
                             f["found"] + ((" Affected: " + _affected(f, PDF_AFFECTED)) if _affected(f) else ""), f["fix"]])
        score = ("%s (%s)" % (p["score"], p["grade"])) if p.get("rated") and p.get("score") is not None else "not verified"
        out.append((p["url"], score, p["spend"], rows))
    return out


def pdf(report):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from .event_intel_pdf import BODY, BOLD, _fonts, text_markup
    _fonts()
    ink, mute, rule = colors.HexColor("#1B2333"), colors.HexColor("#5B6577"), colors.HexColor("#D7DDE8")
    base = ParagraphStyle("b", fontName=BODY, fontSize=8.4, leading=11.4, textColor=ink)
    small = ParagraphStyle("s", parent=base, fontSize=7.6, leading=10, textColor=mute)
    h1 = ParagraphStyle("h1", parent=base, fontName=BOLD, fontSize=17, leading=21, spaceBefore=4, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=base, fontName=BOLD, fontSize=12, leading=15, spaceBefore=10, spaceAfter=4, keepWithNext=1)
    head = ParagraphStyle("th", parent=base, fontName=BOLD, textColor=colors.white, fontSize=7.8)
    page = landscape(A4)
    width = page[0] - 28 * mm
    tint = {"Fail": "#FFD9D6", "Warning": "#FFF1CC", "Pass": "#EAF6ED", "Could not check": "#ECEFF5"}

    def P(t, st=base):
        return Paragraph(text_markup(str(t if t is not None else "")), st)

    def table(header, rows, shade_col=None, weights=None):
        n = len(header)
        weights = weights or [max(1.0, min(4.0, max([len(str(header[i]))] + [min(len(str(r[i])), 200) / 12.0 for r in rows if i < len(r)] or [1]) / 6.0)) for i in range(n)]
        cw = [width * w / sum(weights) for w in weights]
        data = [[Paragraph(text_markup(h), head) for h in header]] + [[P(str(c)[:CELL_CHARS]) for c in r] for r in rows]
        t = Table(data, colWidths=cw, repeatRows=1, splitInRow=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A44")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.25, rule),
                 ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]
        if shade_col is None:
            style.append(("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FA")]))
        else:
            for i, r in enumerate(rows, start=1):
                if tint.get(r[shade_col]):
                    style.append(("BACKGROUND", (shade_col, i), (shade_col, i), colors.HexColor(tint[r[shade_col]])))
        t.setStyle(TableStyle(style))
        return t

    all_tables = tables(report)
    acct = report.get("account") or {}
    name = acct.get("name") or acct.get("id") or ""
    catalog = report.get("catalog") or {}
    story = [P("Landing Page & Relevancy Check", h1), P("%s (%s)" % (name, acct.get("id") or ""), h2),
             P("Read on %s | report made on %s | %s" % ((report.get("readAt") or "")[:16].replace("T", " "), (report.get("generatedAt") or "")[:10], report.get("window") or ""), small), Spacer(1, 6)]
    head_t, ov_rows = all_tables["Summary"]
    story += [P("Summary", h2), table(head_t, ov_rows)]
    if "Fix first" in all_tables:
        fix_rows = []
        for i, f in enumerate(report.get("fixFirst") or []):
            f = _full(f, catalog)
            fix_rows.append([i + 1, f.get("pageName") or f.get("page"), "%s %s" % (f["id"], f["title"]), SEVERITY_TEXT.get(f["severity"], f["severity"]), RESULT_TEXT.get(f["result"], f["result"]),
                             f["found"] + ((" Affected: " + _affected(f, PDF_AFFECTED)) if _affected(f) else ""), f["fix"]])
        story += [PageBreak(), P("Fix first", h2), table(["No.", "Page", "Check", "Severity", "Result", "What was found", "What to do"], fix_rows, shade_col=4,
                                                        weights=[0.5, 1.8, 1.8, 0.8, 0.8, 4.4, 3.0])]
    if "Pages" in all_tables:
        rows = [[r[0], r[1], r[2], r[6], r[7], r[8], r[9], r[10]] for r in all_tables["Pages"][1]]
        story += [PageBreak(), P("Every page checked", h2), table(["Page", "Spend", "Clicks", "Score", "Grade", "Failed", "Warnings", "Could not check"], rows, weights=[6.0, 0.9, 0.8, 0.9, 0.7, 0.7, 0.8, 1.0])]
    for url, score, spend, rows in pdf_pages(report):
        story += [PageBreak(), P("%s (score %s, spend %s)" % (url, score, spend), h2)]
        story += [table(["Result", "Severity", "Check", "What was found", "What to do"], rows, shade_col=0, weights=[0.9, 0.8, 2.2, 5.0, 3.2])] if rows else [P("Nothing failed and nothing is a warning on this page.")]
    if "Notes" in all_tables:
        h, rows = all_tables["Notes"]
        story += [PageBreak(), P("Notes", h2), table(h, rows)]
    out = io.BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(BODY, 7)
        canvas.setFillColor(mute)
        canvas.drawString(14 * mm, 8 * mm, "Position2: %s. Landing Page & Relevancy Check" % clean(name, 70))
        canvas.drawRightString(page[0] - 14 * mm, 8 * mm, "Page %d" % doc.page)
        canvas.restoreState()

    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                            title="Landing Page & Relevancy Check: %s" % name, author="Position2")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
