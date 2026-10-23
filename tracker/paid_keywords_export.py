"""Paid Marketing Agents: Keyword Research, the downloads.

One result, four files, all drawn from `tables(rec)` so they cannot drift apart:

  Excel    every table as a sheet
  CSV      every table as its own CSV in a zip, and one CSV for Google Ads Editor (Campaign, Ad group, Keyword, Criterion Type, Status): only the keywords judged a
           fit, new to the account and not a competitor's, every one Paused so nothing serves until a person has looked at it
  PDF      the result as a document to read and to send (without the full keyword list, which is in the Excel file)

A value that a spreadsheet would run as a formula (it starts with = + - @) is written as plain text in every format (tracker/paid_strategy_export.py).
"""
from __future__ import annotations

import io
import re
import zipfile

from . import paid_keywords_data as kd, paid_strategy as ps
from .paid_strategy_check import text as clean
from .paid_strategy_export import CELL_CHARS, SHEET_NAME_BAD, _csv_text, _join, plain

PDF_SKIP = {"Keywords", "Ads Editor import"}
PDF_ROW_CAP = {"Add to account": 120, "Already in the account": 40, "Meta interests": 60, "Meta ad sets": 40}
PDF_PAGE_BREAK = ("Themes", "Add to account", "Plan", "Meta audiences")
EDITOR_HEADER = ["Campaign", "Ad group", "Keyword", "Criterion Type", "Status"]
FIT_ORDER = {"yes": 0, "maybe": 1, "no": 2, "unjudged": 3}


def file_stem(rec):
    base = re.sub(r"[^A-Za-z0-9]+", "-", rec.get("name") or "client").strip("-")[:50] or "client"
    return "%s-keyword-research" % base


def status_text(r):
    if r.get("status") == "bidding":
        by = r.get("by") or {}
        return "Already a keyword (%s match in %s)" % ((by.get("match") or "").lower() or "any", by.get("campaign") or "the account")
    if r.get("status") == "covered":
        by = r.get("by") or {}
        return "Matched by \"%s\" (%s) in %s" % (by.get("kw"), (by.get("match") or "").lower(), by.get("campaign") or "the account")
    return "New to the account"


def editor_text(r, unread=False):
    """Whether the keyword is in the file for Google Ads Editor, and when it is not though it fits, why."""
    if r.get("include"):
        return "Yes"
    if unread and r["fit"] == "yes":
        return "No, the account's keywords could not be read"
    if r.get("blockedInTarget"):
        return "No, the campaign it would go in already blocks it with a negative keyword"
    if r["fit"] == "yes" and r["intent"] != "competitor" and (r["status"] == "new" or r.get("conv")) and not r.get("editorOk", True):
        return "No, it holds a symbol Google Ads does not accept in a keyword"
    return "No"


def blocked_text(r):
    n = r.get("neg") or {}
    if not n.get("scope"):
        return ""
    return "Blocked in some campaigns by \"%s\"" % ((n.get("by") or [{}])[0].get("kw"))


def pk_platforms(inputs, plan):
    """(google, meta): which platforms this saved result has something for (an older build has no platform: it is Google)."""
    p = (inputs or {}).get("platform") or "google"
    return p in ("google", "both"), p in ("meta", "both")


def _money_cell(v):
    return v if v is not None else ""


def meta_tables(meta):
    """The Meta sheets, each only when it has rows."""
    out = {}
    cur = meta.get("currency") or "USD"
    rows = meta.get("interests") or []
    aud = meta.get("audiences") or []
    if aud:
        out["Meta audiences"] = (["Audience", "Role", "Interests", "Largest interest (people)", "Why", "How to test it"],
                                 [[a["name"], a["role"], ", ".join(("%s (%s)" % (i["name"], i["sizeText"]) if i.get("sizeText") else i["name"]) for i in a["interests"]), a.get("largest") or "", a["why"], a["test"]] for a in aud])
    if rows:
        order = {"yes": 0, "maybe": 1, "no": 2, "unjudged": 3}
        ordered = sorted(rows, key=lambda r: (order.get(r["fit"], 4), 0 if not r.get("inUse") else 1, -(r.get("size") or 0), r["name"]))
        out["Meta interests"] = (["Interest", "Type", "Fit", "Intent", "Audience", "Audience size", "Aimed at by ad sets", "Spent (%s)" % cur, "Results", "Cost per result (%s)" % cur, "Found from", "Where Meta lists it", "Note"],
                                 [[r["name"], r["kind"], r["fit"], r["intent"], r["theme"], r.get("sizeText") or "", (r.get("inUse") or {}).get("adsets", ""), _money_cell((r.get("inUse") or {}).get("cost")),
                                   _money_cell((r.get("inUse") or {}).get("conv")), _money_cell((r.get("inUse") or {}).get("cpa")), r.get("from") or "", " > ".join(r.get("path") or []), r["note"]] for r in ordered])
    if meta.get("inUse"):
        out["Meta in use"] = (["Interest", "Type", "Ad sets", "Spent (%s)" % cur, "Results", "Cost per result (%s)" % cur, "How it is doing", "Why", "Judged"],
                              [[u["name"], u["kind"], u["adsets"], u["cost"], u["conv"], _money_cell(u.get("cpa")), u["verdict"], u["why"], u.get("fit") or ""] for u in meta["inUse"]])
    if meta.get("adsets"):
        out["Meta ad sets"] = (["Ad set", "Campaign", "Status", "Interests", "Behaviors", "Audiences", "Excluded", "Advantage+ audience", "Spent (%s)" % cur, "Results", "Cost per result (%s)" % cur],
                              [[a["name"], a.get("campaign") or "", a.get("status") or "", ", ".join(i["name"] for i in a["interests"]), ", ".join(i["name"] for i in a["behaviors"]), ", ".join(a["audiences"]),
                                ", ".join(a["excluded"]), "Yes" if a.get("advantage") else "No", a["cost"], a["conv"], _money_cell(a.get("cpa"))] for a in meta["adsets"]])
    if meta.get("exclusions"):
        out["Meta exclusions"] = (["Interest", "Audience size", "Why", "Aimed at by an ad set now"], [[e["name"], e.get("sizeText") or "", e["why"], "Yes" if e.get("inUse") else "No"] for e in meta["exclusions"]])
    wr = meta.get("plan") or {}
    prows = [["Quick win", q["what"], q["why"]] for q in wr.get("quickWins") or []] + [["Lookalike source", x["seed"], x["why"]] for x in wr.get("lookalikes") or []]
    prows += [["Risk", r["risk"], r["mitigation"]] for r in wr.get("risks") or []]
    for key, p in (meta.get("plans") or {}).items():
        if p:
            prows.append(["How to: " + p["title"], p.get("worth") or "", " ".join("%d. %s" % (i, t) for i, t in enumerate(p["steps"], 1))])
    if prows:
        out["Meta plan"] = (["Kind", "Item", "Detail"], prows)
    if wr.get("questions"):
        out["Meta questions"] = (["Area", "Question", "Why it matters", "When"], [[q["area"], q["question"], q["why"], q["urgency"]] for q in wr["questions"]])
    chk = meta.get("checks") or {}
    crows = [["Not available", s] for s in chk.get("missing") or []] + [["Meta ad account", w] for w in chk.get("accountNotes") or []] + [["Interest search", w] for w in chk.get("interestNotes") or []]
    crows += [["A number in the write-up that is not in the research: check it", "%s: %s (%s)" % (n["where"], n["number"], n["text"])] for n in chk.get("numbers") or []]
    crows += [["Meta had no interest for the phrase", p] for p in chk.get("emptyPhrases") or []]
    if chk.get("unjudged"):
        crows.append(["Interests the model did not judge", str(chk["unjudged"])])
    if chk.get("droppedAnswers"):
        crows.append(["Answers left out because they named an interest that was not in the list", str(chk["droppedAnswers"])])
    if chk.get("cut"):
        crows.append(["More interests were found than were judged", "%d found, the first %d judged (the ones in use first)." % (chk.get("found") or 0, len(rows))])
    if chk.get("advantageAdsets"):
        crows.append(["Ad sets on Advantage+ audience", "%d: the interests they name are suggestions to Meta, not limits." % chk["advantageAdsets"]])
    crows += [["Assumption", a] for a in chk.get("assumptions") or []]
    if crows:
        out["Meta checks"] = (["Check", "Detail"], crows)
    return out


def tables(rec):
    """{sheet name: (header, rows)} in the order they are shown. Empty tables are left out. Numbers stay numbers."""
    plan = rec.get("plan") or {}
    inputs = rec.get("inputs") or {}
    research = rec.get("research") or {}
    rows = plan.get("keywords") or []
    themes = plan.get("themes") or []
    cov = plan.get("coverage") or {}
    wr = plan.get("plan") or {}
    cur = plan.get("currency") or inputs.get("currency") or "USD"
    names = {t["key"]: t["name"] for t in themes}
    out = {}

    google_on, meta_on = pk_platforms(inputs, plan)
    meta = plan.get("meta") or {}
    ov = [["Client", inputs.get("name")], ["Website", inputs.get("website")], ["Industry", inputs.get("industry")], ["Goal of the advertising", ps.GOALS.get(inputs.get("goal"), inputs.get("goal"))],
          ["Places", inputs.get("locations")]]
    if google_on:
        ov.append(["Google Ads account", (rec.get("account_name") or rec.get("account_id")) if inputs.get("account") else ""])
    if meta_on:
        ov.append(["Meta ad account", (meta.get("account") or {}).get("name") or inputs.get("metaAccount")])
    ov += [["Built by", rec.get("created_by")], ["Built on", (plan.get("generatedAt") or "")[:10]]]
    if google_on:
        ov += [["Smallest monthly searches looked at", inputs.get("minSearches")],
               ["Keywords judged", cov.get("judged")], ["Judged a fit", cov.get("yes")], ["Fit and new to the account", cov.get("new")], ["Monthly searches of the keywords that fit", cov.get("searches")],
               ["Already covered by the account, of the searches for ideas that share a word with what the client sells", ("%d%%" % cov["poolPct"]) if cov.get("poolPct") is not None else ""],
               ["Already covered by the account, of the keywords judged a fit", ("%d%%" % cov["pct"]) if cov.get("pct") is not None else ""], ["Not yet covered, monthly searches", cov.get("newSearches")]]
    mcov = meta.get("coverage") or {}
    if meta_on and meta:
        ov += [["Meta interests judged", mcov.get("judged")], ["Meta interests that fit", mcov.get("yes")], ["Fit and aimed at by no ad set", mcov.get("new")], ["Interests in use by the ad sets", mcov.get("inUse")],
               ["Ad set spend on interests judged not a fit (%s)" % meta.get("currency"), mcov.get("misfitSpend")]]
        ov += [["Meta: what the numbers say", h] for h in meta.get("headlines") or []] + [["Meta: in short", t] for t in (meta.get("plan") or {}).get("summary") or []]
    for plat, why in (plan.get("errors") or {}).items():
        ov.append(["%s part not written" % plat, why])
    ov += [["In short", s] for s in wr.get("summary") or []]
    out["Overview"] = (["Item", "Detail"], [r for r in ov if r[1] not in (None, "")])

    if themes:
        out["Themes"] = (["Theme", "Intent", "Priority", "Campaign", "An existing campaign", "Keywords that fit", "New to the account", "Monthly searches (fit)", "Not yet covered",
                          "Covered % (of the keywords judged)", "Average low bid (%s)" % cur, "Average high bid (%s)" % cur, "How it moves through the year", "Why"],
                         [[t["name"], t["intent"], t["priority"], t["campaign"], "Yes" if t.get("existing") else "No, a new campaign", t["yes"], t["new"], t["searches"], t["newSearches"],
                           t["coverage"] if t.get("coverage") is not None else "", t.get("avgLow") if t.get("avgLow") is not None else "", t.get("avgHigh") if t.get("avgHigh") is not None else "",
                           t.get("pattern") or "", t.get("why") or ""] for t in themes])

    kw_head = ["Keyword", "Fit", "Priority", "Theme", "Intent", "Suggested match", "Monthly searches", "Competition", "Low bid (%s)" % cur, "High bid (%s)" % cur, "How it moves through the year",
               "In the account", "Found from", "Search term conversions", "Note", "In the Ads Editor file"]

    def kw_row(r):
        return [r["kw"], r["fit"], r["priority"], names.get(r["theme"], "Other"), r["intent"], r["match"], r["avg"] if r["avg"] is not None else "", r["comp"],
                r["lo"] if r["lo"] is not None else "", r["hi"] if r["hi"] is not None else "", r["pattern"], status_text(r) + ((". " + blocked_text(r)) if blocked_text(r) else ""),
                _join(r.get("from")), (r.get("conv") or {}).get("conv") if r.get("conv") else "", r["note"], editor_text(r, bool(cov.get("accountUnread")))]

    if rows:
        ordered = sorted(rows, key=lambda r: (FIT_ORDER.get(r["fit"], 4), 0 if r["status"] == "new" else 1, -(r["avg"] or 0), r["kw"]))
        out["Keywords"] = (kw_head, [kw_row(r) for r in ordered])
        add = sorted((r for r in rows if r.get("include")), key=lambda r: ({"high": 0, "medium": 1, "low": 2}.get(r["priority"], 3), -(r["avg"] or 0), r["kw"]))
        by_key = {t["key"]: t for t in themes}
        if add:
            out["Add to account"] = (["Priority", "Keyword", "Match", "Theme", "Campaign", "Ad group", "Monthly searches", "Competition", "Low bid (%s)" % cur, "High bid (%s)" % cur, "How it moves", "Note"],
                                     [[r["priority"], r["kw"], r["match"], names.get(r["theme"], "Other"), kd.campaign_for(by_key.get(r["theme"]) or {}), kd.ad_group_for(by_key.get(r["theme"]) or {"name": "Other"}, r["intent"]),
                                       r["avg"] if r["avg"] is not None else "", r["comp"], r["lo"] if r["lo"] is not None else "", r["hi"] if r["hi"] is not None else "", r["pattern"], r["note"]] for r in add])
        have = sorted((r for r in rows if r["fit"] == "yes" and r["status"] != "new"), key=lambda r: -(r["avg"] or 0))
        if have:
            out["Already in the account"] = (["Keyword", "Monthly searches", "Status", "Covered by", "Match type", "Campaign", "Theme"],
                                             [[r["kw"], r["avg"] if r["avg"] is not None else "", "Already a keyword" if r["status"] == "bidding" else "Matched by a broader keyword", (r.get("by") or {}).get("kw", ""),
                                               (r.get("by") or {}).get("match", ""), (r.get("by") or {}).get("campaign", ""), names.get(r["theme"], "Other")] for r in have])
        conv = sorted((r for r in rows if r.get("conv")), key=lambda r: -(r["conv"].get("conv") or 0))
        if conv:
            out["Converting search terms"] = (["Search term", "Conversions", "Clicks", "Cost (%s)" % cur, "Campaign it came through", "Monthly searches", "Judged", "Theme", "Note"],
                                             [[r["kw"], r["conv"].get("conv"), r["conv"].get("clicks"), r["conv"].get("cost"), r["conv"].get("campaign") or "", r["avg"] if r["avg"] is not None else "",
                                               r["fit"], names.get(r["theme"], "Other"), r["note"]] for r in conv])
        rivals = sorted((r for r in rows if r["intent"] == "competitor"), key=lambda r: -(r["avg"] or 0))
        if rivals:
            out["Competitor keywords"] = (["Keyword", "Competitor named", "Monthly searches", "Competition", "Note"],
                                          [[r["kw"], r.get("rival") or "", r["avg"] if r["avg"] is not None else "", r["comp"], "Whether to bid on a rival's name is the client's decision. Not in the Ads Editor file."] for r in rivals])

    if wr.get("negatives"):
        out["Negatives to add"] = (["Negative keyword", "Match", "Why", "Note"], [[n["kw"], n["match"], n["why"], n.get("note") or ""] for n in wr["negatives"]])
    head, season = (plan.get("seasonal") or {}).get("head"), (plan.get("seasonal") or {}).get("rows")
    if head and season:
        out["Seasonality"] = (["Theme"] + head + ["How it moves"], season)
    prows = [["Quick win", q["what"], q["why"]] for q in wr.get("quickWins") or []] + [["Test", t["hypothesis"], "%s. Metric: %s" % (t["how"], t["metric"])] for t in wr.get("tests") or []]
    prows += [["Season", s["window"], s["action"]] for s in wr.get("seasonal") or []] + [["Risk", r["risk"], r["mitigation"]] for r in wr.get("risks") or []]
    if prows:
        out["Plan"] = (["Kind", "Item", "Detail"], prows)
    if wr.get("questions"):
        out["Questions for the client"] = (["Area", "Question", "Why it matters", "When"], [[q["area"], q["question"], q["why"], q["urgency"]] for q in wr["questions"]])
    editor = plan.get("editor") or []
    if editor:
        out["Ads Editor import"] = (EDITOR_HEADER, editor)

    if meta:
        out.update(meta_tables(meta))
    chk = plan.get("checks") or {}
    stats = plan.get("stats") or {}
    crows = [["Where searches were counted", chk["geography"]]] if chk.get("geography") else []
    crows += [["Google Ads Editor file not made", chk["editorWithheld"]]] if chk.get("editorWithheld") else []
    crows += [["Not available", s] for s in chk.get("missing") or []]
    crows += [["Account", w] for w in chk.get("accountWarnings") or []] + [["Not available from Google Ads", c] for c in chk.get("caveats") or []]
    crows += [["A number in the write-up that is not in the research: check it", "%s: %s (%s)" % (n["where"], n["number"], n["text"])] for n in chk.get("numbers") or []]
    crows += [["Keyword Planner had no ideas for", s] for s in chk.get("emptySeeds") or []]
    if chk.get("unjudged"):
        crows.append(["Keywords the model did not judge", str(chk["unjudged"])])
    for n in chk.get("removedNegatives") or []:
        crows.append(["Negative not suggested", "%s (%s): %s" % (n["kw"], n["match"], n["reason"])])
    for k in chk.get("convertingNotFit") or []:
        crows.append(["A converting search term judged not a fit: check it with the client", k])
    if cov.get("blockedInTarget"):
        crows.append(["Keywords left out of the Editor file because the campaign they would go in blocks them", str(cov["blockedInTarget"])])
    if chk.get("blockedInSomeCampaigns"):
        crows.append(["Keywords blocked by a negative in only some campaigns", str(chk["blockedInSomeCampaigns"])])
    for label, key in (("Ideas below the smallest search volume", "belowFloor"), ("Ideas on the avoid list", "avoided"), ("Ideas blocked in every campaign by the account's negatives", "blockedEverywhere"),
                       ("Ideas merged with a close variant", "merged")):
        if stats.get(key):
            crows.append([label, str(stats[key])])
    for a in chk.get("assumptions") or []:
        crows.append(["Assumption", a])
    if crows:
        out["Checks"] = (["Check", "Detail"], crows)
    return out


def xlsx(rec):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head_fill, head_font, wrap = PatternFill("solid", fgColor="1F2A44"), Font(bold=True, color="FFFFFF"), Alignment(wrap_text=True, vertical="top")
    used = set()
    for name, (header, rows) in tables(rec).items():
        name = SHEET_NAME_BAD.sub("", name)[:31] or "Sheet"
        base, n = name, 2
        while name in used:
            name, n = base[:28] + "-%d" % n, n + 1
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append([plain(h) for h in header])
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, wrap
        for r in rows:
            ws.append([plain(v) for v in r])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"
                c.alignment = wrap
        for i in range(1, len(header) + 1):
            longest = max([len(str(header[i - 1]))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None] or [10])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, longest + 2))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def csv_editor(rec):
    """The keywords to add, for Google Ads Editor's import, or None when there are none."""
    editor = (rec.get("plan") or {}).get("editor") or []
    return _csv_text(EDITOR_HEADER, editor) if editor else None


def csv_zip(rec):
    out = io.BytesIO()
    used = set()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for sheet, (header, rows) in tables(rec).items():
            stem = re.sub(r"[^A-Za-z0-9]+", "-", sheet).strip("-").lower() or "table"
            cand, n = stem, 2
            while cand in used:
                cand, n = "%s-%d" % (stem, n), n + 1
            used.add(cand)
            z.writestr(cand + ".csv", _csv_text(header, rows))
    return out.getvalue()


def pdf_tables(rec):
    """[(name, header, rows shown, note when rows were left out or None)] for the PDF: every table but the full keyword list and the Editor file (they are in the Excel and CSV
    downloads), the long ones cut to PDF_ROW_CAP rows and saying so."""
    out = []
    for name, (header, rows) in tables(rec).items():
        if name in PDF_SKIP:
            continue
        cap = PDF_ROW_CAP.get(name)
        shown = rows[:cap] if cap else rows
        out.append((name, header, shown, "%d of %d rows shown. All of them are in the Excel file." % (cap, len(rows)) if cap and len(rows) > cap else None))
    return out


def pdf(rec):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from .event_intel_pdf import BODY, BOLD, _fonts, text_markup
    _fonts()
    inputs = rec.get("inputs") or {}
    ink, mute, rule = colors.HexColor("#1B2333"), colors.HexColor("#5B6577"), colors.HexColor("#D7DDE8")
    base = ParagraphStyle("b", fontName=BODY, fontSize=8.4, leading=11.4, textColor=ink)
    small = ParagraphStyle("s", parent=base, fontSize=7.6, leading=10, textColor=mute)
    h1 = ParagraphStyle("h1", parent=base, fontName=BOLD, fontSize=17, leading=21, spaceBefore=4, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=base, fontName=BOLD, fontSize=12, leading=15, spaceBefore=10, spaceAfter=4)
    head = ParagraphStyle("th", parent=base, fontName=BOLD, textColor=colors.white, fontSize=7.8)
    page = landscape(A4)
    width = page[0] - 28 * mm

    def P(t, st=base):
        return Paragraph(text_markup(str(t if t is not None else "")), st)

    def table(header, rows):
        n = len(header)
        weights = [max(1.0, min(4.0, max([len(str(header[i]))] + [min(len(str(r[i])), 200) / 12.0 for r in rows if i < len(r)] or [1]) / 6.0)) for i in range(n)]
        cw = [width * w / sum(weights) for w in weights]
        data = [[Paragraph(text_markup(h), head) for h in header]] + [[P(str(c)[:CELL_CHARS]) for c in r] for r in rows]
        t = Table(data, colWidths=cw, repeatRows=1, splitInRow=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A44")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FA")]), ("LINEBELOW", (0, 0), (-1, -1), 0.25, rule),
                               ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        return t

    story = [P("Keyword Research", h1), P(inputs.get("name"), h2), P("%s | %s | goal of the advertising: %s" % (inputs.get("website") or "no website", inputs.get("industry") or "",
                                                                                                           ps.GOALS.get(inputs.get("goal"), "")), small), Spacer(1, 6)]
    for name, header, shown, cut in pdf_tables(rec):
        story += [PageBreak() if name in PDF_PAGE_BREAK else Spacer(1, 4), P(name, h2), table(header, shown)]
        if cut:
            story.append(P(cut, small))
    story.append(P("The full keyword list, with every figure, and the file for Google Ads Editor are in the Excel and CSV downloads.", small))
    out = io.BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(BODY, 7)
        canvas.setFillColor(mute)
        canvas.drawString(14 * mm, 8 * mm, "Position2: %s. Keyword Research" % clean(inputs.get("name"), 60))
        canvas.drawRightString(page[0] - 14 * mm, 8 * mm, "Page %d" % doc.page)
        canvas.restoreState()

    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                            title="Keyword Research: %s" % inputs.get("name"), author="Position2")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
