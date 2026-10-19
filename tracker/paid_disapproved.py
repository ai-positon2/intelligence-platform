"""Paid Marketing Agents: Disapproved Ads Agent, reading and the logic.

Read only. For one ad account it returns what Google or Meta has disapproved, and only what is live: an ad counts when its campaign, its ad group and the ad
itself are all enabled (active, on Meta). The rest of this file is what is done with it, with no database and no clock of its own (so it is tested without a
network): the words for a policy topic, what a finding is and when it is new, fixed or dismissed, the sheet and the alert.

A finding: {fingerprint, platform, accountId, accountName, where, objectId, objectName, adType, campaign, adGroup, url, severity, reason, topics}
(`where` says what was disapproved: an ad, a Performance Max asset group, or an asset in one).
"""
from __future__ import annotations

import hashlib
import html
import io

from . import paid_domains as dom

LIVE = "campaign.status = 'ENABLED'"
ADS_GAQL = (
    "SELECT campaign.id, campaign.name, ad_group.id, ad_group.name, ad_group_ad.ad.id, ad_group_ad.ad.name, ad_group_ad.ad.type, ad_group_ad.ad.final_urls, "
    "ad_group_ad.policy_summary.approval_status, ad_group_ad.policy_summary.review_status, ad_group_ad.policy_summary.policy_topic_entries "
    "FROM ad_group_ad WHERE ad_group_ad.status = 'ENABLED' AND ad_group.status = 'ENABLED' AND campaign.status = 'ENABLED' "
    "AND ad_group_ad.policy_summary.approval_status = 'DISAPPROVED'")
ASSET_GROUPS_GAQL = (
    "SELECT campaign.id, campaign.name, asset_group.id, asset_group.name, asset_group.final_urls, asset_group.primary_status, asset_group.primary_status_reasons "
    "FROM asset_group WHERE asset_group.status = 'ENABLED' AND campaign.status = 'ENABLED' AND asset_group.primary_status = 'NOT_ELIGIBLE'")
ASSETS_GAQL = (
    "SELECT campaign.id, campaign.name, asset_group.id, asset_group.name, asset_group_asset.field_type, asset_group_asset.policy_summary.approval_status, "
    "asset_group_asset.policy_summary.policy_topic_entries, asset.id, asset.type, asset.name, asset.text_asset.text FROM asset_group_asset "
    "WHERE asset_group_asset.status = 'ENABLED' AND asset_group.status = 'ENABLED' AND campaign.status = 'ENABLED' "
    "AND asset_group_asset.policy_summary.approval_status = 'DISAPPROVED'")
META_FIELDS = ("id,name,status,effective_status,ad_review_feedback,issues_info,campaign{name,status,effective_status},adset{name,status,effective_status}")
META_FIELDS_PLAIN = "id,name,status,effective_status,campaign{name,status},adset{name,status}"
META_PAGE = 100

AD_TYPE_NAMES = {"RESPONSIVE_SEARCH_AD": "Responsive Search Ad", "EXPANDED_TEXT_AD": "Expanded Text Ad", "DEMAND_GEN_MULTI_ASSET_AD": "Demand Gen ad",
                 "DEMAND_GEN_VIDEO_RESPONSIVE_AD": "Demand Gen video ad", "DEMAND_GEN_CAROUSEL_AD": "Demand Gen carousel ad", "RESPONSIVE_DISPLAY_AD": "Responsive Display Ad",
                 "SHOPPING_PRODUCT_AD": "Shopping ad", "CALL_AD": "Call ad", "APP_AD": "App ad", "IMAGE_AD": "Image ad", "VIDEO_AD": "Video ad",
                 "VIDEO_RESPONSIVE_AD": "Video responsive ad", "RESPONSIVE_SEARCH": "Responsive Search Ad", "LOCAL_AD": "Local ad", "SMART_CAMPAIGN_AD": "Smart campaign ad"}
TOPIC_TYPES = {"PROHIBITED": "prohibited", "LIMITED": "limited", "FULLY_LIMITED": "fully limited", "DESCRIPTIVE": "descriptive", "BROADENING": "broadening",
               "AREA_OF_INTEREST_ONLY": "area of interest only"}
SEVERITY_ORDER = {"high": 0, "medium": 1}
SHOWN = 15
SHEET_NAME = "disapproved-ads.xlsx"
SHEET_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _g(row, *path):
    for p in path:
        if not isinstance(row, dict):
            return None
        row = row.get(p)
    return row


def words(code):
    """A Google policy code in words: DESTINATION_NOT_WORKING -> Destination not working."""
    s = str(code or "").replace("_", " ").strip().lower()
    return s[:1].upper() + s[1:]


def evidence_of(entry):
    """The short pieces of evidence in a policy topic entry: the offending text, the website, the destination that did not work (capped)."""
    out = []
    for ev in entry.get("evidences") or []:
        for key, sub in (("textList", "texts"), ("websiteList", "websites"), ("destinationTextList", "destinationTexts")):
            for t in _g(ev, key, sub) or []:
                if t and t not in out:
                    out.append(str(t)[:120])
        nw = ev.get("destinationNotWorking") or {}
        if nw:
            out.append("%s%s" % (words(nw.get("device")) + ": " if nw.get("device") else "", words(nw.get("dnsErrorType") or nw.get("httpErrorCode") or "destination not working")))
    return out[:4]


def topics_of(entries):
    """[{topic, type, evidence}] from Google's policy_topic_entries."""
    out = []
    for e in entries or []:
        out.append({"topic": words(e.get("topic") or "Policy issue"), "type": TOPIC_TYPES.get(e.get("type"), words(e.get("type"))), "evidence": evidence_of(e)})
    return out


def reason_text(topics):
    """One sentence for what is wrong: 'Destination not working (prohibited); Misrepresentation (limited): "free gift"'."""
    parts = []
    for t in topics:
        s = "%s (%s)" % (t["topic"], t["type"]) if t.get("type") else t["topic"]
        if t.get("evidence"):
            s += ": " + ", ".join('"%s"' % e if len(e) < 60 else e for e in t["evidence"][:2])
        parts.append(s)
    return "; ".join(parts) or "Disapproved (no reason given)"


def fingerprint(platform, account_id, where, object_id):
    return hashlib.sha1(("da|%s|%s|%s|%s" % (platform, account_id, where, object_id)).encode()).hexdigest()[:24]


def _item(account, where, object_id, object_name, ad_type, campaign, ad_group, url, severity, topics, reason=None):
    return {"fingerprint": fingerprint(account["platform"], account["id"], where, object_id), "platform": account["platform"], "accountId": account["id"],
            "accountName": account["name"], "where": where, "objectId": str(object_id), "objectName": object_name, "adType": ad_type, "campaign": campaign,
            "adGroup": ad_group, "url": url, "severity": severity, "reason": reason or reason_text(topics), "topics": topics}


# ── Reading ─────────────────────────────────────────────────────────────────

def google_items(account, search):
    """(items, {ads, assetGroups, assets}): what is disapproved among the live ads, Performance Max asset groups and their assets of a Google Ads account.
    A failed read is raised: the scan names the account as unreadable instead of counting it clean."""
    out = []
    ads = search(account["id"], ADS_GAQL)
    for r in ads:
        ad = _g(r, "adGroupAd", "ad") or {}
        topics = topics_of(_g(r, "adGroupAd", "policySummary", "policyTopicEntries"))
        kind = AD_TYPE_NAMES.get(ad.get("type"), words(ad.get("type") or "Ad"))
        urls = ad.get("finalUrls") or []
        out.append(_item(account, "ad", ad.get("id"), "%s %s" % (kind, ad.get("id")), kind, _g(r, "campaign", "name") or "", _g(r, "adGroup", "name") or "",
                         urls[0] if urls else "", "high", topics))
    groups = search(account["id"], ASSET_GROUPS_GAQL)
    gcount = 0
    for r in groups:
        g = r.get("assetGroup") or {}
        reasons = [x for x in (g.get("primaryStatusReasons") or []) if "DISAPPROV" in str(x)]
        if not reasons:
            continue
        gcount += 1
        out.append(_item(account, "asset group", g.get("id"), g.get("name") or "Asset group %s" % g.get("id"), "Performance Max asset group", _g(r, "campaign", "name") or "", "",
                         (g.get("finalUrls") or [""])[0], "high", [], "Google marks the asset group not eligible: " + ", ".join(words(x) for x in reasons)))
    assets = search(account["id"], ASSETS_GAQL)
    for r in assets:
        a, aga, g = r.get("asset") or {}, r.get("assetGroupAsset") or {}, r.get("assetGroup") or {}
        topics = topics_of(_g(aga, "policySummary", "policyTopicEntries"))
        field = words(aga.get("fieldType") or a.get("type") or "asset")
        label = _g(a, "textAsset", "text") or a.get("name") or "%s %s" % (field, a.get("id"))
        out.append(_item(account, "asset", "%s:%s:%s" % (g.get("id"), a.get("id"), aga.get("fieldType") or ""), label, field, _g(r, "campaign", "name") or "", g.get("name") or "",
                         "", "medium", topics))
    return out, {"ads": len(ads), "assetGroups": gcount, "assets": len(assets)}


def meta_items(account, pages):
    """(items, {ads}): the live Meta ads Meta has disapproved. `pages(path, params)` is the Graph call that gives every row of a list. An ad counts only when it is
    active in an active ad set in an active campaign."""
    import json
    params = {"effective_status": json.dumps(["DISAPPROVED"]), "limit": META_PAGE}
    try:
        rows = pages(account["id"] + "/ads", dict(params, fields=META_FIELDS))
    except Exception:                                                              # noqa: BLE001 (a field this token may not read: ask for the plain ones)
        rows = pages(account["id"] + "/ads", dict(params, fields=META_FIELDS_PLAIN))
    out = []
    for r in rows:
        if r.get("status") != "ACTIVE" or _g(r, "campaign", "status") != "ACTIVE" or _g(r, "adset", "status") != "ACTIVE":
            continue
        reasons = []
        for i in r.get("issues_info") or []:
            t = i.get("error_summary") or i.get("error_message")
            if t and t not in reasons:
                reasons.append(str(t)[:200])
        for scope in (r.get("ad_review_feedback") or {}).values():
            for code, text in (scope or {}).items():
                t = text if isinstance(text, str) and text else words(code)
                if t not in reasons:
                    reasons.append(str(t)[:200])
        topics = [{"topic": t, "type": "", "evidence": []} for t in reasons[:4]]
        out.append(_item(account, "ad", r.get("id"), r.get("name") or "Ad %s" % r.get("id"), "Meta ad", _g(r, "campaign", "name") or "", _g(r, "adset", "name") or "", "",
                         "high", topics, "; ".join(reasons[:3]) or "Disapproved by Meta (no reason given)"))
    return out, {"ads": len(rows)}


# ── Findings ────────────────────────────────────────────────────────────────

def apply(items, existing, ok_ids, stamp):
    """The findings after a scan. `items` is what is disapproved and live now, `existing` {fingerprint: finding}. Returns (rows to save, {new, reopened, fixed}).
    A finding not seen again is fixed only if its account was read (`ok_ids`): an account that could not be read is not a clean bill."""
    writes, seen, new, reopened, fixed = [], set(), 0, 0, 0
    for it in items:
        fp = it["fingerprint"]
        seen.add(fp)
        old = existing.get(fp)
        base = dict(it, lastSeen=stamp)
        if old is None:
            writes.append(dict(base, status="open", firstSeen=stamp, alertedAt=None, alert={}, decidedBy=None, decidedAt=None))
            new += 1
        elif old["status"] == "dismissed":
            writes.append(dict(old, reason=it["reason"], topics=it["topics"], lastSeen=stamp))
        elif old["status"] == "fixed":
            writes.append(dict(base, status="open", firstSeen=stamp, alertedAt=None, alert={}, decidedBy=None, decidedAt=None))
            reopened += 1
        else:
            writes.append(dict(base, status="open", firstSeen=old["firstSeen"], alertedAt=old.get("alertedAt"), alert=old.get("alert") or {},
                               decidedBy=old.get("decidedBy"), decidedAt=old.get("decidedAt")))
    for fp, old in existing.items():
        if fp not in seen and old["status"] == "open" and old["accountId"] in ok_ids:
            writes.append(dict(old, status="fixed"))
            fixed += 1
    return writes, {"new": new, "reopened": reopened, "fixed": fixed}


def order(rows):
    return sorted(rows, key=lambda f: (SEVERITY_ORDER.get(f["severity"], 9), f["accountName"].lower(), f["campaign"].lower(), f["adGroup"].lower(), f["objectName"]))


# ── The sheet and the alert ─────────────────────────────────────────────────

SHEET_HEAD = ["Account", "Platform", "Campaign", "Ad group", "What", "Name or id", "Type", "Why it was disapproved", "Final URL", "Disapproved since", "New this scan",
              "Fixed or appealed? (fill in)"]


def sheet(rows, new_fingerprints=()):
    """Every open disapproval as an Excel file (bytes)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Disapproved"
    ws.append(SHEET_HEAD)
    marked = set(new_fingerprints)
    for f in order(rows):
        ws.append([f["accountName"], dom.PLATFORM.get(f["platform"], f["platform"]), f["campaign"], f["adGroup"], f["where"], f["objectName"], f["adType"], f["reason"], f["url"],
                   str(f.get("firstSeen") or "")[:10], "Yes" if f["fingerprint"] in marked else "", ""])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for col, w in zip("ABCDEFGHIJKL", (28, 11, 30, 26, 14, 32, 24, 70, 50, 15, 13, 26)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _what(f):
    return " / ".join(x for x in (f["campaign"], f["adGroup"], f["objectName"]) if x) or f["where"]


def _line(f):
    return "%s (%s) · %s · %s" % (f["accountName"], dom.PLATFORM.get(f["platform"], f["platform"]), _what(f), f["reason"])


def digest(open_rows, new_rows, link):
    """{subject, text, html, slack_text, slack_blocks, attachments} for a scan that found new disapproved ads. The sheet holds every open one, the words list the
    new ones."""
    accounts = len({f["accountId"] for f in new_rows})
    n, total = len(new_rows), len(open_rows)
    subject = "%d ad%s newly disapproved in %d ad account%s (%d disapproved in all)" % (n, "" if n == 1 else "s", accounts, "" if accounts == 1 else "s", total)
    shown = order(new_rows)[:SHOWN]
    more = n - len(shown)
    head = "%d new disapproval%s on live ads, %d in all. Only enabled campaigns and ad groups are checked. The attached sheet lists every one with its campaign, ad and reason." % (
        n, "" if n == 1 else "s", total)
    text = "\n".join([head, ""] + ["- " + _line(f) for f in shown] + (["...and %d more in the sheet." % more] if more else []) + ["", "Review them: " + link])
    e = html.escape
    rows = "".join("<li><b>%s</b> (%s)<br>%s<br>%s</li>" % (e(f["accountName"]), e(dom.PLATFORM.get(f["platform"], f["platform"])), e(_what(f)), e(f["reason"])) for f in shown)
    html_body = ("<div style=\"font-family:Arial,sans-serif;font-size:14px;line-height:1.5\"><p>%d new disapproval%s on live ads, %d in all. Only enabled campaigns and ad groups are "
                 "checked. The attached sheet (<b>%s</b>) lists every one with its campaign, ad and reason.</p><ul>%s</ul>%s<p><a href=\"%s\">Open the Disapproved Ads Agent</a></p></div>"
                 % (n, "" if n == 1 else "s", total, SHEET_NAME, rows, ("<p>...and %d more in the sheet.</p>" % more) if more else "", e(link)))
    blocks = [{"type": "header", "text": {"type": "plain_text", "text": subject[:150]}}]
    for f in shown[:8]:
        blocks.append({"type": "section", "text": {"type": "mrkdwn", "text": "*%s* (%s)\n%s\n%s" % (f["accountName"], dom.PLATFORM.get(f["platform"], f["platform"]), _what(f), f["reason"])}})
    blocks.append({"type": "section", "text": {"type": "mrkdwn", "text": ("...and %d more. " % (n - 8) if n > 8 else "") + "The sheet was emailed. <%s|Open the Disapproved Ads Agent>" % link}})
    return {"subject": subject, "text": text, "html": html_body, "slack_text": "*%s*" % subject, "slack_blocks": blocks,
            "attachments": [(SHEET_NAME, sheet(open_rows, [f["fingerprint"] for f in new_rows]), SHEET_TYPE)]}
