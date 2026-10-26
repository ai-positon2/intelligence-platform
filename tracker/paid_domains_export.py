"""Paid Marketing Agents: Unknown Domains Alert, the download.

One Excel workbook with everything the agent holds, so someone can take a dump and check it outside the page:

  Findings            every finding of every status (open, approved, dismissed, gone), with why, where and when
  Domain database     every approved domain and who it is for
  In use today        the domains each account's ads pointed at at the last scan, and whether the database covers each for that account
  Links to fix        every broken or redirecting link the URL Checker holds, of every status
  Scans               the recent scans: what each read, found and alerted

Nothing is filtered by what the page is showing; it is the whole record.
"""
from __future__ import annotations

import io

from . import paid_domains as dom
from . import paid_links as pl

FILENAME = "unknown-domains-export.xlsx"
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
KIND = {"unknown": "Unknown domain", "wrong_account": "Wrong account"}
STATUS = {"open": "Open", "approved": "Approved", "dismissed": "Dismissed", "gone": "Gone", "fixed": "Fixed"}


def _day(v):
    return str(v or "")[:19].replace("T", " ")


def _sheet(wb, title, head, rows, widths):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet(title)
    ws.append(head)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F2937")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def workbook(store, accounts, groups):
    """The whole record as an Excel file (bytes). `accounts` and `groups` give the names of who an entry is for."""
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    names = {x["id"]: x["name"] for x in list(accounts) + list(groups)}
    platform = dom.PLATFORM

    findings = sorted(store.findings_all(), key=lambda f: (f["status"] != "open", dom.SEVERITY_ORDER.get(f["severity"], 9), f["accountName"].lower(), f["domain"]))
    _sheet(wb, "Findings", ["Status", "Severity", "Kind", "Account", "Platform", "Account id", "Where", "Ad or asset", "Campaign", "Ad group", "Address", "Domain", "Why",
                            "Looks like", "Redirect", "Approved for", "First seen", "Last seen", "Alerted", "Decided by", "Decided at"],
           [[STATUS.get(f["status"], f["status"]), f["severity"], KIND.get(f["kind"], f["kind"]), f["accountName"], platform.get(f["platform"], f["platform"]), f["accountId"],
             f["where"], f["objectName"], f["campaign"], f["adGroup"], f["url"], f["domain"], f["why"], f.get("lookalike") or "", "Yes" if f.get("redirect") else "",
             ", ".join(f.get("owners") or []), _day(f.get("firstSeen")), _day(f.get("lastSeen")), _day(f.get("alertedAt")), f.get("decidedBy") or "", _day(f.get("decidedAt"))]
            for f in findings], (11, 10, 16, 28, 11, 14, 22, 28, 28, 24, 52, 24, 60, 22, 9, 24, 19, 19, 19, 26, 19))

    known = store.known_list()
    _sheet(wb, "Domain database", ["Domain", "For", "For (id)", "Note", "Added by", "Added at"],
           [[k["domain"], "Everyone" if k["scope"] == dom.EVERYONE else names.get(k["scope"], k["scope"]), "" if k["scope"] == dom.EVERYONE else k["scope"], k.get("note") or "",
             k.get("added_by") or "", _day(k.get("added_at"))] for k in known], (30, 32, 16, 40, 28, 19))

    snap = (store.state_get("inuse") or {}).get("accounts") or {}
    use_rows = []
    for aid, a in sorted(snap.items(), key=lambda kv: kv[1]["name"].lower()):
        acct = {"id": aid, "platform": a["platform"], "group": a.get("group") or ""}
        for d, n in sorted(a["domains"].items()):
            use_rows.append([a["name"], platform.get(a["platform"], a["platform"]), aid, d, n, "Yes" if dom.judge(d, acct, known)[0] == "ok" else "No"])
    _sheet(wb, "In use today", ["Account", "Platform", "Account id", "Domain its ads point at", "Addresses", "In the database for this account"], use_rows, (30, 11, 16, 34, 11, 18))

    links = sorted(store.links_all(), key=lambda f: (f["status"] != "open", f["accountName"].lower(), f["campaign"].lower(), f["url"]))
    _sheet(wb, "Links to fix", ["Status", "Type", "Account", "Platform", "Campaign", "Ad group", "Where", "Ad or asset", "Link in the ad", "Problem", "Redirects to", "Status code",
                                "First seen", "Last seen", "Alerted"],
           [[STATUS.get(f["status"], f["status"]), pl.KIND_TEXT.get(f["kind"], f["kind"]), f["accountName"], platform.get(f["platform"], f["platform"]), f["campaign"], f["adGroup"],
             f["where"], f["objectName"], f["url"], f["problem"], f.get("target") or "", f.get("code") or "", _day(f.get("firstSeen")), _day(f.get("lastSeen")), _day(f.get("alertedAt"))]
            for f in links], (11, 24, 28, 11, 28, 24, 20, 28, 56, 44, 56, 11, 19, 19, 19))

    runs = store.state_get("runs") or []
    _sheet(wb, "Scans", ["Finished", "Started by", "Status", "Accounts", "Read", "Unreadable", "Addresses", "New findings", "Gone", "Alerted", "Email", "Slack", "URLs checked",
                         "Broken links", "Redirects", "Links alerted", "Possible takeover"],
           [[_day(r.get("finishedAt")), "the schedule" if r.get("trigger") == "schedule" else str(r.get("trigger") or "").replace("page by ", ""), r.get("status"), r.get("accounts"),
             r.get("read"), r.get("unreadableCount"), r.get("sightings"), r.get("new"), r.get("gone"), r.get("alerted"), (r.get("email") or {}).get("status", ""),
             (r.get("slack") or {}).get("status", ""), (r.get("links") or {}).get("urls"), (r.get("links") or {}).get("broken"), (r.get("links") or {}).get("redirects"),
             (r.get("links") or {}).get("alerted"), ", ".join(t["accountName"] for t in r.get("takeover") or [])] for r in runs],
           (19, 28, 10, 10, 8, 11, 11, 13, 8, 9, 14, 14, 13, 13, 11, 13, 36))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
