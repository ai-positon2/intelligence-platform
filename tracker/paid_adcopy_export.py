"""Paid Marketing Agents: Ad Copy & Creative Generator, the downloads.

One result, four files, all drawn from `tables(rec)` so they cannot drift apart:

  Excel    every table as a sheet: one sheet for each ad format with every line, its angle, its length against the limit and whether it is ready or needs the client to confirm a
           claim; the creative briefs; what to confirm; the checks; the facts the copy was allowed to use
  CSV      every table as its own CSV in a zip, and one CSV for Google Ads Editor (responsive search ads only): the lines that are ready, in the order they were written, the ad
           Paused so nothing serves until a person has looked at it. A line that needs the client to confirm a claim is left out of it
  PDF      the result as a document to read and to send

A value that a spreadsheet would run as a formula (it starts with = + - @) is written as plain text in every format (tracker/paid_strategy_export.py).
"""
from __future__ import annotations

import io
import re
import zipfile

from . import paid_adcopy as pa, paid_adcopy_check as ck, paid_strategy as ps
from .paid_strategy_check import text as clean
from .paid_strategy_export import CELL_CHARS, SHEET_NAME_BAD, _csv_text, _join, plain

EDITOR_HEADERS = 15
EDITOR_DESCRIPTIONS = 4
STATUS_TEXT = {ck.READY: "Ready", ck.CONFIRM: "Confirm with the client"}


def file_stem(rec):
    base = re.sub(r"[^A-Za-z0-9]+", "-", rec.get("name") or "client").strip("-")[:50] or "client"
    return "%s-ad-copy" % base


def _formats(rec):
    plan = rec.get("plan") or {}
    formats = plan.get("formats") or {}
    return [formats[k] for k in plan.get("order") or [] if k in formats]


def _format_rows(res):
    rows = []
    for f in res["fields"]:
        for i in f["items"]:
            rows.append([f["label"], i["angle"], i["text"], i["chars"], f["limit"], STATUS_TEXT.get(i["status"], i["status"]),
                         _join(i["claims"]), _join(i["flags"])])
    paths = res.get("paths") or {}
    for k, label in (("path1", "Display path 1"), ("path2", "Display path 2")):
        if paths.get(k):
            rows.append([label, "", paths[k], len(paths[k]), pa.FORMATS[res["key"]]["paths"], "Ready", "", ""])
    if res.get("cta"):
        rows.append(["Button", "", res["cta"], len(res["cta"]), "", "Ready", "", res.get("ctaNote") or ""])
    return rows


def tables(rec):
    """{sheet name: (header, rows)} in the order they are shown. Empty tables are left out. Numbers stay numbers."""
    plan = rec.get("plan") or {}
    inputs = rec.get("inputs") or {}
    chk = plan.get("checks") or {}
    results = _formats(rec)
    totals = chk.get("totals") or {}
    out = {}

    ov = [["Client", inputs.get("name")], ["Website", inputs.get("website")], ["Industry", inputs.get("industry")], ["What the ads are for", inputs.get("product")],
          ["Goal of the advertising", ps.GOALS.get(inputs.get("goal"), inputs.get("goal"))], ["Who it is for", inputs.get("audience")], ["Tone", pa.TONES.get(inputs.get("tone"), inputs.get("tone"))],
          ["Angles", _join([pa.ANGLES.get(a, a) for a in inputs.get("angles") or []])], ["Language", inputs.get("language")], ["Where", inputs.get("locations")],
          ["Formats", _join([r["label"] for r in results])], ["Google Ads account read", rec.get("account_name") or rec.get("account_id")], ["Built by", rec.get("created_by")],
          ["Built on", (plan.get("generatedAt") or "")[:10]], ["Lines ready to use", totals.get("ready")], ["Lines to confirm with the client first", totals.get("confirm")]]
    out["Overview"] = (["Item", "Detail"], [r for r in ov if r[1] not in (None, "")])

    head = ["Field", "Angle", "Text", "Characters", "Limit", "Status", "What to confirm", "Notes"]
    for res in results:
        rows = _format_rows(res)
        if rows:
            out[res["short"]] = (head, rows)

    brief = []
    for res in results:
        b = res.get("brief")
        if not b:
            continue
        for label, key in (("Concept", "concept"), ("Hook", "hook"), ("Visual direction", "visual")):
            if b.get(key):
                brief.append([res["short"], label, b[key]])
        for n, s in enumerate(b.get("script") or [], 1):
            brief.append([res["short"], "Beat %d" % n, " | ".join(x for x in (s.get("beat"), ("On screen: " + s["onScreen"]) if s.get("onScreen") else "",
                                                                              ("Said: " + s["voiceover"]) if s.get("voiceover") else "") if x)])
        for label, key in (("Ask the client for", "assets"), ("Do", "dos"), ("Do not", "donts")):
            for x in b.get(key) or []:
                brief.append([res["short"], label, x])
        for s in res.get("specs") or []:
            brief.append([res["short"], "Size: " + s["asset"], "%s. %s" % (s["size"], s["note"])])
    if brief:
        out["Creative briefs"] = (["Format", "Item", "Detail"], brief)

    confirm = [[c["format"], c["field"], c["text"], _join(c["claims"]), ""] for c in chk.get("claims") or []]
    if confirm:
        out["To confirm"] = (["Format", "Field", "Text", "Claim the evidence does not hold", "Confirmed by the client? (fill in)"], confirm)

    facts = chk.get("factsUsed") or {}
    frows = [["Confirmed by the person", t] for t in facts.get("confirmed") or []] + [["Offer on the website", t] for t in facts.get("siteOffers") or []]
    if facts.get("siteRating"):
        r = facts["siteRating"]
        frows.append(["Rating on the website", "%s from %s reviews" % (r.get("value"), r.get("count"))])
    if frows:
        out["Facts used"] = (["Where from", "Fact"], frows)

    crows = [["Lines ready", totals.get("ready")], ["Lines to confirm", totals.get("confirm")]]
    for k, n in sorted((chk.get("dropped") or {}).items(), key=lambda kv: -kv[1]):
        crows.append(["Dropped: %s" % ck.REASONS.get(k, k), n])
    for s in chk.get("shortfalls") or []:
        crows.append(["Short", s])
    kw = chk.get("keywords")
    if kw:
        crows.append(["Keywords in a Google headline", _join(kw.get("covered"))])
        if kw.get("missing"):
            crows.append(["Keywords in no Google headline", _join(kw["missing"])])
    for m in chk.get("missing") or []:
        crows.append(["Not available", m])
    out["Checks"] = (["Check", "Detail"], [r for r in crows if r[1] not in (None, "")])
    return out


def xlsx(rec):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head_fill, head_font, wrap = PatternFill("solid", fgColor="1F2A44"), Font(bold=True, color="FFFFFF"), Alignment(wrap_text=True, vertical="top")
    used = set()
    for name, (header, rows) in tables(rec).items():
        name = SHEET_NAME_BAD.sub("", name)[:31] or "Sheet"
        base, n = name, 2
        while name in used:
            name, n = base[:28] + "-%d" % n, n + 1
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append([plain(h) for h in header])
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, wrap
        for r in rows:
            ws.append([plain(v) for v in r])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, str):
                    c.data_type = "s"
                c.alignment = wrap
        for i in range(1, len(header) + 1):
            longest = max([len(str(header[i - 1]))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r) and r[i - 1] is not None] or [10])
            ws.column_dimensions[get_column_letter(i)].width = min(70, max(10, longest + 2))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def csv_editor(rec):
    """The responsive search ad for Google Ads Editor's import, or None when there is none that is ready: the lines that are ready only, at least three headlines and two
    descriptions, the ad Paused."""
    inputs = rec.get("inputs") or {}
    for res in _formats(rec):
        if res["key"] != "google_rsa":
            continue
        heads = [i["text"] for i in res["fields"][0]["items"] if i["status"] == ck.READY][:EDITOR_HEADERS]
        descs = [i["text"] for i in res["fields"][1]["items"] if i["status"] == ck.READY][:EDITOR_DESCRIPTIONS]
        if len(heads) < 3 or len(descs) < 2:
            return None
        paths = res.get("paths") or {}
        header = (["Campaign", "Ad group", "Ad type", "Status", "Final URL", "Path 1", "Path 2"] + ["Headline %d" % n for n in range(1, EDITOR_HEADERS + 1)]
                  + ["Description %d" % n for n in range(1, EDITOR_DESCRIPTIONS + 1)])
        row = ([inputs.get("campaignName") or "New campaign", inputs.get("adGroupName") or "New ad group", "Responsive search ad", "Paused", inputs.get("website") or "",
                paths.get("path1") or "", paths.get("path2") or ""] + heads + [""] * (EDITOR_HEADERS - len(heads)) + descs + [""] * (EDITOR_DESCRIPTIONS - len(descs)))
        return _csv_text(header, [row])
    return None


def csv_zip(rec):
    out = io.BytesIO()
    used = set()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for sheet, (header, rows) in tables(rec).items():
            stem = re.sub(r"[^A-Za-z0-9]+", "-", sheet).strip("-").lower() or "table"
            cand, n = stem, 2
            while cand in used:
                cand, n = "%s-%d" % (stem, n), n + 1
            used.add(cand)
            z.writestr(cand + ".csv", _csv_text(header, rows))
    return out.getvalue()


def pdf(rec):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from .event_intel_pdf import BODY, BOLD, _fonts, text_markup
    _fonts()
    inputs = rec.get("inputs") or {}
    ink, mute, rule = colors.HexColor("#1B2333"), colors.HexColor("#5B6577"), colors.HexColor("#D7DDE8")
    base = ParagraphStyle("b", fontName=BODY, fontSize=8.4, leading=11.4, textColor=ink)
    small = ParagraphStyle("s", parent=base, fontSize=7.6, leading=10, textColor=mute)
    h1 = ParagraphStyle("h1", parent=base, fontName=BOLD, fontSize=17, leading=21, spaceBefore=4, spaceAfter=6)
    h2 = ParagraphStyle("h2", parent=base, fontName=BOLD, fontSize=12, leading=15, spaceBefore=10, spaceAfter=4)
    head = ParagraphStyle("th", parent=base, fontName=BOLD, textColor=colors.white, fontSize=7.8)
    page = landscape(A4)
    width = page[0] - 28 * mm

    def P(t, st=base):
        return Paragraph(text_markup(str(t if t is not None else "")), st)

    def table(header, rows):
        n = len(header)
        weights = [max(1.0, min(4.0, max([len(str(header[i]))] + [min(len(str(r[i])), 200) / 12.0 for r in rows if i < len(r)] or [1]) / 6.0)) for i in range(n)]
        cw = [width * w / sum(weights) for w in weights]
        data = [[Paragraph(text_markup(h), head) for h in header]] + [[P(str(c)[:CELL_CHARS]) for c in r] for r in rows]
        t = Table(data, colWidths=cw, repeatRows=1, splitInRow=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A44")), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                               ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6FA")]), ("LINEBELOW", (0, 0), (-1, -1), 0.25, rule),
                               ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
        return t

    story = [P("Ad Copy and Creative Briefs", h1), P(inputs.get("name"), h2),
             P("%s | goal of the advertising: %s" % (inputs.get("website") or "no website", ps.GOALS.get(inputs.get("goal"), "")), small), Spacer(1, 6)]
    first = True
    for name, (header, rows) in tables(rec).items():
        story += [Spacer(1, 4) if first else PageBreak(), P(name, h2), table(header, rows)]
        first = False
    out = io.BytesIO()

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont(BODY, 7)
        canvas.setFillColor(mute)
        canvas.drawString(14 * mm, 8 * mm, "Position2: %s. Ad copy" % clean(inputs.get("name"), 60))
        canvas.drawRightString(page[0] - 14 * mm, 8 * mm, "Page %d" % doc.page)
        canvas.restoreState()

    doc = SimpleDocTemplate(out, pagesize=page, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=14 * mm,
                            title="Ad copy: %s" % inputs.get("name"), author="Position2")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return out.getvalue()
