"""Paid Marketing Agents: Disapproved Ads Agent.

Pinned here: the reading asks only for live ads (an enabled campaign, an enabled ad group and an enabled ad; active on Meta) and only for what is disapproved;
Google's policy topics become words with their evidence; a Performance Max asset group is reported only when Google says it is not eligible because it is
disapproved; Meta ads that are paused, or in a paused ad set or campaign, are never reported and a field the token may not read falls back to the plain ones;
findings (new, fixed, back, dismissed; an unreadable account never marks its findings fixed); the sheet and the alert words; and the scan: one email with the
sheet and one Slack message, once per finding, nothing lost when a channel fails, the Unknown Domains Alert's recipients used until this agent has its own, an
account that cannot be read named, and two scans never at once. No network.
"""
import io
import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from domains_fakes import Deliver, MemStore  # noqa: E402
from tracker import paid_disapproved as da  # noqa: E402
from tracker import paid_disapproved_scan as ds  # noqa: E402

G = {"id": "1112223330", "name": "Acme Garden", "platform": "google", "group": "mcc-1"}
H = {"id": "4445556660", "name": "Bolt Dental", "platform": "google", "group": "mcc-2"}
M = {"id": "act_77", "name": "Acme Garden", "platform": "meta", "group": ""}
LINK = "https://p.example/p2/paid-marketing/disapproved-ads"


class DaStore(MemStore):
    """The domains double with this agent's findings kept apart from that agent's."""

    def __init__(self):
        super().__init__()
        self.da = {}

    def findings_all(self):
        return [dict(f) for f in sorted(self.da.values(), key=lambda f: f["lastSeen"], reverse=True)]

    def finding_get(self, fp):
        return dict(self.da[fp]) if fp in self.da else None

    def findings_put(self, rows):
        for f in rows:
            self.da[f["fingerprint"]] = dict(f)


# ── The reading asks for live ads only ──────────────────────────────────────

def test_the_queries_hold_every_level_to_enabled_and_to_disapproved():
    assert all(x in da.ADS_GAQL for x in ("ad_group_ad.status = 'ENABLED'", "ad_group.status = 'ENABLED'", "campaign.status = 'ENABLED'",
                                           "ad_group_ad.policy_summary.approval_status = 'DISAPPROVED'"))
    assert "asset_group.status = 'ENABLED'" in da.ASSET_GROUPS_GAQL and "campaign.status = 'ENABLED'" in da.ASSET_GROUPS_GAQL
    assert all(x in da.ASSETS_GAQL for x in ("asset_group_asset.status = 'ENABLED'", "asset_group.status = 'ENABLED'", "campaign.status = 'ENABLED'", "= 'DISAPPROVED'"))


# ── Words ───────────────────────────────────────────────────────────────────

def test_policy_topics_in_words_with_their_evidence():
    entries = [{"topic": "DESTINATION_NOT_WORKING", "type": "PROHIBITED", "evidences": [{"destinationNotWorking": {"device": "MOBILE", "httpErrorCode": 404}}]},
               {"topic": "MISREPRESENTATION", "type": "LIMITED", "evidences": [{"textList": {"texts": ["free gift", "free gift", "act now"]}}, {"websiteList": {"websites": ["x.example"]}}]}]
    t = da.topics_of(entries)
    assert t[0] == {"topic": "Destination not working", "type": "prohibited", "evidence": ["Mobile: 404"]}
    assert t[1]["evidence"] == ["free gift", "act now", "x.example"] and t[1]["type"] == "limited"
    assert da.reason_text(t) == 'Destination not working (prohibited): "Mobile: 404"; Misrepresentation (limited): "free gift", "act now"'
    assert da.reason_text([]) == "Disapproved (no reason given)"
    assert da.words("DESTINATION_NOT_WORKING") == "Destination not working"


# ── Google ──────────────────────────────────────────────────────────────────

def fake_search(rows):
    def search(cid, gaql):
        for key, name in (("FROM ad_group_ad", "ads"), ("FROM asset_group_asset", "assets"), ("FROM asset_group ", "groups")):
            if key in gaql:
                return rows.get(name, [])
        raise AssertionError(gaql)
    return search


def test_google_disapprovals_ads_asset_groups_and_assets():
    rows = {"ads": [{"campaign": {"name": "Brand"}, "adGroup": {"name": "Core"},
                     "adGroupAd": {"ad": {"id": "11", "type": "RESPONSIVE_SEARCH_AD", "finalUrls": ["https://acme-garden.com/a"]},
                                   "policySummary": {"approvalStatus": "DISAPPROVED", "policyTopicEntries": [{"topic": "TRADEMARKS", "type": "LIMITED"}]}}}],
            "groups": [{"campaign": {"name": "PMax"}, "assetGroup": {"id": "7", "name": "Landscaping", "finalUrls": ["https://acme-garden.com/"], "primaryStatusReasons": ["ASSET_GROUP_DISAPPROVED"]}},
                       {"campaign": {"name": "PMax"}, "assetGroup": {"id": "8", "name": "Other", "primaryStatusReasons": ["ASSET_GROUP_PAUSED"]}}],
            "assets": [{"campaign": {"name": "PMax"}, "assetGroup": {"id": "7", "name": "Landscaping"}, "asset": {"id": "5", "type": "TEXT", "textAsset": {"text": "Free estimates"}},
                        "assetGroupAsset": {"fieldType": "HEADLINE", "policySummary": {"policyTopicEntries": [{"topic": "SUPERLATIVES", "type": "LIMITED"}]}}}]}
    items, n = da.google_items(G, fake_search(rows))
    assert n == {"ads": 1, "assetGroups": 1, "assets": 1}
    by = {i["where"]: i for i in items}
    assert set(by) == {"ad", "asset group", "asset"}                              # the group Google paused for another reason is not a disapproval
    assert by["ad"]["severity"] == "high" and by["ad"]["adType"] == "Responsive Search Ad" and by["ad"]["campaign"] == "Brand" and by["ad"]["url"] == "https://acme-garden.com/a"
    assert by["ad"]["reason"] == "Trademarks (limited)"
    assert by["asset group"]["reason"] == "Google marks the asset group not eligible: Asset group disapproved"
    assert by["asset"]["severity"] == "medium" and by["asset"]["objectName"] == "Free estimates" and by["asset"]["adGroup"] == "Landscaping"
    assert len({i["fingerprint"] for i in items}) == 3


def test_a_failed_read_is_raised_not_counted_clean():
    def search(cid, gaql):
        if "asset_group_asset" in gaql:
            raise RuntimeError("quota")
        return []
    with pytest.raises(RuntimeError):
        da.google_items(G, search)


# ── Meta ────────────────────────────────────────────────────────────────────

def meta_row(i, status="ACTIVE", campaign="ACTIVE", adset="ACTIVE", **kw):
    return dict({"id": i, "name": "Ad " + i, "status": status, "effective_status": "DISAPPROVED", "campaign": {"name": "C", "status": campaign}, "adset": {"name": "S", "status": adset}}, **kw)


def test_meta_reports_only_live_ads_and_gives_the_reasons():
    rows = [meta_row("1", issues_info=[{"error_summary": "Ad does not follow policy", "error_message": "long text"}], ad_review_feedback={"global": {"1815": "Personal attributes"}}),
            meta_row("2", status="PAUSED"), meta_row("3", campaign="PAUSED"), meta_row("4", adset="PAUSED"), meta_row("5")]
    asked = []
    items, n = da.meta_items(M, lambda path, params: asked.append((path, params)) or rows)
    assert n == {"ads": 5} and [i["objectId"] for i in items] == ["1", "5"]
    assert items[0]["reason"] == "Ad does not follow policy; Personal attributes" and items[1]["reason"] == "Disapproved by Meta (no reason given)"
    assert asked[0][0] == "act_77/ads" and "DISAPPROVED" in asked[0][1]["effective_status"] and items[0]["campaign"] == "C" and items[0]["adGroup"] == "S"


def test_meta_asks_again_with_plain_fields_when_the_token_may_not_read_the_rest():
    calls = []

    def pages(path, params):
        calls.append(params["fields"])
        if len(calls) == 1:
            raise RuntimeError("(#10) Application does not have permission for this action")
        return [meta_row("1")]
    items, _ = da.meta_items(M, pages)
    assert len(calls) == 2 and calls[1] == da.META_FIELDS_PLAIN and len(items) == 1


# ── Findings ────────────────────────────────────────────────────────────────

def item(account, oid, where="ad", reason="Trademarks (limited)"):
    return da._item(account, where, oid, "Ad " + oid, "Responsive Search Ad", "Camp", "Group", "https://x.example/", "high", [], reason)


def test_findings_are_new_fixed_back_or_dismissed():
    a, b = item(G, "1"), item(G, "2")
    w, c = da.apply([a, b], {}, {G["id"]}, "T1")
    assert c == {"new": 2, "reopened": 0, "fixed": 0} and all(x["status"] == "open" and x["alertedAt"] is None and x["firstSeen"] == "T1" for x in w)
    existing = {x["fingerprint"]: dict(x, alertedAt="T1") for x in w}
    w, c = da.apply([a], existing, {G["id"]}, "T2")
    by = {x["objectId"]: x for x in w}
    assert c["new"] == 0 and c["fixed"] == 1 and by["1"]["alertedAt"] == "T1" and by["1"]["firstSeen"] == "T1" and by["1"]["lastSeen"] == "T2" and by["2"]["status"] == "fixed"
    existing = {x["fingerprint"]: x for x in w}
    w, c = da.apply([a, b], existing, {G["id"]}, "T3")
    assert c == {"new": 0, "reopened": 1, "fixed": 0} and next(x for x in w if x["objectId"] == "2")["alertedAt"] is None
    dismissed = {x["fingerprint"]: dict(x, status="dismissed") for x in w}
    w, c = da.apply([a, b], dismissed, {G["id"]}, "T4")
    assert {x["status"] for x in w} == {"dismissed"} and c == {"new": 0, "reopened": 0, "fixed": 0}


def test_an_account_that_could_not_be_read_does_not_get_its_findings_fixed():
    w, _ = da.apply([item(G, "1")], {}, {G["id"]}, "T1")
    existing = {x["fingerprint"]: x for x in w}
    assert da.apply([], existing, set(), "T2")[1]["fixed"] == 0
    assert da.apply([], existing, {G["id"]}, "T2")[1]["fixed"] == 1


# ── The sheet and the alert ─────────────────────────────────────────────────

def rows(*accounts_and_ids):
    items = [item(a, i) for a, i in accounts_and_ids]
    return da.apply(items, {}, {a["id"] for a, _ in accounts_and_ids}, "2026-10-19T00:00:00+00:00")[0]


def test_the_sheet_and_the_alert_words():
    from openpyxl import load_workbook
    r = rows((G, "1"), (H, "2"), (G, "3"))
    ws = load_workbook(io.BytesIO(da.sheet(r, [r[0]["fingerprint"]]))).active
    got = [[c.value for c in x] for x in ws.iter_rows()]
    assert got[0][:3] == ["Account", "Platform", "Campaign"] and len(got) == 4 and got[1][7] == "Trademarks (limited)" and got[1][2] == "Camp"
    m = da.digest(r, r[:2], LINK)
    assert m["subject"] == "2 ads newly disapproved in 2 ad accounts (3 disapproved in all)"
    assert "Only enabled campaigns and ad groups are checked" in m["text"] and LINK in m["text"] and m["attachments"][0][0] == "disapproved-ads.xlsx" and m["attachments"][0][1][:2] == b"PK"
    assert m["slack_blocks"][0]["type"] == "header" and "The sheet was emailed" in m["slack_blocks"][-1]["text"]["text"]
    assert "—" not in m["text"] + m["html"]
    assert da.digest(r[:1], r[:1], LINK)["subject"].startswith("1 ad newly disapproved in 1 ad account")
    many = rows(*[(G, str(i)) for i in range(20)])
    assert "...and 5 more in the sheet." in da.digest(many, many, LINK)["text"] and len(da.digest(many, many, LINK)["slack_blocks"]) <= 12


# ── The scan ────────────────────────────────────────────────────────────────

class World:
    def __init__(self, **disapproved):
        self.d, self.accounts, self.fail = dict(disapproved), [G, H], set()

    def inventory(self, store):
        return list(self.accounts), [], ["Meta Ads is not connected: off"]

    def read(self, account, store):
        if account["id"] in self.fail:
            raise RuntimeError("cannot read")
        out = [item(account, i) for i in self.d.get(account["id"], [])]
        return out, {"ads": len(out), "assetGroups": 0, "assets": 0}


def scan(store, world, deliver):
    return ds.run_scan(store, trigger="test", link=LINK, deliver=deliver, inventory_fn=world.inventory, read_fn=world.read, workers=2)


def configured():
    s = DaStore()
    s.state_set("da_settings", {"emails": ["ppc@example.com"], "slackChannel": "C0TEST"})
    return s


def test_a_scan_alerts_once_by_email_with_the_sheet_and_by_slack():
    s, dv = configured(), Deliver()
    w = World(**{G["id"]: ["1", "2"], H["id"]: ["3"]})
    run = scan(s, w, dv)
    assert run["new"] == 3 and run["alerted"] == 3 and run["disapproved"] == 3 and run["read"] == 2 and run["problems"] == ["Meta Ads is not connected: off"]
    assert len(dv.emails) == 1 and dv.emails[0]["to"] == ["ppc@example.com"] and dv.emails[0]["attachments"][0][0] == "disapproved-ads.xlsx"
    assert len(dv.slacks) == 1 and dv.slacks[0]["channel"] == "C0TEST"
    again = scan(s, w, dv)
    assert again["new"] == 0 and again["alerted"] == 0 and len(dv.emails) == 1
    w.d[G["id"]].append("9")
    assert scan(s, w, dv)["alerted"] == 1 and len(dv.emails) == 2


def test_until_it_has_its_own_recipients_it_uses_the_domain_alerts():
    s, dv = DaStore(), Deliver()
    s.state_set("settings", {"emails": ["domains@example.com"], "slackChannel": "C0DOMAINS"})
    assert ds.settings(s) == {"emails": ["domains@example.com"], "slackChannel": "C0DOMAINS", "own": False}
    scan(s, World(**{G["id"]: ["1"]}), dv)
    assert dv.emails[0]["to"] == ["domains@example.com"] and dv.slacks[0]["channel"] == "C0DOMAINS"
    ds.save_settings(s, {"emails": "mine@example.com", "slackChannel": ""})
    assert ds.settings(s) == {"emails": ["mine@example.com"], "slackChannel": "", "own": True}
    with pytest.raises(ds.BadRequest):
        ds.save_settings(s, {"emails": "nope"})


def test_nowhere_to_send_waits_and_a_failed_send_is_tried_again():
    s, dv = DaStore(), Deliver()
    w = World(**{G["id"]: ["1"]})
    run = scan(s, w, dv)
    assert run["alerted"] == 0 and run["email"]["status"] == "not_configured" and run["slack"]["status"] == "not_configured"
    s.state_set("da_settings", {"emails": ["ppc@example.com"], "slackChannel": ""})
    dv.fail = {"email"}
    run = scan(s, w, dv)
    assert run["alerted"] == 0 and run["email"]["status"] == "failed" and "smtp down" in run["email"]["error"]
    dv.fail = set()
    assert scan(s, w, dv)["alerted"] == 1


def test_an_unreadable_account_is_named_and_its_findings_stay_open():
    s, dv = configured(), Deliver()
    w = World(**{G["id"]: ["1"], H["id"]: ["2"]})
    scan(s, w, dv)
    w.fail.add(G["id"])
    w.d[H["id"]] = []
    run = scan(s, w, dv)
    assert run["read"] == 1 and run["unreadableCount"] == 1 and "cannot read" in run["unreadable"][0]["error"] and run["fixed"] == 1
    st = {f["accountId"]: f["status"] for f in s.da.values()}
    assert st == {G["id"]: "open", H["id"]: "fixed"}


def test_fixed_back_and_dismissed_across_scans():
    s, dv = configured(), Deliver()
    w = World(**{G["id"]: ["1"]})
    scan(s, w, dv)
    w.d[G["id"]] = []
    assert scan(s, w, dv)["fixed"] == 1
    w.d[G["id"]] = ["1"]
    run = scan(s, w, dv)
    assert run["reopened"] == 1 and run["alerted"] == 1
    fp = next(iter(s.da))
    ds.decide(s, fp, "dismissed", "me")
    w.d[G["id"]] = ["1", "2"]
    run = scan(s, w, dv)
    assert run["alerted"] == 1 and s.da[fp]["status"] == "dismissed"
    ds.decide(s, fp, "open", "me")
    assert s.da[fp]["status"] == "open"
    for bad in (("fixed", fp), ("dismissed", "nope")):
        with pytest.raises(ds.BadRequest):
            ds.decide(s, bad[1], bad[0], "me")


def test_two_scans_never_run_at_once_and_a_broken_one_is_recorded():
    s, dv = configured(), Deliver()
    w = World()
    s.claim(ds.LOCK, 60)
    assert scan(s, w, dv) == {"status": "busy"}
    s.release(ds.LOCK)
    assert scan(s, w, dv)["status"] == "done" and ds.LOCK not in s.locks and s.state_get("da_progress")["running"] is False

    def boom(store):
        raise RuntimeError("inventory down")
    run = ds.run_scan(s, trigger="t", link=LINK, deliver=dv, inventory_fn=boom, read_fn=w.read)
    assert run["status"] == "failed" and "inventory down" in run["error"] and ds.LOCK not in s.locks and s.state_get("da_runs")[0]["status"] == "failed"


def test_a_quota_error_stops_reading_google():
    class Quota(Exception):
        kind = "quota"
    s, dv = configured(), Deliver()
    seen = []

    def read(account, store):
        seen.append(account["id"])
        raise Quota("out of reads")
    run = ds.run_scan(s, trigger="t", link=LINK, deliver=dv, inventory_fn=World().inventory, read_fn=read, workers=1)
    assert seen == [G["id"]] and run["unreadableCount"] == 2 and "budget" in run["unreadable"][1]["error"]


def test_the_snapshot_the_test_alert_and_the_sheet():
    s, dv = configured(), Deliver()
    scan(s, World(**{G["id"]: ["1"]}), dv)
    snap = ds.snapshot(s, dv)
    assert snap["counts"] == {"open": 1, "dismissed": 0, "fixed": 0} and snap["findings"][0]["reason"] == "Trademarks (limited)" and snap["runs"][0]["status"] == "done"
    assert snap["senders"] == {"email": True, "slack": True} and snap["settings"]["own"] is True
    assert ds.open_sheet(s)[:2] == b"PK"
    res = ds.test_alert(s, LINK, dv)
    assert res["sent"] and dv.emails[-1]["subject"].startswith("TEST: ") and dv.slacks[-1]["text"].startswith("TEST")
    assert ds.test_alert(DaStore(), LINK, dv)["sent"] is False
