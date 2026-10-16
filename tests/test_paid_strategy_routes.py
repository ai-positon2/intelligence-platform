"""Strategy & Campaign Book: the routes. Staff only, same-origin JSON for anything that changes, a daily cap, and downloads that open."""
import io
import json
import os
import re
import sys
import zipfile

import openpyxl
import pytest

os.environ.setdefault("GOOGLE_CLIENT_ID", "test")
os.environ.setdefault("GOOGLE_CLIENT_SECRET", "test")
os.environ.setdefault("FLASK_SECRET_KEY", "test")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import app as appmod  # noqa: E402
from tests import strategy_fakes as F  # noqa: E402
from tracker import paid_gads_segments as gads  # noqa: E402
from tracker import paid_qa_book, paid_strategy as ps, paid_strategy_store as store  # noqa: E402

STAFF, OTHER, ADMIN = "someone@position2.com", "another@position2.com", "reporting@position2.com"
SAME = {"Origin": "http://localhost"}
BOM = chr(0xFEFF).encode()
API = "/p2/paid-marketing/api/strategy"
FORM = {"name": "Acme Dental", "website": "acmedental.com", "account": "1234567890", "budget": "6000", "locations": "Austin, TX", "services": "dental implants",
        "launch": "2026-11-02", "research": {"social": False}}


class World:
    def __init__(self):
        self.store = F.FakeStrategyStore()
        self.spawned = []
        self.sci = F.FakeSci(runs={})


@pytest.fixture
def world(monkeypatch):
    w = World()
    for name in ("create", "get", "get_light", "list_recent", "begin", "set_steps", "set_step", "save_research", "save_plan_part", "save_plan", "finish", "fail", "archive", "touch",
                 "record_run", "runs_today"):
        monkeypatch.setattr(store, name, getattr(w.store, name))
    monkeypatch.setattr(appmod, "_strat_spawn", lambda fn: (w.spawned.append(fn), fn()))
    monkeypatch.setattr(appmod, "_strat_deps", lambda rec, account: F.deps(sci=w.sci, account=bool(account), email=rec["created_by"]))
    monkeypatch.setattr(gads, "list_accounts", lambda **kw: [{"id": "1234567890", "name": "Acme Dental", "currency": "USD"}])
    monkeypatch.setattr(appmod, "_STRAT_LAST", {})
    monkeypatch.setenv("ANTHROPIC_API_KEY", "x")
    monkeypatch.setattr(ps, "_today", lambda: __import__("datetime").date(2026, 10, 16))
    return w


def client_as(email):
    c = appmod.app.test_client()
    if email:
        with c.session_transaction() as sess:
            sess["google_user"] = {"email": email, "name": "T"}
    return c


def post(c, path, body=None, headers=SAME):
    return c.post(API + path, data=json.dumps(body if body is not None else {}), content_type="application/json", headers=headers)


def built(c, **form):
    r = post(c, "/start", dict(FORM, **form))
    assert r.status_code == 201, r.get_json()
    return r.get_json()["id"]


# ── who may use it ──────────────────────────────────────────────────────────

def test_the_page_and_every_api_are_for_position2_staff_only(world):
    anon, outsider = client_as(None), client_as("x@gmail.com")
    for path in ("/p2/paid-marketing/campaign-book", API + "/options", API + "/list", API + "/" + "a" * 32, API + "/" + "a" * 32 + "/status", API + "/" + "a" * 32 + "/download?file=xlsx"):
        assert anon.get(path).status_code == 302, path
        r = outsider.get(path)
        assert r.status_code == 302 and "/app" in r.headers["Location"], path
    assert anon.post(API + "/start", data="{}", content_type="application/json", headers=SAME).status_code == 302


def test_the_page_renders_for_staff(world):
    html = client_as(STAFF).get("/p2/paid-marketing/campaign-book").get_data(as_text=True)
    assert "Strategy &amp; Campaign Book" in html and "campaign_book.js" in html and 'id="cards"' in html and "pm_picker.js" in html


def test_options_describe_the_cards_and_say_whether_the_model_is_there(world, monkeypatch):
    j = client_as(STAFF).get(API + "/options").get_json()
    assert [c["key"] for c in j["cards"]] == ["website", "competitors", "social", "history", "keywords", "seasonality", "market"] and j["modelReady"] is True
    assert j["goals"]["leads"] and j["kinds"]["pmax"] == "Performance Max" and j["buildsPerDay"] == 12
    monkeypatch.delenv("ANTHROPIC_API_KEY")
    assert client_as(STAFF).get(API + "/options").get_json()["modelReady"] is False


def test_everything_that_changes_needs_same_origin_json(world):
    c = client_as(STAFF)
    assert c.post(API + "/start", data=json.dumps(FORM), content_type="application/json", headers={"Origin": "https://evil.example"}).status_code == 403
    assert c.post(API + "/start", data="name=x", content_type="application/x-www-form-urlencoded", headers=SAME).status_code == 415
    sid = built(c)
    assert c.delete(API + "/" + sid, headers={"Origin": "https://evil.example"}).status_code == 403
    assert c.post(API + "/%s/rebuild" % sid, data="{}", headers=SAME).status_code == 415


# ── starting a build ────────────────────────────────────────────────────────

def test_a_build_is_started_run_and_can_be_read_back(world):
    c = client_as(STAFF)
    r = post(c, "/start", FORM)
    assert r.status_code == 201 and r.get_json()["started"] is True
    sid = r.get_json()["id"]
    rec = c.get(API + "/" + sid).get_json()
    assert rec["status"] == "done" and rec["name"] == "Acme Dental" and rec["accountName"] == "Acme Dental" and rec["createdBy"] == STAFF
    assert rec["inputs"]["currency"] == "USD" and rec["inputs"]["account"] == "1234567890" and [c["name"] for c in rec["plan"]["campaigns"]][0].startswith("P2_Search_")
    assert "text" not in rec["research"]["website"]["data"]["pages"][0] and rec["cost"]["calls"] == 0
    light = c.get(API + "/%s/status" % sid).get_json()
    assert light["status"] == "done" and "plan" not in light and light["progress"]["steps"][-1]["key"] == "f:finish"


def test_bad_details_are_refused_with_a_sentence(world):
    c = client_as(STAFF)
    for body, msg in (({"website": "x.com"}, "name"), (dict(FORM, goal="fame"), "goal"), (dict(FORM, budget="lots"), "number"), ("not an object", "form")):
        r = post(c, "/start", body)
        assert r.status_code == 400 and msg in r.get_json()["error"], body
    r = post(c, "/start", dict(FORM, account="1111111111"))
    assert r.status_code == 400 and "not one this platform can read" in r.get_json()["error"]
    assert world.store.rows == {}


def test_a_build_without_an_account_works_and_takes_the_currency_given(world):
    c = client_as(STAFF)
    sid = built(c, account="", currency="cad")
    rec = c.get(API + "/" + sid).get_json()
    assert rec["inputs"]["currency"] == "CAD" and rec["status"] == "done" and rec["research"]["history"]["state"] == "skipped" if "history" in rec["research"] else True


def test_the_list_shows_every_build_newest_first(world):
    c = client_as(STAFF)
    a, b = built(c, name="First Client"), built(c, name="Second Client")
    rows = client_as(OTHER).get(API + "/list").get_json()["builds"]
    assert [r["name"] for r in rows] == ["Second Client", "First Client"] and rows[0]["status"] == "done" and rows[0]["campaigns"] == 2 and rows[0]["createdBy"] == STAFF


def test_a_person_can_only_start_so_many_builds_a_day(world, monkeypatch):
    monkeypatch.setattr(ps, "BUILDS_PER_DAY", 2)
    c = client_as(STAFF)
    built(c), built(c)
    r = post(c, "/start", FORM)
    assert r.status_code == 429 and r.get_json()["kind"] == "cap" and "2 builds in a day" in r.get_json()["error"]
    assert len(client_as(STAFF).get(API + "/list").get_json()["builds"]) == 2           # the refused one is not left in the list
    other = post(client_as(OTHER), "/start", FORM)
    assert other.status_code == 201


def test_building_again_counts_towards_the_days_builds(world, monkeypatch):
    monkeypatch.setattr(ps, "BUILDS_PER_DAY", 2)
    c = client_as(STAFF)
    sid = built(c)
    appmod._STRAT_LAST.clear()
    assert post(c, "/%s/rebuild" % sid, {}).status_code == 200
    appmod._STRAT_LAST.clear()
    r = post(c, "/%s/rebuild" % sid, {})
    assert r.status_code == 429 and r.get_json()["kind"] == "cap"
    assert world.store.runs_today(STAFF) == 2


def test_only_the_person_who_started_a_build_or_an_admin_can_build_it_again(world):
    sid = built(client_as(STAFF))
    appmod._STRAT_LAST.clear()
    r = post(client_as(OTHER), "/%s/rebuild" % sid, {"resume": False})
    assert r.status_code == 403 and "Only whoever started this build" in r.get_json()["error"]
    assert world.store.get(sid)["plan"]["campaigns"]                          # nothing was wiped
    assert post(client_as(ADMIN), "/%s/rebuild" % sid, {"resume": True}).status_code == 200
    appmod._STRAT_LAST.clear()
    assert post(client_as(STAFF), "/%s/rebuild" % sid, {}).status_code == 200


def test_a_build_that_is_running_is_not_started_twice(world):
    c = client_as(STAFF)
    rec = world.store.create(by=STAFF, inputs=F.clean())
    world.store.begin(rec["id"])
    r = post(c, "/%s/rebuild" % rec["id"], {})
    assert r.status_code == 409 and "already running" in r.get_json()["error"]


def test_the_page_polls_a_light_status_and_unknown_ids_are_404(world):
    c = client_as(STAFF)
    assert c.get(API + "/" + "0" * 32 + "/status").status_code == 404 and c.get(API + "/not-an-id").status_code == 404 and c.get(API + "/" + "0" * 32).status_code == 404


# ── building again and removing ─────────────────────────────────────────────

def test_build_again_is_rate_limited_and_resume_keeps_the_work(world):
    c = client_as(STAFF)
    sid = built(c)
    first = post(c, "/%s/rebuild" % sid, {"resume": False})
    assert first.status_code == 200 and first.get_json() == {"id": sid, "started": True}
    again = post(c, "/%s/rebuild" % sid, {"resume": True})
    assert again.status_code == 429 and again.get_json()["kind"] == "cooldown" and again.get_json()["retryIn"] >= 1
    appmod._STRAT_LAST.clear()
    calls = []
    world.spawned.clear()
    appmod._strat_deps = lambda rec, account: F.deps(sci=world.sci, account=bool(account), call=F.model(calls=calls), email=rec["created_by"])
    assert post(c, "/%s/rebuild" % sid, {"resume": True}).status_code == 200 and calls == []        # everything was already done
    assert c.get(API + "/" + sid).get_json()["status"] == "done"


def test_only_the_person_who_started_a_build_or_an_admin_can_remove_it(world):
    sid = built(client_as(STAFF))
    assert client_as(OTHER).delete(API + "/" + sid, headers=SAME, content_type="application/json").status_code == 403
    assert client_as(STAFF).delete(API + "/" + sid, headers=SAME, content_type="application/json").get_json() == {"ok": True}
    assert client_as(STAFF).get(API + "/" + sid).status_code == 404 and client_as(STAFF).get(API + "/list").get_json()["builds"] == []
    other = built(client_as(STAFF))
    assert client_as(ADMIN).delete(API + "/" + other, headers=SAME, content_type="application/json").status_code == 200


def test_a_store_that_is_down_is_said_in_words(world, monkeypatch):
    def down(*a, **k):
        raise store.StoreUnavailable("no db")
    for name in ("list_recent", "create", "get", "get_light"):
        monkeypatch.setattr(store, name, down)
    c = client_as(STAFF)
    assert c.get(API + "/list").status_code == 503 and post(c, "/start", FORM).status_code == 503
    assert c.get(API + "/" + "a" * 32).status_code == 503 and c.get(API + "/" + "a" * 32 + "/status").status_code == 503
    assert "store is unavailable" in c.get(API + "/list").get_json()["error"]


# ── the downloads ───────────────────────────────────────────────────────────

def dl(c, sid, file, **q):
    qs = "&".join("%s=%s" % kv for kv in dict(file=file, **q).items())
    return c.get(API + "/%s/download?%s" % (sid, qs))


def test_excel_csv_pdf_and_zip_download_as_files_that_open(world):
    c = client_as(STAFF)
    sid = built(c)
    x = dl(c, sid, "xlsx")
    assert x.status_code == 200 and x.headers["Content-Type"].startswith("application/vnd.openxmlformats") and 'filename="Acme-Dental-strategy-and-campaign-book.xlsx"' in x.headers["Content-Disposition"]
    assert x.headers["Cache-Control"] == "private, no-store"
    wb = openpyxl.load_workbook(io.BytesIO(x.data))
    assert "Overview" in wb.sheetnames and any(n.startswith("Book ") for n in wb.sheetnames)
    v = dl(c, sid, "csv")
    assert v.status_code == 200 and v.headers["Content-Type"].startswith("text/csv") and v.data.startswith(BOM) and b"Campaign,Section,Field,Value" in v.data
    assert b"P2_Search_Implants_20261102" in v.data and b"P2_PMax_Local_20261102" in v.data
    p = dl(c, sid, "pdf")
    assert p.status_code == 200 and p.data[:5] == b"%PDF-" and p.headers["Content-Type"] == "application/pdf" and ".pdf" in p.headers["Content-Disposition"]
    z = dl(c, sid, "zip")
    names = zipfile.ZipFile(io.BytesIO(z.data)).namelist()
    assert z.status_code == 200 and "overview.csv" in names and any(n.startswith("campaign-book-p2-search") for n in names)


def test_a_campaigns_book_downloads_in_the_layout_the_qa_agent_reads(world):
    c = client_as(STAFF)
    sid = built(c)
    key = c.get(API + "/" + sid).get_json()["plan"]["campaigns"][0]["key"]
    x = dl(c, sid, "book-xlsx", campaign=key)
    book = paid_qa_book.parse_xlsx(x.data, "book.xlsx")
    assert x.status_code == 200 and len(book["rows"]) == 61 and paid_qa_book.lookup(book, "Campaign Name") == "P2_Search_Implants_20261102"
    assert paid_qa_book.lookup(book, "Bid Strategy Type").startswith("Maximize conversions") and paid_qa_book.lookup(book, "Location Settings") == "Austin, Texas"
    v = dl(c, sid, "book-csv", campaign=key)
    book2 = paid_qa_book.parse_csv(v.data.lstrip(BOM), "book.csv")
    assert v.status_code == 200 and v.data.startswith(BOM) and len(book2["rows"]) == 61 and paid_qa_book.lookup(book2, "Devices") == "All devices, no bid adjustments"
    assert dl(c, sid, "book-xlsx", campaign="nope").status_code == 404 and dl(c, sid, "book-csv").status_code == 404


def test_bad_download_requests_are_refused(world):
    c = client_as(STAFF)
    sid = built(c)
    assert dl(c, sid, "exe").status_code == 400 and c.get(API + "/%s/download" % sid).status_code == 400
    assert dl(c, "0" * 32, "xlsx").status_code == 404 and dl(c, "nope", "xlsx").status_code == 404
    empty = world.store.create(by=STAFF, inputs=F.clean())
    r = dl(c, empty["id"], "xlsx")
    assert r.status_code == 409 and "nothing to download" in r.get_json()["error"]


def test_a_file_that_cannot_be_made_is_said_in_words(world, monkeypatch):
    from tracker import paid_strategy_export as export
    c = client_as(STAFF)
    sid = built(c)
    monkeypatch.setattr(export, "pdf", lambda rec: (_ for _ in ()).throw(RuntimeError("fonts missing")))
    r = dl(c, sid, "pdf")
    assert r.status_code == 500 and "could not be made (RuntimeError)" in r.get_json()["error"]


# ── the wiring ──────────────────────────────────────────────────────────────

def test_every_address_the_page_script_uses_is_a_route():
    js = open(os.path.join(ROOT, "static", "js", "campaign_book.js"), encoding="utf-8").read()
    rules = {r.rule for r in appmod.app.url_map.iter_rules()}
    used = set(re.findall(r'API\+"(/[a-z]+)"', js)) | {"/<sid>/status", "/<sid>/rebuild", "/<sid>/download"}
    for path in used:
        assert any(r.startswith("/p2/paid-marketing/api/strategy") and r.endswith(path) for r in rules), path
    assert "/p2/paid-marketing/api/segments/accounts" in rules


def test_the_agent_is_on_the_paid_marketing_page_the_hub_and_the_palette():
    root = os.path.join(ROOT, "templates")
    for name in ("paid_marketing.html", "_paid_kpal.html"):
        assert "/p2/paid-marketing/campaign-book" in open(os.path.join(root, name), encoding="utf-8").read(), name


def test_the_real_dependencies_wire_up_without_touching_the_network(monkeypatch):
    monkeypatch.undo()
    monkeypatch.setattr(gads, "list_accounts", lambda **kw: [{"id": "1234567890", "name": "Acme", "currency": "USD"}])
    monkeypatch.setenv("ANTHROPIC_API_KEY", "x")
    deps = appmod._strat_deps({"created_by": STAFF}, None)
    assert deps["planner_account"] == "1234567890" and deps["model_ready"] is True and deps["email"] == STAFF
    assert all(callable(deps[k]) for k in ("search", "geo_names", "sqr_read", "cached", "post", "season_report", "spawn"))
    assert deps["sci_store"].__name__ == "tracker.sci_store" and deps["sci_pipeline"].__name__ == "tracker.sci_pipeline"
    assert appmod._strat_deps({"created_by": STAFF}, {"id": "2345678901"})["planner_account"] == "2345678901"


def test_the_accounts_own_currency_wins_over_the_one_typed_in(world):
    c = client_as(STAFF)
    sid = built(c, currency="cad")
    assert c.get(API + "/" + sid).get_json()["inputs"]["currency"] == "USD"
