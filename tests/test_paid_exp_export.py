"""Test & Experiment Tracker: the downloads (tracker/paid_exp_export.py): every file is drawn from the same tables, opens, and cannot run a formula."""
import csv
import io
import json
import zipfile

import openpyxl
import pytest

from tests.test_paid_exp import EARLY, FLAT, LOSE, NOW, TODAY, WIN, ad, fresh, group, made, read_with  # noqa: F401  (fresh: the autouse fixture that forgets saved reads)
from tracker import paid_exp as ex_
from tracker import paid_exp_export as ex


def make_report():
    exps = [made(**WIN), made(**LOSE), made(**FLAT), made(start="2026-10-08", window=("2026-10-08", "2026-10-17"), **EARLY), made(end="2026-10-10", **WIN),
            made(measured=False, control=None, trial=None, typ="AD_VARIATION"), made(state="ended", **WIN)]
    exps[0]["name"] = "=cmd|' /C calc'!A0"
    for i, e in enumerate(exps):
        e["id"] = str(i + 1)
        if i:
            e["name"] = "Test number %d" % (i + 1)
    groups = [group(ad(1, 20000, 1000, 100), ad(2, 20000, 1000, 50), ad(3, 20000, 1000, 55)), group(ad(4, 20000, 1000, 50), ad(5, 20000, 1000, 52), name="Cleanings", campaign="Search Cleanings"),
              group(ad(6, 900, 40, 0), ad(7, 800, 30, 0), name="Braces", campaign="Search Braces")]
    for i, g in enumerate(groups):
        g["id"] = str(10 + i)
    groups[0]["ads"][0]["headlines"] = ["+cmd headline", "@sum"]
    return ex_.assemble(read_with(exps, groups, tests=3), now=NOW, today=TODAY)


@pytest.fixture()
def report():
    return make_report()


def test_the_tables_come_in_the_order_they_are_shown(report):
    assert list(ex.tables(report)) == ["Summary", "Needs a decision", "Experiments", "Experiment days", "Ad tests", "Ads", "Notes"]


def test_the_summary_says_who_what_and_how_many(report):
    header, rows = ex.tables(report)["Summary"]
    d = dict((r[0], r[1]) for r in rows)
    assert header == ["Item", "Detail"]
    assert d["Google Ads account"] == "Acme (1112223334)" and d["Report made on"] == "2026-10-18" and d["Read on"] == "2026-10-18 12:00" and d["Ad tests cover"] == "the last 30 days"
    assert d["Experiments"] == 7 and d["Tests that need a decision"] == report["summary"]["needsAction"] and d["Ad groups with two or more live ads"] == 3
    assert d["Spend in experiments (USD)"] == report["summary"]["experimentSpend"]


def test_every_row_is_as_wide_as_its_header(report):
    for name, (header, rows) in ex.tables(report).items():
        assert rows, name
        for r in rows:
            assert len(r) == len(header), (name, r)
    assert [r[0] for r in ex.tables(report)["Needs a decision"][1]] == list(range(1, len(ex.tables(report)["Needs a decision"][1]) + 1))


def test_the_needs_sheet_names_the_decision_the_test_and_the_money(report):
    header, rows = ex.tables(report)["Needs a decision"]
    first = dict(zip(header, rows[0]))
    assert first["What to do"] == "Decide now" and first["Kind of test"] == "Experiment" and first["Test"] and first["Detail"] and first["Spend (USD)"] > 0
    assert {r[2] for r in rows} == {"Experiment", "Ad test"}


def test_the_experiments_sheet_has_one_row_per_experiment_with_the_verdict_and_both_arms(report):
    header, rows = ex.tables(report)["Experiments"]
    assert len(rows) == 7
    by = {r[0]: dict(zip(header, r)) for r in rows}
    win = next(v for v in by.values() if v["Verdict"] == "Trial is ahead" and v["Status"] == "Running")
    assert win["Compared on"] == "Conversion rate" and win["Control rate %"] == 5.0 and win["Trial rate %"] == 9.0 and win["Trial change %"] == 80.0 and win["Confidence %"] > 95
    assert win["Control clicks"] == 1000 and win["Trial conversions"] == 90 and win["Start"] == "2026-09-20" and win["Days of results"] == 28 and win["What to do"]
    assert {v["Verdict"] for v in by.values()} == {"Trial is ahead", "Trial is behind", "No difference", "Too early", "Not measured"}
    unmeasured = next(v for v in by.values() if v["Verdict"] == "Not measured")
    assert unmeasured["Control rate %"] is None and unmeasured["Type"] == "Search custom"
    assert {v["Status"] for v in by.values()} == {"Running", "Ended"}


def test_the_experiment_days_sheet_has_every_arms_every_day(report):
    header, rows = ex.tables(report)["Experiment days"]
    assert header[:4] == ["Experiment", "Arm", "Role", "Date"] and {r[2] for r in rows} == {"Control", "Trial"} and rows[0][3] == "2026-10-01"
    assert len(rows) == 2 * 6                                           # six measured experiments, one day each in the test data, two arms


def test_the_ad_sheets_say_which_ad_leads_and_which_are_behind(report):
    th, tr = ex.tables(report)["Ad tests"]
    t = {r[0]: dict(zip(th, r)) for r in tr}
    assert t["Implants"]["Verdict"] == "Clear leader" and t["Implants"]["Live ads"] == 3 and t["Implants"]["Compared on"] == "Conversion rate" and t["Implants"]["Ads compared"] == 3
    assert t["Cleanings"]["Verdict"] == "No difference" and t["Braces"]["Verdict"] == "Too early"
    ah, ar = ex.tables(report)["Ads"]
    ads = [dict(zip(ah, r)) for r in ar if r[0] == "Implants"]
    lead = next(a for a in ads if a["Leader"] == "Yes")
    assert lead["Verdict"] == "Clear leader" and lead["Conversion rate %"] == 10.0 and lead["Impressions"] == 20000 and sum(1 for a in ads if a["Verdict"] == "Behind") == 2
    behind = next(a for a in ads if a["Verdict"] == "Behind")
    assert behind["Against the leader %"] < 0 and behind["Confidence %"] > 90


def test_a_value_a_spreadsheet_would_run_never_runs(report):
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(report)))
    for ws in wb:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith(("=", "+", "@"))) or c.data_type == "s", (ws.title, c.value)
    cell = next(c for row in wb["Experiments"].iter_rows(min_row=2) for c in row if isinstance(c.value, str) and "calc" in c.value)
    assert cell.data_type == "s"
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(report)))
    cells = [c for name in z.namelist() for r in csv.reader(io.StringIO(z.read(name).decode())) for c in r]
    assert not any(c.startswith(("=", "+", "@")) for c in cells)
    assert any("calc" in c and c.startswith("'") for c in cells) and any("cmd headline" in c and c.startswith("'") for c in cells)


def test_the_workbook_has_a_sheet_per_table_with_verdicts_coloured(report):
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(report)))
    assert wb.sheetnames == list(ex.tables(report))
    ws = wb["Experiments"]
    header = [c.value for c in ws[1]]
    col = header.index("Verdict") + 1
    fills = {ws.cell(row=r, column=col).value: ws.cell(row=r, column=col).fill.fgColor.rgb for r in range(2, ws.max_row + 1)}
    assert fills["Trial is ahead"].endswith("DDF3E2") and fills["Trial is behind"].endswith("FFD9D6") and fills["Too early"].endswith("E1ECFA") and fills["No difference"].endswith("ECEFF5")
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and wb["Needs a decision"].cell(row=2, column=1).value == 1
    needs = wb["Needs a decision"]
    nh = [c.value for c in needs[1]]
    kinds = {needs.cell(row=r, column=nh.index("What to do") + 1).value: needs.cell(row=r, column=nh.index("What to do") + 1).fill.fgColor.rgb for r in range(2, needs.max_row + 1)}
    assert kinds["Decide now"].endswith("FFD9D6") and kinds["Apply it"].endswith("DDF3E2")


def test_the_csv_zip_has_one_file_per_table(report):
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(report)))
    assert z.namelist() == [t.lower().replace(" ", "-") + ".csv" for t in ex.tables(report)]
    rows = list(csv.reader(io.StringIO(z.read("experiments.csv").decode())))
    assert rows[0][0] == "Experiment" and len(rows) == 8


def test_a_report_with_nothing_to_track_still_makes_every_file():
    empty = ex_.assemble(read_with([], []), now=NOW, today=TODAY)
    assert list(ex.tables(empty)) == ["Summary", "Notes"]
    assert ex.xlsx(empty)[:2] == b"PK" and ex.csv_zip(empty)[:2] == b"PK" and ex.pdf(empty)[:5] == b"%PDF-"


def test_the_pdf_is_a_document(report):
    pdf = ex.pdf(report)
    assert pdf[:5] == b"%PDF-" and len(pdf) > 5000


def test_the_pdf_lists_only_the_ad_tests_that_spent_most():
    groups = [group(ad(i * 2, 20000, 1000, 100, cost=float(i + 1)), ad(i * 2 + 1, 20000, 1000, 50), name="G%d" % i) for i in range(ex.PDF_AD_TESTS + 5)]
    for i, g in enumerate(groups):
        g["id"] = str(i)
    r = ex_.assemble(read_with([], groups), now=NOW, today=TODAY)
    assert len(r["adTests"]) == ex.PDF_AD_TESTS + 5 and ex.pdf(r)[:5] == b"%PDF-"
    rows, total = ex.pdf_ad_rows(r)
    assert ex.PDF_AD_TESTS == 40 and len(rows) == 40 and total == 45 and all(len(x) == 5 for x in rows)
    few, n = ex.pdf_ad_rows(make_report())
    assert len(few) == n == 3 and few[0][0] == "Implants (Search Implants)" and few[0][1] == 3 and few[0][3] == "Clear leader"
    assert ex.pdf_ad_rows({}) == ([], 0)


def test_the_file_stem_names_the_account_safely(report):
    assert ex.file_stem(report) == "Acme-test-and-experiment-tracker"
    assert ex.file_stem({"account": {}}) == "account-test-and-experiment-tracker" and ex.file_stem({}) == "account-test-and-experiment-tracker"
    assert ex.file_stem({"account": {"name": "../../etc/passwd  &x", "id": "1"}}) == "etc-passwd-x-test-and-experiment-tracker"
    assert ex.file_stem({"account": {"id": "1112223334"}}) == "1112223334-test-and-experiment-tracker"
    assert ex.file_stem({"account": {"name": "x" * 100}}) == "x" * 60 + "-test-and-experiment-tracker"


def test_no_em_or_en_dash_in_any_file_text(report):
    for name, (header, rows) in ex.tables(report).items():
        for cell in header + [c for r in rows for c in r]:
            assert "—" not in str(cell) and "–" not in str(cell), (name, cell)
    assert "—" not in json.dumps(ex.tables(report), ensure_ascii=False)


def test_rates_are_rounded_to_two_places_and_empty_stays_empty():
    assert ex._rate(5.126) == 5.13 and ex._rate(None) is None and ex._rate("7") == 7.0


def test_the_summary_lines_carry_the_report_numbers_in_the_right_places(report):
    s, d = report["summary"], dict(ex.tables(report)["Summary"][1])
    assert d["Running and the trial is ahead"] == s["winners"] and d["Running and the trial is behind"] == s["behind"] and d["Running with no difference"] == s["noDifference"]
    assert d["Running and too early to call"] == s["tooEarly"] and d["Ad tests with a clear leader"] == s["adConclusive"] and d["Ad tests with no difference"] == s["adNoDifference"]
    assert d["Ad tests too early to call"] == s["adTooEarly"] and d["Ad tests shown"] == s["adTestsShown"] and d["Spend in the ad tests shown (USD)"] == s["adSpend"] and d["Experiments running"] == s["running"]
    assert d["Running past their end date"] == s["overdue"] and d["Ending within 7 days"] == s["endsSoon"] and d["Experiments ended"] == s["ended"]
    assert s["winners"] != s["behind"]
    mapped = dict(ex.tables({"summary": {"adConclusive": 7, "adNoDifference": 2, "adTooEarly": 3, "winners": 5, "behind": 4}, "account": {"id": "1"}})["Summary"][1])
    assert mapped["Ad tests with a clear leader"] == 7 and mapped["Ad tests with no difference"] == 2 and mapped["Ad tests too early to call"] == 3
    assert mapped["Running and the trial is ahead"] == 5 and mapped["Running and the trial is behind"] == 4


def test_a_summary_leaves_out_lines_the_report_has_no_number_for():
    rows = ex.tables({"summary": {}, "account": {"id": "1", "name": "A"}})["Summary"][1]
    assert [r[0] for r in rows] == ["Google Ads account"] or all(r[1] not in (None, "") for r in rows)
    assert all(r[1] not in (None, "") for r in rows)


def test_each_arms_numbers_land_in_its_own_columns(report):
    header, rows = ex.tables(report)["Experiments"]
    win = next(dict(zip(header, r)) for r in rows if r[6] == "Trial is ahead" and r[1] == "Running")
    assert win["Control conversions"] == 50 and win["Trial conversions"] == 90 and win["Control spend"] == 100.0 and win["Trial impressions"] == 20000


def test_a_day_row_says_whether_it_was_the_control_or_the_trial(report):
    header, rows = ex.tables(report)["Experiment days"]
    assert {(r[1], r[2]) for r in rows} == {("Base", "Control"), ("Trial", "Trial")}


def test_the_ad_test_row_counts_every_live_ad_not_only_the_ones_shown():
    g = group(ad(1, 20000, 1000, 100), ad(2, 20000, 1000, 50))
    g["adCount"] = 9
    r = ex_.assemble(read_with([], [g]), now=NOW, today=TODAY)
    header, rows = ex.tables(r)["Ad tests"]
    assert dict(zip(header, rows[0]))["Live ads"] == 9


def test_a_report_made_from_an_old_saved_read_says_so_in_the_notes(report):
    report["stale"] = [{"platform": "google", "savedAt": "2026-10-10T08:00:00+00:00"}]
    notes = ex.tables(report)["Notes"][1]
    assert ["Saved read", "The read is the one saved on 2026-10-10 08:00 because Google Ads would not answer now."] in notes


def test_the_pdf_says_when_it_lists_fewer_ad_tests_than_the_excel_file_holds():
    groups = [group(ad(i * 2, 20000, 1000, 100), ad(i * 2 + 1, 20000, 1000, 50), name="G%d" % i) for i in range(ex.PDF_AD_TESTS + 1)]
    for i, g in enumerate(groups):
        g["id"] = str(i)
    assert ex.pdf_ad_note(ex_.assemble(read_with([], groups), now=NOW, today=TODAY)) == "The 40 ad tests that spent most are listed here. The Excel file has every ad test that was read."
    assert ex.pdf_ad_note(make_report()) == "" and ex.pdf_ad_note({}) == ""


def _pdf_text(report):
    import pdfplumber
    return "\n".join(p.extract_text() or "" for p in pdfplumber.open(io.BytesIO(ex.pdf(report))).pages)


def test_the_pdf_shows_the_summary_the_decisions_the_experiments_and_the_ad_tests(report):
    text = _pdf_text(report)
    for want in ("Test & Experiment Tracker", "Acme (1112223334)", "Summary", "Needs a decision", "Experiments", "Ad copy tests", "Trial is ahead", "Clear leader", "Decide now",
                 "Implants (Search Implants)", "Test number 2"):
        assert want in text, want
    assert "Notes" in text


def test_the_pdf_carries_the_note_about_the_ad_tests_it_left_out():
    groups = [group(ad(i * 2, 20000, 1000, 100), ad(i * 2 + 1, 20000, 1000, 50), name="G%d" % i) for i in range(ex.PDF_AD_TESTS + 1)]
    for i, g in enumerate(groups):
        g["id"] = str(i)
    assert "ad tests that spent most are listed here" in " ".join(_pdf_text(ex_.assemble(read_with([], groups), now=NOW, today=TODAY)).split())
    assert "listed here" not in _pdf_text(make_report())
