"""Business Research: the downloads (tracker/paid_business_export.py). One set of tables drawn as Excel, CSV and PDF."""
import csv
import io
import json
import zipfile

import openpyxl

from tests import business_fakes as B, strategy_fakes as F
from tracker import paid_business as pb, paid_business_export as ex, paid_business_job as job


def built(**dkw):
    store = F.FakeStrategyStore()
    inp = B.clean()
    rec = store.create(by="ana@position2.com", inputs=inp, account_name="Acme Dental", steps=pb.steps_for(inp))
    store.begin(rec["id"])
    job.run_build(rec["id"], store=store, deps=B.deps(**dkw))
    return store.get(rec["id"])


def test_the_tables_cover_the_whole_report_in_reading_order():
    t = ex.tables(built())
    assert list(t) == ["Overview", "Business", "Offerings", "USPs and proof", "Audiences", "Journey and objections", "Voice, places and seasons", "Rules that touch the ads",
                       "Facts the ads may state", "What the site gets wrong", "Reputation", "Tracking and set-up", "What it means for paid media", "Competitors", "Social media",
                       "Questions for the client", "Risks and next steps", "Checks"]
    overview = dict(t["Overview"][1])
    assert overview["Client"] == "Acme Dental" and overview["Google Ads account"] == "Acme Dental" and overview["Goal of the advertising"] == "Leads (forms and enquiries)"
    assert [r[0] for r in t["Overview"][1]].count("In short") == 8
    assert t["Offerings"][1][0][:6] == ["Dental implants", "Single and full-arch implants", "Adults missing teeth", "From $1,900 per tooth", "https://acmedental.com/services/dental-implants", "high"]
    assert [r[6] for r in t["USPs and proof"][1] if r[0] == "USP"] == ["Yes", "Not found on a page", ""]
    assert [r[1] for r in t["USPs and proof"][1] if r[0] == "Proof"] == ["4.9 rating from 312 reviews"]
    assert [r[3] for r in t["Facts the ads may state"][1]] == ["Yes", "Not found on a page"]
    assert [r[0] for r in t["Questions for the client"][1]] == ["Goals and targets", "Offers and pricing", "Goals and targets"]
    kinds = {r[0] for r in t["Reputation"][1]}
    assert {"Rating", "Customers praise", "Customers complain about", "Customers ask", "Press", "Award", "Person", "Source", "Read of it"} <= kinds
    assert any(r[0] == "Not seen" and "call tracking" in r[1] for r in t["Tracking and set-up"][1])
    assert [r[0] for r in t["Checks"][1]] == ["Facts not found on a page", "Points of difference whose quote was not found on a page"]


def test_a_report_with_little_in_it_leaves_out_the_empty_tables():
    rec = built()
    rec["plan"], rec["research"] = {"core": {"summary": "A dentist"}}, {}
    t = ex.tables(rec)
    assert list(t) == ["Overview", "Business"] and t["Business"][1] == [["Summary", "A dentist"]]
    assert ex.tables({}) == {"Overview": (["Item", "Detail"], []), "Business": (["Item", "Detail"], [])}


def test_the_workbook_opens_with_a_sheet_for_every_table_and_text_stays_text():
    rec = built()
    rec["plan"]["core"]["summary"] = "=HYPERLINK(\"http://evil.example\",\"x\")"
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(rec)))
    assert wb.sheetnames == [n[:31] for n in ex.tables(rec)]
    cell = [r for r in wb["Business"].iter_rows(min_row=2, values_only=False) if r[0].value == "Summary"][0][1]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    assert wb["Offerings"]["A1"].value == "Offering" and wb["Offerings"].freeze_panes == "A2"


def test_the_zip_holds_every_table_as_a_csv_with_formulas_defused():
    rec = built()
    rec["plan"]["implications"]["questions"][0]["question"] = "+SUM(1,1) is a formula"
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(rec)))
    assert len(z.namelist()) == len(ex.tables(rec)) and "overview.csv" in z.namelist() and "questions-for-the-client.csv" in z.namelist()
    rows = list(csv.reader(io.StringIO(z.read("questions-for-the-client.csv").decode())))
    assert rows[0] == ["Area", "Question", "Why it matters", "When"] and rows[1][1] == "'+SUM(1,1) is a formula"


def test_the_questions_alone_are_one_csv_or_nothing():
    rec = built()
    rows = list(csv.reader(io.StringIO(ex.csv_questions(rec))))
    assert len(rows) == 4 and rows[2][3] == "First month"
    rec["plan"]["implications"]["questions"] = []
    assert ex.csv_questions(rec) is None


def test_the_pdf_is_a_real_document_that_holds_the_report():
    rec = built()
    body = ex.pdf(rec)
    assert body[:5] == b"%PDF-" and len(body) > 5000
    huge = built()
    huge["plan"]["core"]["offerings"][0]["description"] = "word " * 20000                        # a cell taller than a page must still print
    assert ex.pdf(huge)[:5] == b"%PDF-"


def test_file_names_are_safe():
    assert ex.file_stem({"name": "A/B  Dental!!"}) == "A-B-Dental-business-research" and ex.file_stem({}) == "client-business-research"
    assert len(ex.file_stem({"name": "x" * 500})) <= 50 + len("-business-research")


def test_no_dash_is_in_any_table():
    rec = built()
    text = json.dumps(ex.tables(rec), ensure_ascii=False)
    assert chr(0x2014) not in text and chr(0x2013) not in text


def test_a_cell_too_long_for_a_page_is_cut_in_the_pdf_only():
    import pdfplumber
    from tracker.paid_strategy_export import CELL_CHARS
    rec = built()
    rec["plan"]["core"]["offerings"][0]["description"] = "Z" * (CELL_CHARS + 5000)
    shown = "\n".join(p.extract_text() or "" for p in pdfplumber.open(io.BytesIO(ex.pdf(rec))).pages)
    assert shown.count("Z") == CELL_CHARS                                                                        # the PDF holds CELL_CHARS of it
    book = openpyxl.load_workbook(io.BytesIO(ex.xlsx(rec)))
    assert len(book["Offerings"]["B2"].value) == CELL_CHARS + 5000                                                # the workbook holds all of it


def test_the_export_says_where_each_proof_came_from_and_which_ratings_nothing_confirms():
    rec = built()
    rec["plan"]["core"]["usps"][0].update(verified=False, viaSearch=True)
    rec["research"]["reputation"]["data"]["ratings"][0]["confirmed"] = True
    rec["research"]["reputation"]["data"]["ratings"][1]["confirmed"] = False
    t = ex.tables(rec)
    assert [r[6] for r in t["USPs and proof"][1] if r[0] == "USP"] == ["Only in what the web search found", "Not found on a page", ""]
    ratings = [r[1] for r in t["Reputation"][1] if r[0] == "Rating"]
    assert ratings[0] == "Google: 4.8 stars from 1,204 reviews" and ratings[1] == "Yelp: 4.5 stars from 212 reviews (no page the search read confirms this)"
