"""Ad Copy & Creative Generator: the downloads (tracker/paid_adcopy_export.py). Every file is drawn from one set of tables; these tests build a real result with the fake model
and read the files back."""
import copy
import csv
import io
import zipfile

import openpyxl
import pytest

from tests import adcopy_fakes as A, strategy_fakes as F
from tracker import paid_adcopy as pa, paid_adcopy_export as ex, paid_adcopy_job as job, paid_strategy as ps


def build(make=None, model=None):
    store = F.FakeStrategyStore()
    inputs = A.clean(**(make or {}))
    sid = store.create(by="ana@position2.com", inputs=inputs, account_name="Acme Dental", steps=pa.steps_for(inputs))["id"]
    store.begin(sid)
    job.run_build(sid, store=store, deps=A.deps(model or A.Model()))
    rec = store.get(sid)
    assert rec["status"] == "done", rec["error"]
    return rec


@pytest.fixture(scope="module")
def rec():
    over = A.answer("google_rsa")
    over["headlines"][3] = {"text": "Free Whitening Today", "angle": "offer"}                  # a claim to confirm
    return build({"formats": ["google_rsa", "google_pmax", "meta_feed"], "campaignName": "P2_Search_Implants_20261101", "adGroupName": "Implants"}, A.Model({"google_rsa": over}))


def rows_of(table, rec):
    return ex.tables(rec)[table][1]


def test_the_tables_are_the_sheets_in_order(rec):
    assert list(ex.tables(rec)) == ["Overview", "Google Search", "Performance Max", "Meta feed", "Creative briefs", "To confirm", "Facts used", "Checks"]
    ov = dict(ex.tables(rec)["Overview"][1])
    assert ov["Client"] == "Acme Dental" and ov["Goal of the advertising"] == ps.GOALS["leads"] and ov["Lines to confirm with the client first"] == 1
    assert ov["Formats"].startswith("Google Search ad")


def test_a_format_sheet_lists_every_line_with_its_length_limit_and_status(rec):
    rows = rows_of("Google Search", rec)
    heads = [r for r in rows if r[0] == "Headlines"]
    assert len(heads) == 15 and all(r[3] == len(r[2]) and r[4] == 30 for r in heads)
    free = next(r for r in heads if r[2] == "Free Whitening Today")
    assert free[5] == "Confirm with the client" and free[6]
    assert next(r for r in heads if r[2] == "Dental Implants Austin")[5] == "Ready"
    assert {r[0] for r in rows} == {"Headlines", "Descriptions", "Display path 1", "Display path 2"}
    assert next(r for r in rows if r[0] == "Display path 1")[2:5] == ["Dental-Implants", 15, 15]
    pmax = rows_of("Performance Max", rec)
    assert {r[0] for r in pmax} == {"Headlines", "Long headlines", "Descriptions", "Short description", "Button"}
    assert next(r for r in pmax if r[0] == "Button")[2] == "Book now"


def test_the_briefs_and_what_to_confirm(rec):
    briefs = rows_of("Creative briefs", rec)
    assert {r[0] for r in briefs} == {"Meta feed"} and {"Concept", "Hook", "Visual direction", "Beat 1", "Ask the client for", "Do", "Do not"} <= {r[1] for r in briefs}
    assert any(r[1].startswith("Size: Image") and "1080 x 1080" in r[2] for r in briefs)
    confirm = rows_of("To confirm", rec)
    assert [(r[0], r[2]) for r in confirm] == [("Google Search", "Free Whitening Today")] and confirm[0][3] and confirm[0][4] == ""


def test_the_checks_and_the_facts(rec):
    checks = dict(map(tuple, ex.tables(rec)["Checks"][1]))
    assert checks["Lines ready"] > 20 and checks["Lines to confirm"] == 1 and "Keywords in a Google headline" in checks
    facts = ex.tables(rec)["Facts used"][1]
    assert ["Confirmed by the person", "Rated 4.9 stars from 312 reviews."] in facts
    assert any(r[0] == "Rating on the website" and "4.9" in r[1] for r in facts)


def test_the_excel_file_opens_has_a_sheet_per_table_and_never_runs_a_formula(rec):
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(rec)))
    assert wb.sheetnames == list(ex.tables(rec))
    ws = wb["Google Search"]
    assert [c.value for c in ws[1]] == ["Field", "Angle", "Text", "Characters", "Limit", "Status", "What to confirm", "Notes"] and ws.freeze_panes == "A2"
    assert ws.max_row == 1 + 15 + 4 + 2
    evil = copy.deepcopy(rec)
    evil["plan"]["formats"]["google_rsa"]["fields"][0]["items"][0]["text"] = '=HYPERLINK("http://x","click")'
    evil["name"] = "=cmd"
    cell = next(c for row in openpyxl.load_workbook(io.BytesIO(ex.xlsx(evil)))["Google Search"].iter_rows(min_row=2) for c in row if "HYPERLINK" in str(c.value))
    assert cell.data_type == "s"
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(evil)))
    body = z.read("google-search.csv").decode()
    assert "'=HYPERLINK" in body or '"=HYPERLINK' not in body.replace("'=HYPERLINK", "")


def test_the_csv_zip_has_a_file_for_every_table(rec):
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(rec)))
    assert z.namelist() == ["overview.csv", "google-search.csv", "performance-max.csv", "meta-feed.csv", "creative-briefs.csv", "to-confirm.csv", "facts-used.csv", "checks.csv"]
    raw = z.read("google-search.csv").decode()
    assert "\r\n" in raw and len(list(csv.reader(io.StringIO(raw)))) == 1 + 15 + 4 + 2


def test_the_editor_file_holds_the_ready_lines_of_the_search_ad_paused(rec):
    got = list(csv.reader(io.StringIO(ex.csv_editor(rec))))
    assert len(got) == 2 and got[0][:7] == ["Campaign", "Ad group", "Ad type", "Status", "Final URL", "Path 1", "Path 2"]
    row = dict(zip(got[0], got[1]))
    assert row["Campaign"] == "P2_Search_Implants_20261101" and row["Ad group"] == "Implants" and row["Status"] == "Paused" and row["Ad type"] == "Responsive search ad"
    assert row["Final URL"] == "https://acmedental.com/" and (row["Path 1"], row["Path 2"]) == ("Dental-Implants", "Austin")
    heads = [row["Headline %d" % n] for n in range(1, 16) if row["Headline %d" % n]]
    assert "Free Whitening Today" not in heads and len(heads) == 14 and heads[0] == "Dental Implants Austin"
    assert [row["Description %d" % n] for n in range(1, 5)] == A.DESCS[:4]


def test_the_editor_file_defaults_the_names_and_is_withheld_when_too_little_is_ready():
    plain = build({"formats": ["google_rsa"]})
    row = dict(zip(*csv.reader(io.StringIO(ex.csv_editor(plain)))))
    assert row["Campaign"] == "New campaign" and row["Ad group"] == "New ad group"
    thin = copy.deepcopy(plain)
    for i in thin["plan"]["formats"]["google_rsa"]["fields"][0]["items"][2:]:
        i["status"] = "confirm"
    assert ex.csv_editor(thin) is None
    assert ex.csv_editor(build({"formats": ["meta_feed"]})) is None


def test_the_editor_file_is_the_search_ad_wherever_it_comes_in_the_order():
    both = build({"formats": ["google_pmax", "meta_feed", "google_rsa"]})
    got = list(csv.reader(io.StringIO(ex.csv_editor(both))))
    row = dict(zip(got[0], got[1]))
    assert row["Headline 1"] == "Dental Implants Austin" and row["Ad type"] == "Responsive search ad"
    assert row["Headline 15"] and row["Description 1"] == A.DESCS[0]                           # the search ad's own fields, not the Performance Max group's


def test_a_format_with_nothing_kept_has_no_sheet():
    got = build({"formats": ["google_rsa", "meta_feed"]}, A.Model({"google_rsa": {}}))
    assert "Google Search" not in ex.tables(got) and "Meta feed" in ex.tables(got)
    assert any("Google Search" in r[1] for r in ex.tables(got)["Checks"][1] if r[0] == "Short")


def test_the_pdf_is_a_pdf_that_says_what_it_holds(rec):
    import pdfplumber
    body = ex.pdf(rec)
    assert body[:4] == b"%PDF"
    with pdfplumber.open(io.BytesIO(body)) as pdf:
        text = " ".join(p.extract_text() or "" for p in pdf.pages)
    assert "Ad Copy and Creative Briefs" in text and "Acme Dental" in text and "Dental Implants" in text and "Confirm with" in text and "Free Whitening" in text      # table cells wrap, so whole lines are not searched for


def test_the_stem_and_an_empty_result():
    assert ex.file_stem({"name": "Acme Dental & Co."}) == "Acme-Dental-Co-ad-copy" and ex.file_stem({}) == "client-ad-copy"
    empty = {"id": "x", "name": "Acme", "inputs": A.clean(), "plan": {}}
    assert list(ex.tables(empty)) == ["Overview", "Checks"] and ex.csv_editor(empty) is None
    assert ex.xlsx(empty)[:2] == b"PK" and ex.pdf(empty)[:4] == b"%PDF"
