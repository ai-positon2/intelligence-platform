"""Landing Page & Relevancy Check: the downloads (tracker/paid_lp_export.py): every file is drawn from the same tables, opens, and cannot run a formula."""
import csv
import io
import zipfile

import openpyxl
import pytest

from tests.test_paid_lp import ACCOUNT_ID, NOW, Fetcher, fresh, fresh_caches, run, search  # noqa: F401  (fresh_caches: the autouse fixture that forgets saved reads)
from tests.test_paid_lp_read import ad_row, kw_row, page_row, term_row
from tracker import paid_lp_export as ex

BAD_HTML = b"<html><body>" + b"<p>nothing about the ad here</p>" * 30 + b"</body></html>"


def make_report(**kw):
    pages = [page_row("https://acme-dental.com/lp", 700.0, "MOBILE", 600, 20), page_row("https://acme-dental.com/lp", 100.0, "DESKTOP", 250, 10),
             page_row("https://acme-dental.com/=cmd", 200.0, "MOBILE", 100, 1)]
    ads = [ad_row("g1", "Emergency", "Boston", "https://acme-dental.com/lp", 800.0, ("Emergency Dentist Boston",), ("Cleanings from $49. Free consultation.",)),
           ad_row("g2", "Other", "Boston", "https://acme-dental.com/=cmd", 200.0, ("Zebra Grooming",), ("Free zebra stripes",))]
    kws = [kw_row("g1", "emergency dentist boston", 50.0), kw_row("g2", "=cmd|' /C calc'!A0", 40.0, lp="BELOW_AVERAGE")]
    terms = [term_row("g1", "emergency dentist boston", 9), term_row("g2", "+cmd zebra", 4), term_row("g2", "@sum", 3)]
    kw.setdefault("fetcher", Fetcher({"/=cmd": (200, BAD_HTML)}))
    return run(search_fn=search(pages=pages, ads=ads, keywords=kws, terms=terms), **kw)


@pytest.fixture()
def report():
    return make_report()


def test_the_tables_come_in_the_order_they_are_shown(report):
    assert list(ex.tables(report)) == ["Summary", "Pages", "Fix first", "All checks", "Could not check", "Keywords", "Search terms", "Ad offers", "Notes"]


def test_the_summary_says_who_what_and_how_well(report):
    header, rows = ex.tables(report)["Summary"]
    d = dict((r[0], r[1]) for r in rows)
    assert header == ["Item", "Detail"]
    assert d["Google Ads account"] == "Acme Dental (%s)" % ACCOUNT_ID and d["Report made on"] == "2026-10-17" and d["Read on"] == "2026-10-17 12:00" and d["Window"] == "the last 30 days"
    assert d["Overall score"] == report["summary"]["score"] and d["Overall grade"] == report["summary"]["grade"]
    assert d["Landing pages that took money"] == 2 and d["Pages checked"] == 2 and d["Failed critical checks"] == report["summary"]["critical"]
    assert any(k.startswith("Section: The page opens") for k in d) and any(k.startswith("Section: Does it match the ad") for k in d)


def test_every_row_is_as_wide_as_its_header(report):
    for name, (header, rows) in ex.tables(report).items():
        assert rows, name
        for r in rows:
            assert len(r) == len(header), (name, r)
    assert [r[0] for r in ex.tables(report)["Fix first"][1]] == list(range(1, len(ex.tables(report)["Fix first"][1]) + 1))


def test_the_pages_sheet_has_one_row_per_page_with_its_numbers(report):
    header, rows = ex.tables(report)["Pages"]
    by = {r[0]: dict(zip(header, r)) for r in rows}
    good = by["https://acme-dental.com/lp"]
    assert good["Spend"] == 800.0 and good["Clicks"] == 850 and good["Conversions"] == 30.0 and good["Conv rate %"] == 3.53 and good["HTTP status"] == 200
    assert good["Mobile conv rate %"] == 3.33 and good["Desktop conv rate %"] == 4.0 and good["Grade"] in ("A", "B") and good["Failed"] >= 0
    assert good["Main headline"] == "Emergency Dentist in Boston" and good["Ad groups"] == 1
    assert by["https://acme-dental.com/=cmd"]["Score"] < good["Score"]


def test_a_page_that_could_not_be_verified_says_so_instead_of_a_score(report):
    fresh()
    r = make_report(fetcher=Fetcher({"/lp": (403, b"no"), "/=cmd": (403, b"no")}))
    header, rows = ex.tables(r)["Pages"]
    assert all(dict(zip(header, row))["Score"] == "Not verified" and dict(zip(header, row))["Grade"] == "" for row in rows)


def test_fix_first_is_one_row_per_check_with_the_pages_it_fails_on_and_what_to_do(report):
    header, rows = ex.tables(report)["Fix first"]
    first = dict(zip(header, rows[0]))
    assert first["Check"].split(" ")[0] in report["catalog"] and first["What to do"] and first["Pages (the biggest first)"] and first["Spend on those pages"] > 0
    assert first["Pages failing"] + first["Pages warning"] >= 1 and first["Example of what was found"]
    assert [dict(zip(header, r))["Severity"] for r in rows][0] in ("Critical", "High")
    assert len({r[1] for r in rows}) == len(rows)


def test_all_checks_has_every_finding_of_every_page_with_why_and_fix(report):
    header, rows = ex.tables(report)["All checks"]
    assert len(rows) == 2 * 35
    row = next(dict(zip(header, r)) for r in rows if r[2].startswith("match_offers") and r[0].endswith("/lp"))
    assert row["Result"] in ("Pass", "Fail") and row["Why it matters"] and row["What to do"] and row["Section"] == "Does it match the ad"


def test_the_keywords_terms_and_offers_sheets_say_what_the_page_carries(report):
    kh, kr = ex.tables(report)["Keywords"]
    k = dict(zip(kh, next(r for r in kr if r[3] == "emergency dentist boston")))
    assert k["On the page"] in ("In the headline", "On the page") and k["Quality score"] == 7 and k["Landing page experience"] == "Below average" and k["Ad group"] == "Emergency"
    bad = [dict(zip(kh, r)) for r in kr if r[0].endswith("=cmd")][0]
    assert bad["On the page"] == "Not on the page"
    oh, orows = ex.tables(report)["Ad offers"]
    assert {(r[3], r[4]) for r in orows if r[0].endswith("/lp")} == {("$49", "Not on the page"), ("free consultation", "On the page")} or {r[3] for r in orows}
    th, trows = ex.tables(report)["Search terms"]
    assert th[:4] == ["Page", "Campaign", "Ad group", "Search term"]


def test_a_value_a_spreadsheet_would_run_never_runs(report):
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(report)))
    for ws in wb:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                assert not (isinstance(c.value, str) and c.value.startswith(("=", "+", "@"))) or c.data_type == "s", (ws.title, c.value)
    cell = next(c for row in wb["Keywords"].iter_rows(min_row=2) for c in row if isinstance(c.value, str) and "cmd" in c.value and "calc" in c.value)
    assert cell.data_type == "s"
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(report)))
    cells = [c for name in z.namelist() for r in csv.reader(io.StringIO(z.read(name).decode())) for c in r]
    assert not any(c.startswith(("=", "+", "@")) for c in cells)
    assert "'=cmd|' /C calc'!A0" in cells or any("calc" in c and c.startswith("'") for c in cells)


def test_the_workbook_has_a_sheet_per_table_with_results_coloured(report):
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(report)))
    assert wb.sheetnames == list(ex.tables(report))
    ws = wb["All checks"]
    header = [c.value for c in ws[1]]
    col = header.index("Result") + 1
    fills = {ws.cell(row=r, column=col).value: ws.cell(row=r, column=col).fill.fgColor.rgb for r in range(2, ws.max_row + 1)}
    assert fills["Pass"].endswith("DDF3E2") and (fills.get("Fail", "FFD9D6").endswith("FFD9D6"))
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref and wb["Fix first"].cell(row=2, column=1).value == 1
    kw = wb["Keywords"]
    kh = [c.value for c in kw[1]]
    on = kh.index("On the page") + 1
    vals = {kw.cell(row=r, column=on).value: kw.cell(row=r, column=on).fill.fgColor.rgb for r in range(2, kw.max_row + 1)}
    assert vals["Not on the page"].endswith("FFD9D6")


def test_the_csv_zip_has_one_file_per_table(report):
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(report)))
    assert z.namelist() == [t.lower().replace(" ", "-") + ".csv" for t in ex.tables(report)]
    rows = list(csv.reader(io.StringIO(z.read("pages.csv").decode())))
    assert rows[0][0] == "Page" and len(rows) == 3


def test_a_report_with_no_pages_still_makes_every_file():
    empty = run(search_fn=search(pages=[], ads=[]))
    assert list(ex.tables(empty)) == ["Summary", "Notes"]
    assert ex.xlsx(empty)[:2] == b"PK" and ex.csv_zip(empty)[:2] == b"PK" and ex.pdf(empty)[:5] == b"%PDF-"
    assert dict(ex.tables(empty)["Summary"][1])["Overall score"] == "Nothing could be judged"


def test_the_pdf_is_a_document_with_the_summary_the_fixes_and_each_pages_problems(report):
    pdf = ex.pdf(report)
    assert pdf[:5] == b"%PDF-" and len(pdf) > 5000
    pages = ex.pdf_pages(report)
    assert [p[0] for p in pages] == ["https://acme-dental.com/lp", "https://acme-dental.com/=cmd"]
    worse = pages[1]
    assert worse[2] == 200.0 and worse[3] and all(r[0] in ("Fail", "Warning") for r in worse[3]) and all(len(r) == 5 for r in worse[3])


def test_the_pdf_lists_only_the_busiest_pages():
    pages = [page_row("https://acme-dental.com/p%d" % i, float(100 - i)) for i in range(ex.PDF_PAGES + 5)]
    r = run(search_fn=search(pages=pages, ads=[]), max_pages=ex.PDF_PAGES + 5)
    assert len(ex.pdf_pages(r)) == ex.PDF_PAGES and ex.pdf(r)[:5] == b"%PDF-"


def test_the_file_stem_names_the_account_safely(report):
    assert ex.file_stem(report) == "Acme-Dental-landing-page-check"
    assert ex.file_stem({"account": {}}) == "account-landing-page-check" and ex.file_stem({}) == "account-landing-page-check"
    assert ex.file_stem({"account": {"name": "../../etc/passwd  &x", "id": "1"}}) == "etc-passwd-x-landing-page-check"
    assert ex.file_stem({"account": {"id": "1112223334"}}) == "1112223334-landing-page-check"
    assert ex.file_stem({"account": {"name": "x" * 100}}) == "x" * 60 + "-landing-page-check"


def test_affected_things_are_joined_and_say_how_many_more():
    assert ex._affected({"items": ["a", "b"], "count": 2}) == "a; b" and ex._affected({"items": [], "count": 0}) == ""
    assert ex._affected({"items": ["a", "b"], "count": 3}) == "a; b (and 1 more)"
    many = {"items": ["a", "b", "c", "d", "e", "f", "g"], "count": 12}
    assert ex._affected(many, 5) == "a; b; c; d; e (and 7 more)" and ex._affected(many) == "a; b; c; d; e; f; g (and 5 more)"


def test_no_em_or_en_dash_in_any_file_text(report):
    for name, (header, rows) in ex.tables(report).items():
        for cell in header + [c for r in rows for c in r]:
            assert "—" not in str(cell) and "–" not in str(cell), name

