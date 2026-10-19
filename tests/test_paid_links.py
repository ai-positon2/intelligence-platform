"""Paid Marketing Agents: Unknown Domains Alert, the URL Checker.

Pinned here: which addresses are checked (placeholders taken out, templates and Facebook left alone); how an answer is judged (404, 410, 5xx, no such domain,
a bad certificate, no answer, a redirect loop are broken; a site that refuses automated checks is blocked, never broken); a broken answer is asked for again
and must fail both times; what counts as a redirect (another page, another domain) and what does not (http to https, www, a trailing slash, case, a query
string); who is checked when (never checked, broken, or a day old); per-host politeness and the time budget; findings (new, back, fixed, dismissed); the
sheet and the alert words; and the whole thing inside a scan: one email with the sheet attached, once per finding, nothing before the first check is
over, and a check that fails never costs the domain scan. No network.
"""
import io
import os
import sys
import threading
import time

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from domains_fakes import Deliver, MemStore, Site  # noqa: E402
from tracker import paid_domains as dom  # noqa: E402
from tracker import paid_domains_scan as sc  # noqa: E402
from tracker import paid_links as pl  # noqa: E402
from tracker import paid_links_scan as pls  # noqa: E402

A = {"id": "1112223330", "name": "Acme Garden", "platform": "google", "group": "mcc-1"}
B = {"id": "4445556660", "name": "Bolt Dental", "platform": "google", "group": "mcc-2"}
LINK = "https://p.example/p2/paid-marketing/unknown-domains"
NOW = 1_800_000_000.0
NOSLEEP = lambda s: None  # noqa: E731


def sight(account, url, where="ad", oid="1", redirect=False):
    return {"platform": account["platform"], "accountId": account["id"], "accountName": account["name"], "where": where, "objectId": oid,
            "objectName": "Ad " + oid, "campaign": "Camp", "adGroup": "Group", "url": url, "host": dom.host_of(url), "redirect": redirect}


# ── Which addresses ─────────────────────────────────────────────────────────

def test_which_addresses_are_checked():
    assert pl.clean_url("https://acme-garden.com/a?c={campaignid}&x={_custom}") == "https://acme-garden.com/a?c=&x="
    assert pl.clean_url("{lpurl}?x=1") is None and pl.clean_url("tel:+1555") is None and pl.clean_url("acme-garden.com") is None
    assert pl.clean_url("https://www.facebook.com/acme") is None and pl.clean_url("https://apps.apple.com/app/x") is None
    s = [sight(A, "https://acme-garden.com/", oid="1"), sight(A, "https://acme-garden.com/", oid="1"), sight(A, "https://acme-garden.com/", oid="2"),
         sight(A, "https://track.example/r?u={lpurl}", where="tracking template (ad)", redirect=True), sight(A, "https://facebook.com/x", oid="3")]
    occ = pl.occurrences(s)
    assert [o["objectId"] for o in occ] == ["1", "2"]                              # the same ad twice is one place; a template is not a final URL


# ── Judging an answer ───────────────────────────────────────────────────────

@pytest.mark.parametrize("answer, verdict, text", [
    (200, "ok", ""), (301, "ok", ""), (404, "broken", "Page not found (404)"), (410, "broken", "Page not found (410)"),
    (500, "broken", "Server error (500)"), (503, "broken", "Server error (503)"),
    (403, "blocked", "could not be verified"), (429, "blocked", "429"), (401, "blocked", "401"),
    (OSError("[Errno 11001] getaddrinfo failed"), "broken", "does not resolve"),
    (Exception("SSLError certificate verify failed"), "broken", "certificate or SSL"),
    (TimeoutError("timed out"), "broken", "No answer within"),
    (ConnectionRefusedError("refused"), "broken", "Could not connect"),
    (ValueError("Too many redirects."), "broken", "Redirect loop"),
    (ValueError("Only standard web ports are allowed."), "unchecked", "Not checked"),
    (ValueError("Private or reserved page destinations are not allowed."), "unchecked", "Not checked"),
])
def test_how_an_answer_is_judged(answer, verdict, text):
    r = pl.check_url("https://x.example/", Site({"https://x.example/": answer}), NOSLEEP)
    assert r["verdict"] == verdict and text in r["problem"]


def test_a_broken_answer_is_asked_for_again_and_the_second_counts():
    asks = []

    def flaky(url, **kw):
        asks.append(url)
        return Site({"https://x.example/": 503 if len(asks) == 1 else 200})(url, **kw)
    r = pl.check_url("https://x.example/", flaky, lambda s: asks.append("slept"))
    assert r["verdict"] == "ok" and asks == ["https://x.example/", "slept", "https://x.example/"]
    never = Site({"https://x.example/": 404})
    pl.check_url("https://x.example/", never, NOSLEEP)
    assert len(never.asked) == 2
    fine = Site()
    pl.check_url("https://x.example/", fine, NOSLEEP)
    assert len(fine.asked) == 1


def test_the_response_is_closed_without_reading_the_page():
    closed = []

    class R:
        status_code, url = 200, "https://x.example/"

        def close(self):
            closed.append(1)
    pl.check_url("https://x.example/", lambda url, **kw: R(), NOSLEEP)
    assert closed == [1]


def test_confirming_keeps_an_address_broken_only_if_it_fails_twice():
    first = {"https://a.example/": {"verdict": "broken", "problem": "Page not found (404)", "code": 404, "finalUrl": ""},
             "https://b.example/": {"verdict": "broken", "problem": "Page not found (404)", "code": 404, "finalUrl": ""},
             "https://c.example/": {"verdict": "ok", "problem": "", "code": 200, "finalUrl": "https://c.example/"}}
    site = Site({"https://a.example/": 404, "https://b.example/": 200})
    out = pl.confirm(first, site, sleep=NOSLEEP)
    assert out["https://a.example/"]["verdict"] == "broken" and out["https://b.example/"]["verdict"] == "ok" and out["https://c.example/"]["verdict"] == "ok"
    assert "https://c.example/" not in {u.split("?")[0] for u in site.asked}


# ── Redirects ───────────────────────────────────────────────────────────────

@pytest.mark.parametrize("url, final", [
    ("http://acme-garden.com/", "https://acme-garden.com/"), ("https://acme-garden.com/", "https://www.acme-garden.com/"),
    ("https://www.acme-garden.com/a", "https://acme-garden.com/a/"), ("https://acme-garden.com/A", "https://acme-garden.com/a"),
    ("https://acme-garden.com/a?x=1", "https://acme-garden.com/a"), ("https://acme-garden.com/a", ""),
])
def test_a_change_that_only_normalises_the_same_page_is_no_redirect(url, final):
    assert pl.redirect_of(url, final) is None


def test_a_different_page_or_domain_is_a_redirect():
    r = pl.redirect_of("https://acme-garden.com/spring-sale", "https://acme-garden.com/")
    assert r == {"to": "https://acme-garden.com/", "crossDomain": False}
    r = pl.redirect_of("https://acme-garden.com/", "https://somewhere-else.example/landing")
    assert r["crossDomain"] is True and pl.redirect_problem(r) == "Redirects to https://somewhere-else.example/landing (another domain)"
    assert pl.redirect_of("https://shop.acme-garden.com/", "https://acme-garden.com/shop")["crossDomain"] is False      # a sibling subdomain is the same domain


def test_the_guarded_fetch_gives_the_final_address_without_its_query():
    r = pl.check_url("https://acme-garden.com/old?x=1", Site({"https://acme-garden.com/old": "https://acme-garden.com/new?utm=1#top"}), NOSLEEP)
    assert r["verdict"] == "ok" and r["finalUrl"] == "https://acme-garden.com/new"


# ── Who is checked, how politely ────────────────────────────────────────────

def test_who_is_due():
    cache = {"https://fresh.example/": {"verdict": "ok", "checkedAt": NOW - 3600}, "https://old.example/": {"verdict": "ok", "checkedAt": NOW - 21 * 3600},
             "https://broken.example/": {"verdict": "broken", "checkedAt": NOW - 60}, "https://blocked.example/": {"verdict": "blocked", "checkedAt": NOW - 3600}}
    urls = list(cache) + ["https://new.example/"]
    assert pl.due(urls, cache, NOW) == ["https://old.example/", "https://broken.example/", "https://new.example/"]


def test_at_most_two_at_a_time_to_one_host():
    live, peak, lock = [0], [0], threading.Lock()

    def slow(url, **kw):
        with lock:
            live[0] += 1
            peak[0] = max(peak[0], live[0])
        time.sleep(0.02)
        with lock:
            live[0] -= 1
        return Site()(url, **kw)
    out = pl.check_many(["https://one.example/p%d" % i for i in range(12)], slow, workers=8, sleep=NOSLEEP)
    assert len(out) == 12 and peak[0] <= 2


def test_the_time_budget_leaves_the_rest_for_the_next_scan():
    ticks = iter(range(100))
    out = pl.check_many(["https://x%d.example/" % i for i in range(5)], Site(), workers=1, sleep=NOSLEEP, budget_s=2, clock=lambda: next(ticks))
    assert 0 < len(out) < 5


# ── Findings ────────────────────────────────────────────────────────────────

def answer(verdict="ok", problem="", code=200, final=""):
    return {"verdict": verdict, "problem": problem, "code": code, "finalUrl": final}


def test_findings_are_new_back_fixed_or_dismissed():
    occs = pl.occurrences([sight(A, "https://x.example/a", oid="1"), sight(A, "https://x.example/b", oid="2"), sight(B, "https://y.example/", oid="3")])
    cache = {"https://x.example/a": answer("broken", "Page not found (404)", 404), "https://x.example/b": answer(final="https://x.example/"),
             "https://y.example/": answer()}
    writes, c = pl.apply(occs, cache, {}, {A["id"], B["id"]}, "T1")
    assert c == {"new": 2, "reopened": 0, "fixed": 0}
    by = {w["objectId"]: w for w in writes}
    assert by["1"]["kind"] == "broken" and by["1"]["problem"] == "Page not found (404)" and by["1"]["status"] == "open" and by["1"]["alertedAt"] is None
    assert by["2"]["kind"] == "redirect" and by["2"]["target"] == "https://x.example/" and "Redirects to" in by["2"]["problem"]
    existing = {w["fingerprint"]: dict(w, alertedAt="T1") for w in writes}
    writes, c = pl.apply(occs, cache, existing, {A["id"], B["id"]}, "T2")
    assert c["new"] == 0 and all(w["alertedAt"] == "T1" and w["firstSeen"] == "T1" and w["lastSeen"] == "T2" for w in writes)
    cache["https://x.example/a"] = answer()
    cache["https://x.example/b"] = answer()
    writes, c = pl.apply(occs, cache, existing, {A["id"], B["id"]}, "T3")
    assert c["fixed"] == 2 and {w["status"] for w in writes} == {"fixed"}
    existing = {w["fingerprint"]: w for w in writes}
    cache["https://x.example/a"] = answer("broken", "Server error (500)", 500)
    writes, c = pl.apply(occs, cache, existing, {A["id"], B["id"]}, "T4")
    assert c == {"new": 0, "reopened": 1, "fixed": 0} and writes[0]["alertedAt"] is None and writes[0]["problem"] == "Server error (500)"
    dismissed = {w["fingerprint"]: dict(w, status="dismissed") for w in writes}
    writes, c = pl.apply(occs, cache, dismissed, {A["id"], B["id"]}, "T5")
    assert writes[0]["status"] == "dismissed" and c["new"] == 0 and c["reopened"] == 0


def test_an_unreadable_account_does_not_make_its_findings_fixed():
    occs = pl.occurrences([sight(A, "https://x.example/a")])
    writes, _ = pl.apply(occs, {"https://x.example/a": answer("broken", "Page not found (404)", 404)}, {}, {A["id"]}, "T1")
    existing = {w["fingerprint"]: w for w in writes}
    out, c = pl.apply([], {}, existing, set(), "T2")
    assert out == [] and c["fixed"] == 0
    out, c = pl.apply([], {}, existing, {A["id"]}, "T2")
    assert c["fixed"] == 1                                                         # read fine, and the ad is gone


def test_a_broken_link_that_becomes_a_redirect_is_a_new_finding():
    occs = pl.occurrences([sight(A, "https://x.example/a")])
    writes, _ = pl.apply(occs, {"https://x.example/a": answer("broken", "Page not found (404)", 404)}, {}, {A["id"]}, "T1")
    existing = {w["fingerprint"]: w for w in writes}
    out, c = pl.apply(occs, {"https://x.example/a": answer(final="https://x.example/")}, existing, {A["id"]}, "T2")
    assert c == {"new": 1, "reopened": 0, "fixed": 1} and {w["kind"] for w in out} == {"broken", "redirect"}


# ── The sheet and the alert ─────────────────────────────────────────────────

def rows_for(*specs):
    occs = pl.occurrences([sight(A, u, oid=str(i)) for i, (u, _) in enumerate(specs)])
    cache = {u: a for u, a in specs}
    writes, _ = pl.apply(occs, cache, {}, {A["id"]}, "2026-10-19T00:00:00+00:00")
    return writes


def test_the_sheet_lists_every_link_to_fix_broken_first():
    from openpyxl import load_workbook
    rows = rows_for(("https://a.example/old", answer(final="https://other.example/")), ("https://b.example/gone", answer("broken", "Page not found (404)", 404)))
    ws = load_workbook(io.BytesIO(pl.sheet(rows, [rows[1]["fingerprint"]]))).active
    got = [[c.value for c in r] for r in ws.iter_rows()]
    assert got[0][:4] == ["Account", "Platform", "Campaign", "Ad group"] and got[0][-1] == "Fixed or paused? (fill in)"
    assert got[1][7] == "Broken" and got[1][6] == "https://b.example/gone" and got[1][12] == "Yes"
    assert got[2][7] == "Redirect to another domain" and got[2][9] == "https://other.example/" and got[2][12] in (None, "")


def test_the_alert_words_and_attachment():
    rows = rows_for(("https://a.example/old", answer(final="https://a.example/")), ("https://b.example/gone", answer("broken", "Page not found (404)", 404)),
                    ("https://c.example/gone", answer("broken", "Server error (500)", 500)))
    m = pl.digest(rows, rows[:2], LINK)
    assert m["subject"] == "URL Checker: 1 broken link and 1 redirect in 1 ad account: fix or pause the affected ads"
    assert "3 to fix in all" in m["text"] and "Page not found (404)" in m["text"] and LINK in m["text"]
    assert m["attachments"][0][0] == "broken-links.xlsx" and m["attachments"][0][1][:2] == b"PK"
    assert m["slack_blocks"][0]["type"] == "header" and "The sheet was emailed" in m["slack_blocks"][-1]["text"]["text"]
    assert "—" not in m["text"] + m["html"]
    many = rows_for(*[("https://x%d.example/" % i, answer("broken", "Page not found (404)", 404)) for i in range(20)])
    assert "...and 5 more in the sheet." in pl.digest(many, many, LINK)["text"] and len(pl.digest(many, many, LINK)["slack_blocks"]) <= 12


# ── Inside a scan ───────────────────────────────────────────────────────────

class World:
    def __init__(self, **ads):
        self.ads, self.accounts = dict(ads), [A, B]

    def inventory(self, store):
        return list(self.accounts), [], []

    def read(self, account, store):
        out = [sight(account, u, oid=str(i)) for i, u in enumerate(self.ads.get(account["id"], []))]
        return out, {"ads": len(out), "assetGroups": 0, "assets": 0}


def scan(store, world, deliver, site, hours=0):
    """One scan, `hours` after the start of the test's day."""
    from datetime import datetime, timedelta, timezone
    when = datetime(2026, 10, 19, tzinfo=timezone.utc) + timedelta(hours=hours)
    return sc.run_scan(store, trigger="test", link=LINK, deliver=deliver, inventory_fn=world.inventory, read_fn=world.read, workers=2, fetch_fn=site, sleep_fn=NOSLEEP,
                       now_fn=lambda: when)


def configured():
    s = MemStore()
    s.state_set("settings", {"emails": ["ppc@example.com"], "slackChannel": "#alerts"})
    return s


def test_a_scan_emails_a_sheet_of_broken_links_and_redirects_once():
    s, dv = configured(), Deliver()
    w = World(**{A["id"]: ["https://x.example/gone", "https://x.example/moved", "https://x.example/fine"], B["id"]: ["https://y.example/"]})
    site = Site({"https://x.example/gone": 404, "https://x.example/moved": "https://other.example/"})
    run = scan(s, w, dv, site)
    assert run["links"]["urls"] == 4 and run["links"]["broken"] == 1 and run["links"]["redirects"] == 1 and run["links"]["alerted"] == 2
    links = [e for e in dv.emails if e["subject"].startswith("URL Checker")]
    assert len(links) == 1 and links[0]["to"] == ["ppc@example.com"] and links[0]["attachments"][0][0] == "broken-links.xlsx"
    assert [m for m in dv.slacks if "URL Checker" in m["text"]]
    again = scan(s, w, dv, site)
    assert again["links"]["new"] == 0 and again["links"]["alerted"] == 0 and len([e for e in dv.emails if e["subject"].startswith("URL Checker")]) == 1


def test_an_address_is_not_asked_again_inside_a_day_unless_it_is_broken():
    s, dv = configured(), Deliver()
    w = World(**{A["id"]: ["https://x.example/gone", "https://x.example/fine"]})
    site = Site({"https://x.example/gone": 404})
    scan(s, w, dv, site)
    site.asked.clear()
    scan(s, w, dv, site)
    assert {u for u in site.asked} == {"https://x.example/gone"}


def test_a_fixed_link_is_marked_fixed_and_a_new_one_alerts_again():
    s, dv = configured(), Deliver()
    w = World(**{A["id"]: ["https://x.example/gone"]})
    site = Site({"https://x.example/gone": 404})
    scan(s, w, dv, site)
    site.answers["https://x.example/gone"] = 200
    run = scan(s, w, dv, site, hours=4)                                             # it was broken, so it is asked again at once
    assert run["links"]["fixed"] == 1 and run["links"]["broken"] == 0 and [f["status"] for f in s.links.values()] == ["fixed"]
    site.answers["https://x.example/gone"] = 404
    assert scan(s, w, dv, site, hours=8)["links"]["broken"] == 0                    # it worked a few hours ago: not asked again until a day has passed
    run = scan(s, w, dv, site, hours=30)
    assert run["links"]["reopened"] == 1 and run["links"]["alerted"] == 1


def test_a_digest_that_went_nowhere_is_tried_again():
    s, dv = MemStore(), Deliver()
    w = World(**{A["id"]: ["https://x.example/gone"]})
    site = Site({"https://x.example/gone": 404})
    run = scan(s, w, dv, site)
    assert run["links"]["alerted"] == 0 and run["links"]["email"]["status"] == "not_configured"
    s.state_set("settings", {"emails": ["ppc@example.com"], "slackChannel": ""})
    assert scan(s, w, dv, site)["links"]["alerted"] == 1


def test_a_check_that_fails_never_costs_the_domain_scan():
    s, dv = configured(), Deliver()
    w = World(**{A["id"]: ["https://x.example/a"]})

    class Down(MemStore):
        def urls_all(self):
            raise RuntimeError("urls table gone")
    d = Down()
    d.state_set("settings", {"emails": [], "slackChannel": ""})
    run = scan(d, w, dv, Site())
    assert run["status"] == "done" and "urls table gone" in run["links"]["error"] and run["new"] == 1


def test_blocked_and_unchecked_addresses_are_counted_not_reported():
    s, dv = configured(), Deliver()
    w = World(**{A["id"]: ["https://x.example/a", "https://x.example/b"]})
    site = Site({"https://x.example/a": 403, "https://x.example/b": ValueError("Only standard web ports are allowed.")})
    run = scan(s, w, dv, site)
    assert run["links"]["blocked"] == 1 and run["links"]["unchecked"] == 1 and run["links"]["broken"] == 0 and not [e for e in dv.emails if e["subject"].startswith("URL")]


def test_the_page_snapshot_and_dismissing():
    s, dv = configured(), Deliver()
    w = World(**{A["id"]: ["https://x.example/gone"]})
    scan(s, w, dv, Site({"https://x.example/gone": 404}))
    snap = sc.snapshot(s, [A], [], [], dv)
    assert snap["linkCounts"] == {"open": 1, "dismissed": 0, "fixed": 0} and snap["links"][0]["problem"] == "Page not found (404)"
    fp = snap["links"][0]["fingerprint"]
    pls.decide(s, fp, "dismissed", "me")
    assert s.links[fp]["status"] == "dismissed"
    pls.decide(s, fp, "open", "me")
    with pytest.raises(sc.BadRequest):
        pls.decide(s, fp, "fixed", "me")
    with pytest.raises(sc.BadRequest):
        pls.decide(s, "nope", "dismissed", "me")
    assert pls.open_sheet(s)[:2] == b"PK"


def test_an_email_can_carry_an_attachment():
    from tracker import market_radar_deliver as d
    msg = d._message("s", "t", "<p>h</p>", ["a@example.com"], "me@example.com", [("broken-links.xlsx", b"PKdata", pl.SHEET_TYPE)])
    parts = list(msg.iter_attachments())
    assert len(parts) == 1 and parts[0].get_filename() == "broken-links.xlsx" and parts[0].get_content() == b"PKdata"
    assert not list(d._message("s", "t", "<p>h</p>", ["a@example.com"], "me@example.com").iter_attachments())
