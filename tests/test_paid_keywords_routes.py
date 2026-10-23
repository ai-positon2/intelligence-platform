"""Keyword Research: the routes. Staff only, same-origin JSON for anything that changes, a daily cap shared with the other two agents of this kind, a build belongs to one
agent, and downloads that open."""
import csv
import io
import json
import os
import re
import sys
import zipfile

import openpyxl
import pytest

os.environ.setdefault("GOOGLE_CLIENT_ID", "test")
os.environ.setdefault("GOOGLE_CLIENT_SECRET", "test")
os.environ.setdefault("FLASK_SECRET_KEY", "test")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import app as appmod  # noqa: E402
from tests import business_fakes as B, keyword_fakes as K, strategy_fakes as F  # noqa: E402
from tracker import paid_gads_segments as gads, paid_keywords as pk, paid_strategy as ps, paid_strategy_store as store  # noqa: E402

STAFF, OTHER, ADMIN = "someone@position2.com", "another@position2.com", "reporting@position2.com"
SAME = {"Origin": "http://localhost"}
API = "/p2/paid-marketing/api/keywords"
BAPI = "/p2/paid-marketing/api/business"
SAPI = "/p2/paid-marketing/api/strategy"
FORM = {"name": "Acme Dental", "website": "acmedental.com", "account": "1234567890", "locations": "Austin, TX", "services": "dental implants, invisalign", "competitors": ["smileaustin.com"]}


class World:
    def __init__(self):
        self.store = F.FakeStrategyStore()
        self.spawned = []


@pytest.fixture
def world(monkeypatch):
    w = World()
    for name in ("create", "get", "get_light", "list_recent", "begin", "set_steps", "set_step", "save_research", "save_plan_part", "save_plan", "finish", "fail", "archive", "touch",
                 "record_run", "runs_today"):
        monkeypatch.setattr(store, name, getattr(w.store, name))
    monkeypatch.setattr(appmod, "_strat_spawn", lambda fn: (w.spawned.append(fn), fn()))
    monkeypatch.setattr(appmod, "_strat_deps", lambda rec, account: K.deps(store=K.Store([]), account=bool(account), email=rec["created_by"]))
    monkeypatch.setattr(gads, "list_accounts", lambda **kw: [{"id": "1234567890", "name": "Acme Dental", "currency": "USD"}])
    monkeypatch.setattr(appmod, "_STRAT_LAST", {})
    monkeypatch.setenv("ANTHROPIC_API_KEY", "x")
    return w


def client_as(email):
    c = appmod.app.test_client()
    if email:
        with c.session_transaction() as sess:
            sess["google_user"] = {"email": email, "name": "T"}
    return c


def post(c, path, body=None, headers=SAME, api=API):
    return c.post(api + path, data=json.dumps(body if body is not None else {}), content_type="application/json", headers=headers)


def built(c, **form):
    r = post(c, "/start", dict(FORM, **form))
    assert r.status_code == 201, r.get_json()
    return r.get_json()["id"]


def delete(c, path):
    return c.delete(API + path, headers=SAME, content_type="application/json")


# ── who may use it ──────────────────────────────────────────────────────────

def test_the_page_and_every_api_are_for_position2_staff_only(world):
    anon, outsider = client_as(None), client_as("x@gmail.com")
    for path in ("/p2/paid-marketing/keyword-research", API + "/options", API + "/list", API + "/" + "a" * 32, API + "/" + "a" * 32 + "/status", API + "/" + "a" * 32 + "/download?file=xlsx"):
        assert anon.get(path).status_code == 302, path
        r = outsider.get(path)
        assert r.status_code == 302 and "/app" in r.headers["Location"], path
    assert anon.post(API + "/start", data="{}", content_type="application/json", headers=SAME).status_code == 302
    assert anon.post(API + "/" + "a" * 32 + "/rebuild", data="{}", content_type="application/json", headers=SAME).status_code == 302
    assert anon.delete(API + "/" + "a" * 32, headers=SAME).status_code == 302


def test_the_page_renders_for_staff(world):
    html = client_as(STAFF).get("/p2/paid-marketing/keyword-research").get_data(as_text=True)
    assert "<title>Keyword Research" in html and "keyword_research.js" in html and "keyword_research.css" in html and 'id="cards"' in html and "pm_picker.js" in html and 'id="fMin"' in html


def test_options_describe_the_cards_the_floors_and_say_whether_the_model_is_there(world, monkeypatch):
    j = client_as(STAFF).get(API + "/options").get_json()
    assert [c["key"] for c in j["cards"]] == ["business", "website", "account", "metaaccount", "competitors"] and j["modelReady"] is True and j["minSearches"] == [10, 50, 100, 500] and j["defaultMin"] == 10
    assert j["goals"]["leads"] and j["buildsPerDay"] == 12
    monkeypatch.delenv("ANTHROPIC_API_KEY")
    assert client_as(STAFF).get(API + "/options").get_json()["modelReady"] is False


def test_everything_that_changes_needs_same_origin_json(world):
    c = client_as(STAFF)
    assert c.post(API + "/start", data=json.dumps(FORM), content_type="application/json", headers={"Origin": "https://evil.example"}).status_code == 403
    assert c.post(API + "/start", data="name=x", content_type="application/x-www-form-urlencoded", headers=SAME).status_code == 415
    sid = built(c)
    assert c.delete(API + "/" + sid, headers={"Origin": "https://evil.example"}).status_code == 403
    assert c.post(API + "/%s/rebuild" % sid, data="{}", headers=SAME).status_code == 415


# ── starting a build ────────────────────────────────────────────────────────

def test_a_build_is_started_run_and_can_be_read_back(world):
    c = client_as(STAFF)
    r = post(c, "/start", FORM)
    assert r.status_code == 201 and r.get_json()["started"] is True
    sid = r.get_json()["id"]
    rec = c.get(API + "/" + sid).get_json()
    assert rec["status"] == "done" and rec["name"] == "Acme Dental" and rec["accountName"] == "Acme Dental" and rec["createdBy"] == STAFF
    assert rec["inputs"]["kind"] == "keywords" and rec["inputs"]["currency"] == "USD" and rec["inputs"]["account"] == "1234567890" and rec["inputs"]["minSearches"] == 10
    assert rec["plan"]["keywords"] and rec["plan"]["themes"] and rec["plan"]["coverage"]["judged"] and rec["plan"]["plan"]["summary"]
    keys = [s["key"] for s in rec["progress"]["steps"]]
    assert keys == ["r:business", "r:website", "r:account", "r:competitors", "r:planner", "a:themes", "a:classify", "a:plan", "f:finish"]
    light = c.get(API + "/%s/status" % sid).get_json()
    assert light["status"] == "done" and "plan" not in light and light["kind"] == "keywords"


def test_the_page_is_not_sent_the_thousands_of_rows_only_the_build_needs(world):
    c = client_as(STAFF)
    rec = c.get(API + "/" + built(c)).get_json()
    assert "ideas" not in rec["research"]["planner"]["data"] and "termMetrics" not in rec["research"]["planner"]["data"] and rec["research"]["planner"]["data"]["total"] == 25
    assert "keywords" not in rec["research"]["account"]["data"] and "negatives" not in rec["research"]["account"]["data"] and rec["research"]["account"]["data"]["keywordCount"] == 5
    assert "text" not in rec["research"]["website"]["data"]["pages"][0]


def test_bad_details_are_refused_with_a_sentence(world):
    c = client_as(STAFF)
    for body, msg in (({"website": "x.com"}, "name"), (dict(FORM, goal="fame"), "goal"), (dict(FORM, competitors=["a.com", "b.com", "c.com", "d.com"]), "at most 3"), (dict(FORM, minSearches=7), "smallest monthly"),
                      (dict(FORM, website="", account="", services=""), "what the client sells"), ("not an object", "form")):
        r = post(c, "/start", body)
        assert r.status_code == 400 and msg in r.get_json()["error"], body
    r = post(c, "/start", dict(FORM, account="1111111111"))
    assert r.status_code == 400 and "not one this platform can read" in r.get_json()["error"]
    assert world.store.rows == {}


def test_a_build_without_an_account_works_and_takes_the_currency_given(world):
    c = client_as(STAFF)
    rec = c.get(API + "/" + built(c, account="", currency="cad")).get_json()
    assert rec["inputs"]["currency"] == "CAD" and rec["status"] == "done" and "account" not in rec["research"]


def test_the_list_shows_this_agents_builds_newest_first_and_not_the_others(world):
    c = client_as(STAFF)
    built(c, name="First Client"), built(c, name="Second Client")
    world.store.create(by=STAFF, inputs=F.clean(name="A Strategy Build"))
    world.store.create(by=STAFF, inputs=B.clean(name="A Business Build"))
    rows = client_as(OTHER).get(API + "/list").get_json()["builds"]
    assert [r["name"] for r in rows] == ["Second Client", "First Client"] and rows[0]["status"] == "done" and rows[0]["createdBy"] == STAFF and "campaigns" not in rows[0]
    assert [r["name"] for r in client_as(OTHER).get(SAPI + "/list").get_json()["builds"]] == ["A Strategy Build"]                     # and each one's list is its own
    assert [r["name"] for r in client_as(OTHER).get(BAPI + "/list").get_json()["builds"]] == ["A Business Build"]


def test_a_build_belongs_to_one_agent_whichever_address_is_used(world):
    c = client_as(STAFF)
    mine = built(c)
    others = [world.store.create(by=STAFF, inputs=F.clean())["id"], world.store.create(by=STAFF, inputs=B.clean())["id"]]
    for theirs in others:
        for path in ("", "/status", "/download?file=xlsx"):
            assert c.get(API + "/" + theirs + path).status_code == 404, path
        assert post(c, "/%s/rebuild" % theirs, {}).status_code == 404 and delete(c, "/" + theirs).status_code == 404
    for other_api in (SAPI, BAPI):
        for path in ("", "/status", "/download?file=xlsx"):
            assert c.get(other_api + "/" + mine + path).status_code == 404, (other_api, path)
        assert c.post(other_api + "/%s/rebuild" % mine, data="{}", headers=SAME, content_type="application/json").status_code == 404
        assert c.delete(other_api + "/" + mine, headers=SAME, content_type="application/json").status_code == 404
    assert world.store.get(mine)["archived"] is False and all(world.store.get(o)["archived"] is False for o in others)


def test_a_person_can_only_start_so_many_builds_a_day_and_the_count_is_shared(world, monkeypatch):
    monkeypatch.setattr(ps, "BUILDS_PER_DAY", 2)
    c = client_as(STAFF)
    built(c)
    world.store.record_run(STAFF)                                                   # a build of another agent counts too: they spend the same credits
    r = post(c, "/start", FORM)
    assert r.status_code == 429 and r.get_json()["kind"] == "cap" and "2 builds in a day" in r.get_json()["error"]
    assert len(client_as(STAFF).get(API + "/list").get_json()["builds"]) == 1       # the refused one is not left in the list
    assert post(client_as(OTHER), "/start", FORM).status_code == 201


def test_only_the_person_who_started_a_build_or_an_admin_can_build_it_again_or_remove_it(world):
    sid = built(client_as(STAFF))
    appmod._STRAT_LAST.clear()
    r = post(client_as(OTHER), "/%s/rebuild" % sid, {"resume": False})
    assert r.status_code == 403 and "Only whoever started this build" in r.get_json()["error"] and world.store.get(sid)["plan"]["keywords"]
    assert delete(client_as(OTHER), "/" + sid).status_code == 403
    assert post(client_as(ADMIN), "/%s/rebuild" % sid, {"resume": True}).status_code == 200
    appmod._STRAT_LAST.clear()
    assert post(client_as(STAFF), "/%s/rebuild" % sid, {}).status_code == 200
    assert delete(client_as(STAFF), "/" + sid).get_json() == {"ok": True}
    assert client_as(STAFF).get(API + "/" + sid).status_code == 404 and client_as(STAFF).get(API + "/list").get_json()["builds"] == []
    other = built(client_as(STAFF))
    assert delete(client_as(ADMIN), "/" + other).status_code == 200


def test_build_again_is_rate_limited_and_resume_keeps_the_work(world, monkeypatch):
    c = client_as(STAFF)
    sid = built(c)
    first = post(c, "/%s/rebuild" % sid, {"resume": False})
    assert first.status_code == 200 and first.get_json() == {"id": sid, "started": True}
    again = post(c, "/%s/rebuild" % sid, {"resume": True})
    assert again.status_code == 429 and again.get_json()["kind"] == "cooldown" and again.get_json()["retryIn"] >= 1
    appmod._STRAT_LAST.clear()
    calls = []
    monkeypatch.setattr(appmod, "_strat_deps", lambda rec, account: K.deps(store=K.Store([]), account=bool(account), call=K.model(calls=calls), email=rec["created_by"]))
    assert post(c, "/%s/rebuild" % sid, {"resume": True}).status_code == 200 and calls == []        # everything was already done
    assert c.get(API + "/" + sid).get_json()["status"] == "done"


def test_a_build_that_is_running_is_not_started_twice(world):
    rec = world.store.create(by=STAFF, inputs=K.clean())
    world.store.begin(rec["id"])
    r = post(client_as(STAFF), "/%s/rebuild" % rec["id"], {})
    assert r.status_code == 409 and "already running" in r.get_json()["error"]


def test_unknown_ids_are_404(world):
    c = client_as(STAFF)
    assert c.get(API + "/" + "0" * 32 + "/status").status_code == 404 and c.get(API + "/not-an-id").status_code == 404 and c.get(API + "/" + "0" * 32).status_code == 404


def test_a_store_that_is_down_is_said_in_words(world, monkeypatch):
    def down(*a, **k):
        raise store.StoreUnavailable("no db")
    for name in ("list_recent", "create", "get", "get_light"):
        monkeypatch.setattr(store, name, down)
    c = client_as(STAFF)
    assert c.get(API + "/list").status_code == 503 and post(c, "/start", FORM).status_code == 503
    assert c.get(API + "/" + "a" * 32).status_code == 503 and c.get(API + "/" + "a" * 32 + "/status").status_code == 503
    assert "store is unavailable" in c.get(API + "/list").get_json()["error"]


# ── the downloads ───────────────────────────────────────────────────────────

def dl(c, sid, file):
    return c.get(API + "/%s/download?file=%s" % (sid, file))


def test_excel_pdf_zip_and_the_editor_file_download_as_files_that_open(world):
    c = client_as(STAFF)
    sid = built(c)
    x = dl(c, sid, "xlsx")
    assert x.status_code == 200 and x.headers["Content-Type"].startswith("application/vnd.openxmlformats") and 'filename="Acme-Dental-keyword-research.xlsx"' in x.headers["Content-Disposition"]
    assert "Keywords" in openpyxl.load_workbook(io.BytesIO(x.data)).sheetnames and x.headers["Cache-Control"] == "private, no-store"
    p = dl(c, sid, "pdf")
    assert p.status_code == 200 and p.headers["Content-Type"] == "application/pdf" and p.data[:5] == b"%PDF-"
    z = dl(c, sid, "zip")
    assert z.status_code == 200 and "keywords.csv" in zipfile.ZipFile(io.BytesIO(z.data)).namelist() and "all-tables.zip" in z.headers["Content-Disposition"]
    q = dl(c, sid, "csv")
    assert q.status_code == 200 and q.data.startswith(chr(0xFEFF).encode()) and 'ads-editor-keywords.csv"' in q.headers["Content-Disposition"]
    assert list(csv.reader(io.StringIO(q.data.decode("utf-8-sig"))))[0] == ["Campaign", "Ad group", "Keyword", "Criterion Type", "Status"]


def test_a_download_that_cannot_be_made_is_said_in_words(world, monkeypatch):
    from tracker import paid_keywords_export as export
    c = client_as(STAFF)
    sid = built(c)
    assert dl(c, sid, "docx").status_code == 400 and "file should be one of" in dl(c, sid, "docx").get_json()["error"]
    monkeypatch.setattr(export, "pdf", lambda rec: (_ for _ in ()).throw(ValueError("boom")))
    r = dl(c, sid, "pdf")
    assert r.status_code == 500 and r.get_json()["error"] == "The file could not be made (ValueError)."
    rec = world.store.rows[sid]
    rec["plan"]["editor"] = []
    r = dl(c, sid, "csv")
    assert r.status_code == 404 and "every keyword that fits is already in the account" in r.get_json()["error"]
    rec["plan"] = {}
    assert dl(c, sid, "xlsx").status_code == 409
    rec["plan"] = {"themeDef": {"themes": []}}                                       # judged, but not yet checked: no checks sheet, so no files
    for kind in ("xlsx", "pdf", "zip", "csv"):
        assert dl(c, sid, kind).status_code == 409, kind


# ── the wiring ──────────────────────────────────────────────────────────────

def test_every_address_the_page_script_uses_is_a_route():
    js = open(os.path.join(ROOT, "static", "js", "keyword_research.js"), encoding="utf-8").read()
    rules = {r.rule for r in appmod.app.url_map.iter_rules()}
    used = set(re.findall(r'API\+"(/[a-z]+)"', js)) | {"/<sid>/status", "/<sid>/rebuild", "/<sid>/download"}
    for path in used:
        assert any(r.startswith("/p2/paid-marketing/api/keywords") and r.endswith(path) for r in rules), path
    assert "/p2/paid-marketing/api/segments/accounts" in rules and "/p2/paid-marketing/keyword-research" in rules


def test_the_agent_is_on_the_paid_marketing_page_the_palette_and_the_hub_counts_it():
    root = os.path.join(ROOT, "templates")
    for name in ("paid_marketing.html", "_paid_kpal.html"):
        assert "/p2/paid-marketing/keyword-research" in open(os.path.join(root, name), encoding="utf-8").read(), name
    assert "Keyword Research" in open(os.path.join(root, "hub.html"), encoding="utf-8").read()
    css = open(os.path.join(ROOT, "static", "css", "paid_marketing.css"), encoding="utf-8").read()
    assert ".c-kwres" in css and "c-kwres" in open(os.path.join(root, "paid_marketing.html"), encoding="utf-8").read()


def test_the_knowledge_the_chat_agents_hold_describes_it():
    src = open(os.path.join(ROOT, "app.py"), encoding="utf-8").read()
    assert "Keyword Research (/p2/paid-marketing/keyword-research)" in src and "paid_keywords_job.py" in src and "never put in the Google Ads Editor file" in src


def test_the_real_dependencies_start_each_agents_own_runner(monkeypatch):
    monkeypatch.undo()
    from tracker import paid_business_job, paid_keywords_job, paid_strategy_job
    ran = []
    monkeypatch.setattr(paid_business_job, "run_build", lambda sid, **kw: ran.append(("business", sid)))
    monkeypatch.setattr(paid_keywords_job, "run_build", lambda sid, **kw: ran.append(("keywords", sid)))
    monkeypatch.setattr(paid_strategy_job, "run_build", lambda sid, **kw: ran.append(("strategy", sid)))
    fake = F.FakeStrategyStore()
    monkeypatch.setattr(appmod, "_strat_spawn", lambda fn: fn())
    monkeypatch.setattr(appmod, "_pm_user_email", lambda: STAFF)
    for inputs, want in ((B.clean(), "business"), (F.clean(), "strategy"), (K.clean(), "keywords")):
        rec = fake.create(by=STAFF, inputs=inputs)
        assert appmod._strat_begin(fake, rec) == (True, "")
        assert ran[-1] == (want, rec["id"])


def test_the_real_dependencies_carry_what_keyword_research_needs(monkeypatch):
    monkeypatch.undo()
    monkeypatch.setattr(gads, "list_accounts", lambda **kw: [{"id": "9999999999", "name": "First", "currency": "GBP"}])
    d = appmod._strat_deps({"created_by": STAFF}, None)
    assert d["planner_account"] == "9999999999" and d["planner_currency"] == "GBP" and callable(d["business_list"]) and callable(d["business_get"]) and callable(d["sqr_read"]) and callable(d["post"])
    d = appmod._strat_deps({"created_by": STAFF}, {"id": "1234567890", "name": "Acme", "currency": "USD"})
    assert d["planner_account"] == "1234567890" and d["planner_currency"] == "USD"
    monkeypatch.setattr(gads, "list_accounts", lambda **kw: (_ for _ in ()).throw(RuntimeError("down")))
    d = appmod._strat_deps({"created_by": STAFF}, None)
    assert d["planner_account"] is None and d["planner_currency"] is None


def test_a_new_build_is_saved_with_this_agents_own_steps(world, monkeypatch):
    saved = []
    original = world.store.create

    def spy(**kw):
        saved.append(kw)
        return original(**kw)
    monkeypatch.setattr(store, "create", spy)
    built(client_as(STAFF))
    assert [s["key"] for s in saved[0]["steps"]][:2] == ["r:business", "r:website"] and saved[0]["steps"][-1]["key"] == "f:finish" and saved[0]["inputs"]["kind"] == "keywords"


def test_the_real_dependencies_look_for_business_research_among_business_builds(monkeypatch):
    monkeypatch.undo()
    asked = []
    monkeypatch.setattr(store, "list_recent", lambda limit=100, kind="strategy": (asked.append(kind), [])[1])
    monkeypatch.setattr(gads, "list_accounts", lambda **kw: [{"id": "9999999999", "name": "First", "currency": "GBP"}])
    d = appmod._strat_deps({"created_by": STAFF}, None)
    assert d["business_list"]() == [] and d["business_list"](kind="business") == [] and asked == ["business", "business"]


# ── Meta ────────────────────────────────────────────────────────────────────

def meta_world(world, monkeypatch):
    from tests import test_paid_keywords_meta_job as MJ
    monkeypatch.setattr(appmod, "_strat_deps", lambda rec, account: MJ.deps(email=rec["created_by"]))
    monkeypatch.setattr(appmod, "_strat_meta_account", lambda act: {"id": act, "name": "Acme Meta", "currency": "EUR"} if act == "act_1234567" else None)


def test_a_meta_account_is_checked_named_and_read_into_a_keyword_research_build(world, monkeypatch):
    meta_world(world, monkeypatch)
    c = client_as(STAFF)
    r = post(c, "/start", dict(FORM, account="", metaAccount="act_1234567", platform="meta"))
    assert r.status_code == 201
    rec = c.get(API + "/" + r.get_json()["id"]).get_json()
    assert rec["status"] == "done" and rec["accountName"] == "Acme Meta" and rec["inputs"]["currency"] == "EUR" and rec["inputs"]["platform"] == "meta" and rec["inputs"]["metaAccount"] == "act_1234567"
    assert rec["plan"]["meta"]["audiences"] and "keywords" not in rec["plan"] and "r:planner" not in [s["key"] for s in rec["progress"]["steps"]]
    assert rec["research"]["metaaccount"]["state"] == "ok" and rec["research"]["interests"]["state"] == "ok"
    bad = post(c, "/start", dict(FORM, metaAccount="act_999999"))
    assert bad.status_code == 400 and "Meta ad account is not one this platform can read" in bad.get_json()["error"]
    assert post(c, "/start", dict(FORM, metaAccount="nonsense")).status_code == 400


def test_the_options_list_the_meta_card_and_a_meta_result_downloads_without_a_google_file(world, monkeypatch):
    meta_world(world, monkeypatch)
    c = client_as(STAFF)
    assert [x["key"] for x in c.get(API + "/options").get_json()["cards"]] == ["business", "website", "account", "metaaccount", "competitors"]
    sid = post(c, "/start", dict(FORM, account="", metaAccount="act_1234567", platform="meta")).get_json()["id"]
    for kind in ("xlsx", "pdf", "zip"):
        r = c.get(API + "/%s/download?file=%s" % (sid, kind))
        assert r.status_code == 200 and len(r.data) > 500, kind
    wb = openpyxl.load_workbook(io.BytesIO(c.get(API + "/%s/download?file=xlsx" % sid).data))
    assert "Meta audiences" in wb.sheetnames and "Keywords" not in wb.sheetnames
    r = c.get(API + "/%s/download?file=csv" % sid)
    assert r.status_code == 404 and "Meta, which has no keywords" in r.get_json()["error"]


def test_the_interest_search_reads_one_page_never_every_page_behind_it(monkeypatch):
    from tracker import paid_meta
    asked = []
    monkeypatch.setattr(paid_meta, "_get", lambda path, params, store, **kw: asked.append((path, params)) or {"data": [{"id": "1", "name": "Insulin pump"}], "paging": {"next": "https://graph.facebook.com/next"}})
    monkeypatch.setattr(paid_meta, "_pages", lambda *a, **k: (_ for _ in ()).throw(AssertionError("a paging reader was used")))
    deps = appmod._strat_deps({"created_by": STAFF, "inputs": {}}, None)
    rows = deps["meta_get"]("/search", {"type": "adinterest", "q": "x"})
    assert rows == [{"id": "1", "name": "Insulin pump"}] and asked == [("/search", {"type": "adinterest", "q": "x"})]
    monkeypatch.setattr(paid_meta, "_get", lambda *a, **k: {})
    assert deps["meta_get"]("/search", {}) == []
