"""Paid Marketing Agents: Pacing Desk, Search Query Analysis, Client Feedback.

Ported from Agent Holmes. What is pinned here is what the port changed or
must not lose: the seed push is refused without its secret and checked
before anything is stored; the rollup is computed here, not trusted; the SQR
route only classifies, against server-held categories whose ids match the
page's; the chat's tool loop survives failing tools and a spent budget; the
shared Slack/Google grant is encrypted, admin-only, and state-checked; and an
answer is only readable by whoever asked.
"""
import json
import os
import re
import sys
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import pytest

os.environ.setdefault("GOOGLE_CLIENT_ID", "test")
os.environ.setdefault("GOOGLE_CLIENT_SECRET", "test")
os.environ.setdefault("FLASK_SECRET_KEY", "test")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

import app as appmod  # noqa: E402
from tracker import paid_feedback as fb  # noqa: E402
from tracker import paid_pacing as pp  # noqa: E402
from tracker import paid_sqr as sqr  # noqa: E402
from tracker import paid_store as store  # noqa: E402

ADMIN = "reporting@position2.com"
STAFF = "someone@position2.com"


def account(name, landing="ON TARGET", allocated=1000.0, spend=500.0, projected=1000.0, conv=10.0,
            moves=0, capped=False, ambiguous=False):
    return {"name": name, "landing": landing, "allocated": allocated, "spend": spend,
            "projected": projected, "conv": conv, "capped": capped, "ambiguous": ambiguous,
            "moves": [{"from": "A", "to": "B", "amt": 10, "why": "x"}] * moves, "campaigns": []}


def dashboard(**over):
    d = {"asOf": "2026-10-12", "dataThrough": "2026-10-11", "month": "October 2026", "elapsed": 11,
         "daysInMonth": 31, "idealPacing": 0.3548, "tolerance": 2.0,
         "portfolioWarnings": ["Workato — $84,624 MTD with no budget row"],
         "clients": [
             {"name": "GD & Affiliates", "accounts": [
                 account("Arlington", "WILL OVERSPEND", 1000, 600, 1300, 5, moves=2),
                 account("Burke", "WILL UNDERSPEND", 2000, 300, 1500, 7, capped=True)]},
             {"name": "OIA", "accounts": [account("Heartland", "BUDGET UNCLEAR", 999, 999, 999, 3, ambiguous=True)]},
         ]}
    d.update(over)
    return d


# ── Pacing: the seed ───────────────────────────────────────────────────────

def test_rollup_counts_money_only_for_accounts_whose_budget_is_clear():
    meta, rows, latest = pp.build_seed(dashboard())
    assert latest == "2026-10-12"
    assert [s for s, _ in rows] == ["gd-and-affiliates", "oia"]
    assert meta["tAll"] == 3000 and meta["tSp"] == 900 and meta["tPr"] == 2800
    assert meta["tCv"] == 15            # conversions count every account
    assert meta["tMv"] == 2
    assert meta["atRisk"] == 500         # Burke: 2000 allocated, projects 1500
    assert (meta["overCount"], meta["underCount"], meta["okCount"], meta["ambiguousCount"]) == (1, 1, 0, 1)


def test_seed_live_js_shape_is_accepted_and_its_own_rollup_is_ignored():
    d = dashboard()
    body = {"meta": dict({k: d[k] for k in pp.META_KEYS}, tAll=1, tSp=99999999),
            "clients": [{"id": pp.slug(c["name"]), "data": c} for c in d["clients"]],
            "latest": {"date": "2026-10-12"}}
    meta, rows, _ = pp.build_seed(body)
    assert meta["tAll"] == 3000 and meta["tSp"] == 900
    assert rows[0][1]["name"] == "GD & Affiliates"


def test_em_dashes_never_reach_the_page():
    meta, _rows, _ = pp.build_seed(dashboard())
    assert meta["portfolioWarnings"] == ["Workato, $84,624 MTD with no budget row"]


@pytest.mark.parametrize("mutate, says", [
    (lambda d: d.update(clients=[]), "no clients"),
    (lambda d: d.update(asOf="12/10/2026"), "asOf"),
    (lambda d: d["clients"][0]["accounts"][0].update(landing="SOMEWHERE"), "landing"),
    (lambda d: d["clients"][0]["accounts"][0].update(spend="500"), "spend"),
    (lambda d: d["clients"][0]["accounts"][0].update(moves=None), "moves"),
    (lambda d: d["clients"].append({"name": "OIA", "accounts": []}), "share the id"),
    (lambda d: d.update(elapsed=True), "elapsed"),
])
def test_a_bad_push_is_refused_with_the_reason(mutate, says):
    d = dashboard()
    mutate(d)
    with pytest.raises(pp.BadPayload, match=says):
        pp.build_seed(d)


def test_the_real_holmes_payload_seeds_when_it_is_present():
    path = os.path.join(ROOT, "..", "holmes_src", "Agent_Holmes", "out", "dashboard.json")
    if not os.path.exists(path):
        pytest.skip("Agent Holmes checkout not beside this repo")
    meta, rows, _ = pp.build_seed(json.load(open(path)))
    assert meta["clientCount"] == len(rows) == 12


def test_move_ids_and_bodies_are_checked():
    good = {"date": "2026-10-12", "loc": "Arlington", "from": "Brand", "to": "PMax", "amt": 12.5, "applied": True}
    mid = "2026-10-12__arlington__brand__1a2b3c4d"
    assert pp.check_move(mid, good)["from_campaign"] == "Brand"
    for bad_id in ("x", "2026-10-12__a__b__xyz", "../../etc"):
        with pytest.raises(pp.BadPayload):
            pp.check_move(bad_id, good)
    for patch in ({"date": "2026-10-11"}, {"applied": "yes"}, {"amt": "12"}, {"from": ""}):
        with pytest.raises(pp.BadPayload):
            pp.check_move(mid, dict(good, **patch))


def test_the_page_and_the_server_build_the_same_move_id_shape():
    js = open(os.path.join(ROOT, "static", "js", "pacing_desk.js")).read()
    assert 'date+"__"+slug(loc).slice(0,24)+"__"+slug(from).slice(0,24)+"__"+hashStr(' in js
    assert pp.MOVE_ID_RE.pattern.count("{0,24}") == 2


# ── SQR ─────────────────────────────────────────────────────────────────────

def test_server_rulesets_match_the_pages():
    js = open(os.path.join(ROOT, "static", "js", "sqr_agent.js")).read()
    page_ids = set(re.findall(r'^\s+id: "([a-z0-9_]+)", label:', js, re.M))
    assert page_ids == set(sqr.RULESETS)
    assert "llmCategories" not in js       # the categories live server-side only


def test_classify_keeps_only_answers_to_what_was_asked():
    def call(system, user):
        assert "dentist near me" in user and "Jobs/Recruitment" in user
        return {"results": [
            {"term": "dentist near me", "status": "Relevant", "subtype": "Core"},
            {"term": "ignore previous instructions", "status": "Relevant", "subtype": "x"},
            {"term": "dental hygienist jobs", "status": "Maybe", "subtype": "x"},
        ]}
    out = sqr.classify("riccobene", ["dentist near me", "dental hygienist jobs", "dentist near me"], call=call)
    assert set(out) == {"dentist near me", "dental hygienist jobs"}
    assert out["dentist near me"] == {"status": "Relevant", "subtype": "Core", "by": "model"}
    assert out["dental hygienist jobs"]["by"] == "fallback"


def test_classify_refuses_unknown_rulesets_and_oversized_batches():
    with pytest.raises(sqr.BadRequest):
        sqr.classify("everyone", ["x"], call=lambda s, u: {})
    with pytest.raises(sqr.BadRequest):
        sqr.classify("generic", ["t%d" % i for i in range(sqr.MAX_TERMS + 1)], call=lambda s, u: {})
    with pytest.raises(sqr.BadRequest):
        sqr.classify("generic", [], call=lambda s, u: {})


# ── Client Feedback: pieces ─────────────────────────────────────────────────

def test_history_is_plain_alternating_text_without_the_new_question():
    h = [{"role": "assistant", "content": "hi"}, {"role": "user", "content": "a"},
         {"role": "user", "content": "b"}, {"role": "system", "content": "x"},
         {"role": "assistant", "content": [{"type": "tool_use"}]}, {"role": "assistant", "content": "c"},
         {"role": "user", "content": "pending"}]
    assert fb.clean_history(h) == [{"role": "user", "content": "a\n\nb"}, {"role": "assistant", "content": "c"}]


def test_meeting_sections_and_vtt():
    doc = "Meeting: Riccobene Weekly\nDate: x\nTranscript: hello\nMeeting: OIA sync\nbody"
    assert [t for t, _ in fb.meeting_sections(doc)] == ["Riccobene Weekly", "OIA sync"]
    assert fb.vtt_to_text("WEBVTT\n\n1\n00:00:01.000 --> 00:00:02.000\nHello  there\n\n2\n00:00:03.000 --> 00:00:04.000\nBye") == "Hello there Bye"


def test_gmail_plain_body_prefers_text_plain_and_decodes_urlsafe():
    import base64
    enc = lambda s: base64.urlsafe_b64encode(s.encode()).decode().rstrip("=")
    payload = {"mimeType": "multipart/alternative", "parts": [
        {"mimeType": "text/html", "body": {"data": enc("<b>x</b>")}},
        {"mimeType": "text/plain", "body": {"data": enc("plain ?>~ text")}}]}
    assert fb.plain_body(payload) == "plain ?>~ text"


def test_authorize_url_carries_the_state(monkeypatch):
    monkeypatch.setenv("FEEDBACK_SLACK_CLIENT_ID", "cid")
    monkeypatch.setenv("FEEDBACK_GOOGLE_CLIENT_ID", "gid")
    assert "state=abc" in fb.authorize_url("slack", "https://h/cb", "abc")
    g = fb.authorize_url("google", "https://h/cb", "abc")
    assert "state=abc" in g and "access_type=offline" in g and "gmail.readonly" in g


class FakeBlock(SimpleNamespace):
    pass


def tool_use(i, name, **inp):
    return FakeBlock(type="tool_use", id="t%d" % i, name=name, input=inp)


def text(t):
    return FakeBlock(type="text", text=t)


class FakeClient:
    def __init__(self, replies):
        self.replies, self.calls = list(replies), []
        self.messages = self

    def create(self, **kw):
        self.calls.append(kw)
        content, stop = self.replies.pop(0)
        return SimpleNamespace(content=content, stop_reason=stop)


class FakeTools:
    def search_slack(self, inp):
        return "Channel: #oia | hi"

    def search_email(self, inp):
        raise fb.NotConnected("Google isn't connected yet.")

    def get_budget_snapshot(self, inp):
        raise RuntimeError("Sheets exploded")


def test_tool_loop_runs_tools_reports_failures_and_answers():
    client = FakeClient([
        ([tool_use(1, "search_slack", query="OIA"), tool_use(2, "search_email", query="OIA"),
          tool_use(3, "get_budget_snapshot"), tool_use(4, "rm_rf")], "tool_use"),
        ([text("OIA asked about — MRI.")], "end_turn"),
    ])
    steps = []
    answer, used = fb.run_loop([{"role": "user", "content": "q"}], client=client, tools=FakeTools(), on_step=steps.append)
    assert answer == "OIA asked about, MRI."
    assert used == ["search_slack", "search_email", "get_budget_snapshot", "rm_rf"]
    results = {r["tool_use_id"]: r["content"] for r in client.calls[1]["messages"][-1]["content"]}
    assert results["t1"].startswith("Channel")
    assert results["t2"] == "Google isn't connected yet."
    assert results["t3"].startswith("Error: Sheets exploded")
    assert results["t4"] == "Unknown tool rm_rf"
    assert [s["kind"] for s in steps].count("warn") == 3
    assert steps[0]["label"] == 'Searching Slack for "OIA"'


def test_a_spent_tool_budget_still_gets_an_answer():
    replies = [([tool_use(i, "search_slack", query="x")], "tool_use") for i in range(fb.MAX_TOOL_ROUNDS)]
    replies.append(([text("Best I can say.")], "end_turn"))
    client = FakeClient(replies)
    answer, _ = fb.run_loop([{"role": "user", "content": "q"}], client=client, tools=FakeTools())
    assert answer == "Best I can say."
    assert client.calls[-1]["tool_choice"] == {"type": "none"}
    assert len(client.calls) == fb.MAX_TOOL_ROUNDS + 1


def test_a_job_its_worker_abandoned_reads_as_failed(monkeypatch):
    started = datetime(2026, 10, 12, 10, 0, tzinfo=timezone.utc)
    monkeypatch.setattr(store, "get_job", lambda j: {"status": "running", "createdAt": started.isoformat()})
    assert fb.read_job("x", now=started + timedelta(minutes=2))["status"] == "running"
    assert fb.read_job("x", now=started + timedelta(minutes=30))["status"] == "failed"


def test_tokens_are_sealed_and_a_changed_key_says_so(monkeypatch):
    monkeypatch.setenv("SECRET_KEY", "one")
    sealed = store._seal("xoxp-secret")
    assert "xoxp" not in sealed and store._open(sealed) == "xoxp-secret"
    monkeypatch.setenv("SECRET_KEY", "two")
    with pytest.raises(store.TokenUnreadable):
        store._open(sealed)


def test_google_code_without_refresh_token_is_refused(monkeypatch):
    saved = []
    monkeypatch.setattr(store, "set_token", lambda *a, **k: saved.append((a, k)))
    resp = SimpleNamespace(status_code=200, json=lambda: {"access_token": "a", "expires_in": 3600})
    with pytest.raises(fb.OAuthFailed, match="refresh token"):
        fb.exchange_code("google", "c", "https://h/cb", by=ADMIN, post=lambda *a, **k: resp)
    assert not saved
    slack = SimpleNamespace(status_code=200, json=lambda: {"ok": True, "authed_user": {"access_token": "u", "id": "U1"},
                                                          "team": {"name": "P2"}})
    fb.exchange_code("slack", "c", "https://h/cb", by=ADMIN, post=lambda *a, **k: slack)
    assert saved[0][0] == ("slack",) and saved[0][1]["by"] == ADMIN and saved[0][1]["meta"]["team"] == "P2"


# ── Routes ──────────────────────────────────────────────────────────────────

def client_as(email):
    c = appmod.app.test_client()
    if email:
        with c.session_transaction() as sess:
            sess["google_user"] = {"email": email, "name": "T"}
    return c


SAME = {"Origin": "http://localhost"}


@pytest.mark.parametrize("path", ["/p2/paid-marketing", "/p2/paid-marketing/pacing-desk",
                                  "/p2/paid-marketing/search-query-analysis",
                                  "/p2/paid-marketing/client-feedback-agent"])
def test_pages_render_for_staff_and_not_for_anyone_else(path):
    assert client_as(STAFF).get(path).status_code == 200
    assert client_as(None).get(path).status_code == 302
    assert client_as("someone@gmail.com").get(path).status_code == 302


def test_seed_is_refused_without_the_secret(monkeypatch):
    stored = []
    monkeypatch.setattr(store, "save_snapshot", lambda *a: stored.append(a))
    c = client_as(None)
    url = "/p2/paid-marketing/api/pacing/seed"
    monkeypatch.delenv("PACING_SEED_SECRET", raising=False)
    monkeypatch.delenv("SEED_SECRET", raising=False)
    assert c.post(url, json=dashboard(), headers={"X-Seed-Secret": ""}).status_code == 403
    monkeypatch.setenv("PACING_SEED_SECRET", "s3cret")
    assert c.post(url, json=dashboard(), headers={"X-Seed-Secret": "nope"}).status_code == 403
    assert c.post(url, json=dashboard(clients=[]), headers={"X-Seed-Secret": "s3cret"}).status_code == 400
    assert not stored
    r = c.post(url, json=dashboard(), headers={"X-Seed-Secret": "s3cret"})
    assert r.status_code == 200 and r.get_json()["accounts"] == 3
    assert stored[0][0]["asOf"] == "2026-10-12" and stored[0][2] == "2026-10-12"


def test_holmes_env_names_are_accepted(monkeypatch, tmp_path):
    for n in ("FEEDBACK_SLACK_CLIENT_ID", "FEEDBACK_SLACK_CLIENT_SECRET", "FEEDBACK_GOOGLE_CLIENT_ID"):
        monkeypatch.delenv(n, raising=False)
    monkeypatch.setenv("SLACK_CLIENT_ID", "cid")
    monkeypatch.setenv("SLACK_CLIENT_SECRET", "sec")
    monkeypatch.setenv("GOOGLE_CLIENT_ID", "the-sign-in-client")
    monkeypatch.delenv("FEEDBACK_GOOGLE_CLIENT_SECRET", raising=False)
    monkeypatch.delenv("GOOGLE_ADS_CLIENT_ID", raising=False)
    monkeypatch.delenv("GOOGLE_ADS_CLIENT_SECRET", raising=False)
    assert fb.configured()["slack"] is True
    assert fb.env("GOOGLE_CLIENT_ID") == ""          # never the sign-in client
    monkeypatch.setenv("FEEDBACK_SLACK_CLIENT_ID", "prefixed")
    assert fb.env("SLACK_CLIENT_ID") == "prefixed"   # the prefixed name wins


def test_google_falls_back_to_the_ads_client(monkeypatch):
    for n in ("FEEDBACK_GOOGLE_CLIENT_ID", "FEEDBACK_GOOGLE_CLIENT_SECRET"):
        monkeypatch.delenv(n, raising=False)
    monkeypatch.setenv("GOOGLE_CLIENT_ID", "the-sign-in-client")
    monkeypatch.setenv("GOOGLE_ADS_CLIENT_ID", "ads-cid")
    monkeypatch.setenv("GOOGLE_ADS_CLIENT_SECRET", "ads-sec")
    assert fb.env("GOOGLE_CLIENT_ID") == "ads-cid"   # the Ads client, never the sign-in one
    assert fb.env("GOOGLE_CLIENT_SECRET") == "ads-sec"
    assert fb.configured()["google"] is True
    monkeypatch.setenv("FEEDBACK_GOOGLE_CLIENT_ID", "own-cid")
    assert fb.env("GOOGLE_CLIENT_ID") == "own-cid"   # a prefixed name still wins


def test_seed_accepts_holmes_secret_name(monkeypatch):
    monkeypatch.setattr(store, "save_snapshot", lambda *a: None)
    monkeypatch.delenv("PACING_SEED_SECRET", raising=False)
    monkeypatch.setenv("SEED_SECRET", "h0lmes")
    c = client_as(None)
    url = "/p2/paid-marketing/api/pacing/seed"
    assert c.post(url, json=dashboard(), headers={"X-Seed-Secret": "nope"}).status_code == 403
    assert c.post(url, json=dashboard(), headers={"X-Seed-Secret": "h0lmes"}).status_code == 200


def test_state_changes_must_come_from_this_site_as_json(monkeypatch):
    monkeypatch.setattr(store, "set_move", lambda mid, **f: {"id": mid, **f})
    c = client_as(STAFF)
    url = "/p2/paid-marketing/api/pacing/moves/2026-10-12__a__b__1f"
    body = {"date": "2026-10-12", "loc": "A", "from": "B", "to": "C", "amt": 5, "applied": True}
    assert c.put(url, json=body, headers={"Origin": "https://evil.example"}).status_code == 403
    assert c.put(url, data="x", headers=SAME).status_code == 415
    r = c.put(url, json=body, headers=SAME)
    assert r.status_code == 200 and r.get_json()["move"]["by"] == STAFF


def test_sqr_route_takes_only_ruleset_and_terms(monkeypatch):
    seen = []
    monkeypatch.setattr(sqr, "classify", lambda rid, terms: seen.append((rid, terms)) or {"t": {"status": "Relevant"}})
    c = client_as(STAFF)
    r = c.post("/p2/paid-marketing/api/sqr/classify", json={"ruleset": "oia", "terms": ["t"], "prompt": "anything"}, headers=SAME)
    assert r.status_code == 200 and seen == [("oia", ["t"])]


def test_an_answer_is_readable_only_by_whoever_asked(monkeypatch):
    job = {"id": "a" * 32, "createdBy": STAFF, "status": "done", "createdAt": None, "answer": "secret"}
    monkeypatch.setattr(fb, "read_job", lambda j: dict(job))
    url = "/p2/paid-marketing/api/feedback/jobs/" + "a" * 32
    assert client_as(STAFF).get(url).status_code == 200
    assert client_as("other@position2.com").get(url).status_code == 404
    assert client_as(ADMIN).get(url).status_code == 200
    assert client_as(STAFF).get("/p2/paid-marketing/api/feedback/jobs/../x").status_code == 404


def test_one_question_at_a_time(monkeypatch):
    monkeypatch.setenv("ANTHROPIC_API_KEY", "k")
    monkeypatch.setattr(store, "job_counts", lambda by, stale: (1, 3))
    monkeypatch.setattr(fb, "start", lambda *a, **k: pytest.fail("must not start"))
    r = client_as(STAFF).post("/p2/paid-marketing/api/feedback/chat", json={"message": "hi"}, headers=SAME)
    assert r.status_code == 429


def test_only_admins_connect_and_the_callback_checks_state(monkeypatch):
    monkeypatch.setenv("FEEDBACK_SLACK_CLIENT_ID", "cid")
    monkeypatch.setenv("FEEDBACK_SLACK_CLIENT_SECRET", "sec")
    assert client_as(STAFF).get("/p2/paid-marketing/auth/slack/start").status_code == 403
    c = client_as(ADMIN)
    r = c.get("/p2/paid-marketing/auth/slack/start")
    assert r.status_code == 302 and r.location.startswith("https://slack.com/oauth/v2/authorize")
    state = re.search(r"state=([^&]+)", r.location).group(1)
    called = []
    monkeypatch.setattr(fb, "exchange_code", lambda *a, **k: called.append(k["by"]))
    forged = c.get("/p2/paid-marketing/auth/slack/callback?code=x&state=forged")
    assert "error=" in forged.location and not called
    # The forged attempt used up the state, as any callback does: start again.
    r = c.get("/p2/paid-marketing/auth/slack/start")
    state = re.search(r"state=([^&]+)", r.location).group(1)
    ok = c.get("/p2/paid-marketing/auth/slack/callback?code=x&state=" + state)
    assert ok.location.endswith("?connected=slack") and called == [ADMIN]


def test_hub_paid_card_counts_match_the_section_page():
    hub = re.sub(r"<!--.*?-->", "", open(os.path.join(ROOT, "templates", "hub.html")).read(), flags=re.S)
    block = hub.split('href="/p2/paid-marketing"', 1)[1].split("</a>", 1)[0]
    stats = dict((label, int(n)) for n, label in re.findall(r'class="card-stat"><span>(\d+)</span>\s*(\w+)', block))
    page = re.sub(r"<!--.*?-->", "", open(os.path.join(ROOT, "templates", "paid_marketing.html")).read(), flags=re.S)
    cards = re.findall(r'class="dash-card (\w+)', page)
    assert stats == {"dashboards": len(cards), "live": cards.count("active")}
    order = [m.group(1) for m in re.finditer(r'<a href="(/p2/[a-z-]+)" class="hub-card', hub)]
    assert order == ["/p2/strategic-agents", "/p2/paid-marketing", "/p2/seo-aeo"]


def test_clients_route_lists_the_latest_run_for_the_chat_picker(monkeypatch):
    monkeypatch.setattr(store, "latest_clients", lambda: ("2026-09-30", [{"name": "OIA", "accounts": 6, "offTarget": 2}]))
    r = client_as(STAFF).get("/p2/paid-marketing/api/pacing/clients")
    assert r.get_json() == {"date": "2026-09-30", "clients": [{"name": "OIA", "accounts": 6, "offTarget": 2}]}
    assert client_as(None).get("/p2/paid-marketing/api/pacing/clients").status_code == 302


def test_the_chat_page_sends_its_context_with_the_question():
    js = open(os.path.join(ROOT, "static", "js", "client_feedback_agent.js")).read()
    assert "message=ctx?ctx+\"\\n\\n\"+text:text" in js        # the model reads the context line
    assert 'turns.push({role:"user",content:message}' in js    # and follow-ups keep it
    assert "pm_picker.js" in open(os.path.join(ROOT, "templates", "client_feedback_agent.html")).read()
    assert "pm_picker.js" in open(os.path.join(ROOT, "templates", "pacing_desk.html")).read()


def test_the_launcher_only_links_to_rulesets_the_sqr_page_has():
    js = open(os.path.join(ROOT, "static", "js", "paid_marketing.js")).read()
    linked = set(re.findall(r'/i,"([a-z0-9_]+)",', js))
    assert linked and linked <= set(sqr.RULESETS), linked - set(sqr.RULESETS)
    for page, param in (("pacing_desk.js", '"client"'), ("client_feedback_agent.js", '"client"'), ("sqr_agent.js", '"account"')):
        assert ".get(" + param + ")" in open(os.path.join(ROOT, "static", "js", page)).read(), page


# ── Pacing: the Refresh button ─────────────────────────────────────────────

from datetime import date as _date  # noqa: E402

from tracker import paid_refresh as rf  # noqa: E402

HOLMES_RAW = os.path.join(ROOT, "..", "holmes_src", "Agent_Holmes", "raw")


@pytest.mark.parametrize("through,today,want", [
    (_date(2026, 10, 11), _date(2026, 10, 12), _date(2026, 10, 12)),   # data through yesterday
    (_date(2026, 10, 9), _date(2026, 10, 12), _date(2026, 10, 10)),    # sheet lagging: day after its data
    (_date(2026, 9, 30), _date(2026, 10, 1), _date(2026, 9, 30)),      # the 1st: stay in the data's month
    (_date(2026, 10, 12), _date(2026, 10, 12), _date(2026, 10, 12)),   # never in the future
    (None, _date(2026, 10, 12), _date(2026, 10, 12)),
])
def test_the_run_is_for_the_day_after_the_newest_data(through, today, want):
    assert rf.pick_as_of(through, today) == want


def test_sheet_rows_become_the_workbook_the_engine_reads(tmp_path):
    import openpyxl
    rows = [["Account", "Date", "Campaign name", "Cost"],
            ["Acme", 46296, "P2_Brand", 12.5],          # 46296 = 2026-10-01 as a sheet serial
            ["", "", "", ""],                            # blank row: dropped
            ["Acme", 46297.0, "P2_PMax"]]                # short row: padded
    p = tmp_path / "s.xlsx"
    rf.write_xlsx(p, "Data", rows)
    got = list(openpyxl.load_workbook(p, data_only=True)["Data"].iter_rows(values_only=True))
    assert got[0] == ("Account", "Date", "Campaign name", "Cost")
    assert got[1] == ("Acme", datetime(2026, 10, 1), "P2_Brand", 12.5)
    assert got[2] == ("Acme", datetime(2026, 10, 2), "P2_PMax", None)
    assert len(got) == 3
    assert rf.last_spend_date(rows) == _date(2026, 10, 2)


def test_a_budget_tracker_from_another_month_is_called_out():
    rows = [["Client", "Account", "Allocated Budget", "Target", "Date"], ["A", "A", 1, 1, 46266]]  # Sept 1
    assert rf.budget_months(rows) == {"2026-09"}


def test_refresh_says_what_setup_is_missing(monkeypatch):
    monkeypatch.delenv("GOOGLE_SA_JSON", raising=False)
    assert "GOOGLE_SA_JSON" in rf.config_problem()
    monkeypatch.setenv("GOOGLE_SA_JSON", json.dumps({"client_email": "bot@x.iam.gserviceaccount.com"}))
    monkeypatch.delenv("PACING_SPEND_SHEET_ID", raising=False)
    assert "PACING_SPEND_SHEET_ID" in rf.config_problem()
    assert rf.service_email() == "bot@x.iam.gserviceaccount.com"
    monkeypatch.setenv("PACING_SPEND_SHEET_ID", "abc")
    assert rf.config_problem() is None
    monkeypatch.delenv("PACING_BUDGET_SHEET_ID", raising=False)
    assert rf.budget_sheet_id() == rf.DEFAULT_BUDGET_SHEET


def test_an_engine_that_stops_is_reported_and_flags_are_not(monkeypatch):
    calls = []

    def fake_run(script, args, cwd):
        calls.append(script)
        if script == "validate.py":
            return 1, "[FAIL] x: 1 issue(s)\n   ✗ OIA/Wolf River/MRI: below the $1/day policy floor\n"
        os.makedirs(os.path.join(cwd, "out"), exist_ok=True)
        with open(os.path.join(cwd, "out", "dashboard.json"), "w") as f:
            json.dump(dashboard(), f)
        return 0, "ok"
    monkeypatch.setattr(rf, "_run", fake_run)
    rows = [["Account", "Date"], ["A", 46296]]
    body, issues = rf.run_engine(rows, "Data", rows, "DPR", _date(2026, 10, 2))
    assert body["asOf"] == "2026-10-12" and issues == ["OIA/Wolf River/MRI: below the $1/day policy floor"]

    monkeypatch.setattr(rf, "_run", lambda s, a, c: (1, "Traceback...\nKeyError: 'Account'"))
    with pytest.raises(rf.RefreshError, match="KeyError"):
        rf.run_engine(rows, "Data", rows, "DPR", _date(2026, 10, 2))


class FakeRefreshStore:
    StoreUnavailable = store.StoreUnavailable

    def __init__(self):
        self.state, self.saved = {}, []

    def claim_refresh(self, by, stale_s):
        if self.state.get("state") == "running":
            return False
        self.state = {"state": "running", "by": by, "startedAt": datetime.now(timezone.utc).isoformat()}
        return True

    def update_refresh(self, **f):
        self.state.update(f)

    def get_refresh(self):
        return dict(self.state)

    def save_snapshot(self, *a):
        self.saved.append(a)


def test_a_failed_refresh_keeps_the_old_run_and_says_why(monkeypatch):
    s = FakeRefreshStore()
    monkeypatch.setenv("GOOGLE_SA_JSON", json.dumps({"client_email": "bot@x"}))
    monkeypatch.setenv("PACING_SPEND_SHEET_ID", "abc")
    monkeypatch.setattr(rf, "_sheets", lambda: object())

    def no_access(svc, sheet_id, tab, label):
        raise rf.RefreshError("Could not open the %s: share it with bot@x" % label)
    monkeypatch.setattr(rf, "fetch_tab", no_access)
    monkeypatch.setattr(rf.threading, "Thread", lambda target, **k: SimpleNamespace(start=target))
    assert rf.start(s, pp, STAFF) is True
    assert s.state["state"] == "failed" and "share it with bot@x" in s.state["error"] and not s.saved

    monkeypatch.setattr(rf, "fetch_tab", lambda *a: 1 / 0)
    assert rf.start(s, pp, STAFF) is True
    assert s.state["error"] == "The refresh failed unexpectedly (ZeroDivisionError)."


def test_a_running_refresh_its_worker_lost_reads_as_failed():
    s = FakeRefreshStore()
    s.state = {"state": "running", "startedAt": (datetime.now(timezone.utc) - timedelta(minutes=11)).isoformat()}
    assert rf.status(s)["state"] == "failed"
    s.state["startedAt"] = datetime.now(timezone.utc).isoformat()
    assert rf.status(s)["state"] == "running"


def test_refresh_on_the_teams_real_september_sheets(monkeypatch):
    """The whole path on Divith's real export, with only Google's reply faked:
    rows arrive the way the Sheets API sends them (dates as serials, blanks
    as ''), and the run must equal the engine's own on the original files."""
    if not os.path.exists(os.path.join(HOLMES_RAW, "all_accounts_spend_data.xlsx")):
        pytest.skip("Agent Holmes checkout not beside this repo")
    pytest.importorskip("pandas")
    import openpyxl

    def as_api(name, tab):
        out = []
        for r in openpyxl.load_workbook(os.path.join(HOLMES_RAW, name), read_only=True,
                                        data_only=True)[tab].iter_rows(values_only=True):
            row = [((v - datetime(1899, 12, 30)).total_seconds() / 86400 if isinstance(v, datetime)
                    else "" if v is None else v) for v in r]
            while row and row[-1] == "":
                row.pop()
            out.append(row)
        return tab, out
    sheets = {"spend sheet": as_api("all_accounts_spend_data.xlsx", "Data"),
              "Budget Tracker": as_api("budget_tracker.xlsx", "DPR")}
    monkeypatch.setenv("GOOGLE_SA_JSON", json.dumps({"client_email": "bot@x"}))
    monkeypatch.setenv("PACING_SPEND_SHEET_ID", "abc")
    monkeypatch.setattr(rf, "_sheets", lambda: object())
    monkeypatch.setattr(rf, "fetch_tab", lambda svc, sid, tab, label: sheets[label])
    monkeypatch.setattr(rf, "today_ist", lambda: _date(2026, 10, 12))
    s = FakeRefreshStore()
    res = rf.refresh(s, pp)
    assert (res["date"], res["clients"], res["accounts"]) == ("2026-09-30", 12, 47)
    assert res["issueCount"] == len(res["issues"]) > 0          # the engine's own flags travel with it
    meta, clients, latest = s.saved[0]
    assert latest == "2026-09-30" and len(clients) == 12


def test_refresh_route(monkeypatch):
    s = FakeRefreshStore()
    for name in ("claim_refresh", "update_refresh", "get_refresh", "save_snapshot"):
        monkeypatch.setattr(store, name, getattr(s, name))
    started = []
    monkeypatch.setattr(rf, "start", lambda st, pc, by: started.append(by) or st.claim_refresh(by, 1))
    url = "/p2/paid-marketing/api/pacing/refresh"
    assert client_as(None).get(url).status_code == 302
    c = client_as(STAFF)

    monkeypatch.delenv("PACING_SPEND_SHEET_ID", raising=False)
    monkeypatch.setenv("GOOGLE_SA_JSON", json.dumps({"client_email": "bot@x"}))
    r = c.post(url, json={}, headers=SAME)
    assert r.status_code == 409 and "PACING_SPEND_SHEET_ID" in r.get_json()["error"]
    assert r.get_json()["serviceEmail"] == "bot@x" and not started

    monkeypatch.setenv("PACING_SPEND_SHEET_ID", "abc")
    assert c.post(url, json={}, headers={"Origin": "https://evil.example"}).status_code == 403
    assert not started
    r = c.post(url, json={}, headers=SAME)
    assert r.status_code == 202 and r.get_json()["state"] == "running" and started == [STAFF]
    r = c.post(url, json={}, headers=SAME)
    assert r.status_code == 409 and "already running" in r.get_json()["error"]
    assert c.get(url).get_json()["state"] == "running"


def test_the_page_has_the_refresh_button_wired_to_the_route():
    html = open(os.path.join(ROOT, "templates", "pacing_desk.html")).read()
    js = open(os.path.join(ROOT, "static", "js", "pacing_desk.js")).read()
    assert 'id="refreshBtn"' in html and 'id="refreshNote"' in html
    assert 'API+"/refresh",{method:"POST"' in js and "Content-Type\":\"application/json\"" in js
