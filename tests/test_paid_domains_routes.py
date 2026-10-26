"""Paid Marketing Agents: Unknown Domains Alert, the page and its routes.

Pinned here: the page is staff-only; every state-changing route refuses a cross-site request; the scheduled scan is refused without the secret (and always when
the secret is not set) and, with it, needs no session; Scan now starts one scan and refuses a second while one runs; the database's add, remove, import (CSV text,
a name for an account or client, bad rows listed) and export; approving, dismissing and reopening a finding; the baseline needs a scan first; settings are
validated; a test alert goes through the senders; and an unavailable store is a 503. Nothing here talks to Google, Meta, Slack or a mail server.
"""
import os
import sys

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
os.environ.setdefault("GOOGLE_CLIENT_ID", "test")
os.environ.setdefault("GOOGLE_CLIENT_SECRET", "test")
os.environ.setdefault("FLASK_SECRET_KEY", "test")

import app as appmod  # noqa: E402
from domains_fakes import Deliver, MemStore, Site  # noqa: E402
from tracker import paid_domains as d  # noqa: E402
from tracker import paid_domains_scan as sc  # noqa: E402
from tracker import paid_links as pl  # noqa: E402

STAFF = "someone@position2.com"
API = "/p2/paid-marketing/api/domains"
SAME = {"Origin": "http://localhost"}
A = {"id": "1112223330", "name": "Acme Garden", "platform": "google", "group": "mcc-1"}
B = {"id": "4445556660", "name": "Bolt Dental", "platform": "google", "group": "mcc-2"}
GROUPS = [{"id": "mcc-1", "name": "Acme", "accountIds": [A["id"]]}, {"id": "mcc-2", "name": "Bolt", "accountIds": [B["id"]]}]


def client_as(email):
    c = appmod.app.test_client()
    if email:
        with c.session_transaction() as sess:
            sess["google_user"] = {"email": email, "name": "T"}
    return c


def ad(account, url, i=0):
    return {"platform": "google", "accountId": account["id"], "accountName": account["name"], "where": "ad", "objectId": str(i), "objectName": "Ad %d" % i,
            "campaign": "C", "adGroup": "G", "url": url, "host": d.host_of(url), "redirect": False}


@pytest.fixture
def world(monkeypatch):
    store, deliver = MemStore(), Deliver()
    ads = {A["id"]: ["https://acme-garden.com/"], B["id"]: ["https://strange.example/"]}
    monkeypatch.setattr(appmod, "_dom_store", lambda: store)
    monkeypatch.setattr(appmod, "_dom_deliver", lambda: deliver)
    monkeypatch.setattr(appmod, "_dom_inventory", lambda s: ([A, B], GROUPS, []))
    monkeypatch.setattr(appmod, "_dom_spawn", lambda fn: fn())
    monkeypatch.setattr(sc, "read_account", lambda a, s: ([ad(a, u, i) for i, u in enumerate(ads[a["id"]])], {"ads": len(ads[a["id"]]), "assetGroups": 0, "assets": 0}))
    monkeypatch.setattr(pl, "default_fetch", lambda: Site())
    monkeypatch.setenv("DOMAIN_ALERT_SECRET", "s3cret")
    store.state_set("settings", {"emails": ["ppc@example.com"], "slackChannel": "#alerts"})
    return type("W", (), {"store": store, "deliver": deliver, "ads": ads})()


def post(path, body=None, email=STAFF, headers=SAME):
    return client_as(email).post(API + path, json=body if body is not None else {}, headers=headers)


# ── Access ──────────────────────────────────────────────────────────────────

def test_the_page_is_for_staff_only():
    path = "/p2/paid-marketing/unknown-domains"
    assert client_as(STAFF).get(path).status_code == 200
    assert client_as(None).get(path).status_code == 302
    assert client_as("someone@gmail.com").get(path).status_code == 302
    assert b"unknown_domains.js" in client_as(STAFF).get(path).data


@pytest.mark.parametrize("method, path", [("post", "/scan"), ("post", "/known"), ("post", "/baseline"), ("post", "/settings"), ("post", "/test-alert"),
                                          ("post", "/known/import"), ("delete", "/known/1"), ("post", "/findings/%s/dismiss" % ("a" * 24))])
def test_a_cross_site_request_is_refused(world, method, path):
    r = getattr(client_as(STAFF), method)(API + path, json={}, headers={"Origin": "https://evil.example"})
    assert r.status_code == 403
    assert getattr(client_as(None), method)(API + path, json={}).status_code in (302, 401, 403)


def test_the_state_needs_a_staff_login(world):
    assert client_as(None).get(API + "/state").status_code == 302
    r = client_as(STAFF).get(API + "/state")
    assert r.status_code == 200 and {"known", "findings", "settings", "senders", "runs", "accounts", "groups"} <= set(r.get_json())


# ── The scheduled scan ──────────────────────────────────────────────────────

def test_the_scheduled_scan_needs_the_secret(world, monkeypatch):
    c = client_as(None)
    assert c.post(API + "/scheduled").status_code == 403
    assert c.post(API + "/scheduled", headers={"X-Domain-Secret": "wrong"}).status_code == 403
    monkeypatch.delenv("DOMAIN_ALERT_SECRET")
    assert c.post(API + "/scheduled", headers={"X-Domain-Secret": ""}).status_code == 403          # unset means refused, even to an empty secret
    assert world.store.state_get("runs") is None


def test_the_scheduled_scan_runs_without_a_session_and_is_recorded(world):
    r = client_as(None).post(API + "/scheduled", headers={"X-Domain-Secret": "s3cret"})
    assert r.status_code == 202
    run = world.store.state_get("runs")[0]
    assert run["trigger"] == "schedule" and run["read"] == 2 and run["new"] == 2          # the database is empty, so both accounts' ads are unknown


def test_a_scan_that_is_running_is_not_started_again(world):
    from datetime import datetime, timezone
    world.store.state_set("progress", {"running": True, "startedAt": datetime.now(timezone.utc).isoformat(), "done": 1, "total": 2})
    assert post("/scan").status_code == 409
    r = client_as(None).post(API + "/scheduled", headers={"X-Domain-Secret": "s3cret"})
    assert r.status_code == 200 and r.get_json()["started"] is False and world.store.state_get("runs") is None
    world.store.state_set("progress", {"running": True, "startedAt": "2020-01-01T00:00:00+00:00", "done": 1, "total": 2})     # a scan that died long ago
    assert post("/scan").status_code == 202


def test_scan_now_scans_and_reports_the_user(world):
    assert post("/scan").status_code == 202
    assert world.store.state_get("runs")[0]["trigger"] == "page by " + STAFF
    assert world.store.state_get("progress")["running"] is False


# ── The database ────────────────────────────────────────────────────────────

def test_adding_listing_and_removing_a_domain(world):
    r = post("/known", {"domain": "https://www.Acme-Garden.com/", "scope": "mcc-1", "note": "site"})
    assert r.status_code == 201 and r.get_json()["entry"]["domain"] == "acme-garden.com"
    assert post("/known", {"domain": "acme-garden.com", "scope": "mcc-1"}).status_code == 200            # already there
    assert post("/known", {"domain": "not a domain", "scope": ""}).status_code == 400
    assert post("/known", {"domain": "x.example", "scope": "zzz"}).status_code == 400
    entry_id = world.store.known[0]["id"]
    assert client_as(STAFF).get(API + "/state").get_json()["known"][0]["added_by"] == STAFF
    assert client_as(STAFF).delete(API + "/known/%d" % entry_id, json={}, headers=SAME).status_code == 200
    assert client_as(STAFF).delete(API + "/known/%d" % entry_id, json={}, headers=SAME).status_code == 404


def test_importing_a_csv_by_name_and_listing_the_bad_rows(world):
    text = "Domain,Client,Notes\nacme-garden.com,Acme,site\nbolt-dental.com,Bolt Dental,\nlinktree.example,,everyone\nbad one,,\nzed.example,Nobody,\n"
    r = post("/known/import", {"text": text})
    assert r.status_code == 200
    j = r.get_json()
    assert j["added"] == 3 and j["skippedCount"] == 2 and [s["row"] for s in j["skipped"]] == [4, 5]
    got = {(k["domain"], k["scope"]) for k in world.store.known}
    assert got == {("acme-garden.com", "mcc-1"), ("bolt-dental.com", B["id"]), ("linktree.example", "*")}
    assert post("/known/import", {"text": ""}).status_code == 400
    assert post("/known/import", {"sheetUrl": "https://example.com/x"}).status_code == 400


def test_the_export_is_csv_and_a_formula_cell_is_defused(world):
    world.store.known_add("acme-garden.com", "mcc-1", "=HYPERLINK(1)", "me")
    r = client_as(STAFF).get(API + "/known.csv")
    assert r.status_code == 200 and r.mimetype == "text/csv" and "attachment" in r.headers["Content-Disposition"]
    body = r.get_data(as_text=True)
    assert body.splitlines()[0] == "domain,for,for_id,note,added_by"
    assert "acme-garden.com,Acme" in body and "'=HYPERLINK(1)" in body


# ── Findings, baseline, settings ────────────────────────────────────────────

def test_approving_dismissing_and_reopening(world):
    post("/scan")
    f = next(x for x in world.store.found.values() if x["accountId"] == B["id"])
    fp = f["fingerprint"]
    assert post("/findings/%s/dismiss" % fp).status_code == 200 and world.store.found[fp]["status"] == "dismissed"
    assert post("/findings/%s/reopen" % fp).status_code == 200 and world.store.found[fp]["status"] == "open"
    r = post("/findings/%s/approve" % fp, {"who": "account"})
    assert r.status_code == 200 and r.get_json()["approved"] == 1 and world.store.found[fp]["status"] == "approved"
    assert post("/findings/%s/approve" % fp, {"who": "nobody"}).status_code == 400
    assert post("/findings/nothex/dismiss").status_code == 404
    assert post("/findings/%s/delete" % fp).status_code == 404
    assert post("/findings/%s/dismiss" % ("0" * 24)).status_code == 400


def test_the_baseline_needs_a_scan_and_then_starts_the_alerts(world):
    assert post("/baseline", {"trust": "all"}).status_code == 400
    post("/scan")
    assert world.deliver.emails == []                                          # no baseline, nothing sent
    r = post("/baseline", {"trust": [A["id"]]})
    assert r.status_code == 200 and r.get_json()["trusted"] == 1
    assert post("/baseline", {"trust": "nonsense"}).status_code == 400
    world.ads[A["id"]].append("https://new-and-odd.example/")
    post("/scan")
    assert len(world.deliver.emails) == 1 and "new-and-odd.example" in world.deliver.emails[0]["text"]
    assert "p2/paid-marketing/unknown-domains" in world.deliver.emails[0]["text"]


def test_settings_are_validated_and_saved(world):
    r = post("/settings", {"emails": "a@x.com, b@y.org", "slackChannel": "#sec"})
    assert r.status_code == 200 and world.store.state_get("settings") == {"emails": ["a@x.com", "b@y.org"], "slackChannel": "#sec"}
    assert post("/settings", {"emails": "nope"}).status_code == 400
    assert post("/settings", {"emails": "", "slackChannel": "two words"}).status_code == 400


def test_a_test_alert_goes_through_the_senders(world):
    r = post("/test-alert")
    assert r.status_code == 200 and r.get_json()["sent"] is True and world.deliver.emails[0]["subject"].startswith("TEST")


def test_an_unavailable_store_is_a_503(monkeypatch):
    class Down(MemStore):
        def known_list(self):
            raise self.StoreUnavailable("no database")

        def state_get(self, key, default=None):
            raise self.StoreUnavailable("no database")
    monkeypatch.setattr(appmod, "_dom_store", lambda: Down())
    monkeypatch.setattr(appmod, "_dom_inventory", lambda s: ([A], GROUPS, []))
    r = client_as(STAFF).get(API + "/state")
    assert r.status_code == 503 and "unavailable" in r.get_json()["error"]


def test_the_url_checker_state_dismiss_reopen_and_the_sheet(world, monkeypatch):
    monkeypatch.setattr(pl, "default_fetch", lambda: Site({"https://strange.example/": 404, "https://acme-garden.com/": "https://elsewhere.example/"}))
    monkeypatch.setattr(pl.time, "sleep", lambda s: None)
    assert post("/scan").status_code == 202
    j = client_as(STAFF).get(API + "/state").get_json()
    assert j["linkCounts"]["open"] == 2 and {f["kind"] for f in j["links"]} == {"broken", "redirect_domain"}
    broken = next(f for f in j["links"] if f["kind"] == "broken")
    assert post("/links/%s/dismiss" % broken["fingerprint"]).status_code == 200 and world.store.links[broken["fingerprint"]]["status"] == "dismissed"
    assert post("/links/%s/reopen" % broken["fingerprint"]).status_code == 200 and world.store.links[broken["fingerprint"]]["status"] == "open"
    assert post("/links/nothex/dismiss").status_code == 404
    assert post("/links/%s/delete" % broken["fingerprint"]).status_code == 404
    assert post("/links/%s/dismiss" % ("0" * 24)).status_code == 400
    r = client_as(STAFF).get(API + "/links.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.mimetype and r.data[:2] == b"PK" and "broken-links.xlsx" in r.headers["Content-Disposition"]
    assert client_as(None).get(API + "/links.xlsx").status_code == 302
    assert post("/links/%s/dismiss" % broken["fingerprint"], headers={"Origin": "https://evil.example"}).status_code == 403


def test_the_download_is_the_whole_record_as_one_workbook(world, monkeypatch):
    """Every finding of every status, the domain database, what each account uses, the broken links and the scans, whatever the page is showing."""
    import io
    from openpyxl import load_workbook
    monkeypatch.setattr(pl, "default_fetch", lambda: Site({"https://strange.example/": 404}))
    monkeypatch.setattr(pl.time, "sleep", lambda s: None)
    post("/scan")
    fps = [f["fingerprint"] for f in world.store.found.values()]
    post("/findings/%s/dismiss" % fps[0])
    world.store.found[fps[1]]["status"] = "gone"
    post("/known", {"domain": "acme-garden.com", "scope": "mcc-1", "note": "main site"})
    assert client_as(None).get(API + "/export.xlsx").status_code == 302
    r = client_as(STAFF).get(API + "/export.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.mimetype and "unknown-domains-export.xlsx" in r.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(r.data))
    assert wb.sheetnames == ["Findings", "Domain database", "In use today", "Links to fix", "Scans"]
    findings = [[c.value for c in row] for row in wb["Findings"].iter_rows()]
    assert findings[0][:4] == ["Status", "Severity", "Kind", "Account"] 
    assert {"Dismissed", "Gone"} <= {row[0] for row in findings[1:]}                              # the dump holds what the page's default filter hides
    db = [[c.value for c in row] for row in wb["Domain database"].iter_rows()]
    assert db[1][:2] == ["acme-garden.com", "Acme"] and db[1][3] == "main site"
    use = [[c.value for c in row] for row in wb["In use today"].iter_rows()]
    assert use[0][3] == "Domain its ads point at" and {row[3] for row in use[1:]} >= {"acme-garden.com", "strange.example"}
    assert dict((row[3], row[5]) for row in use[1:])["acme-garden.com"] == "Yes" and dict((row[3], row[5]) for row in use[1:])["strange.example"] == "No"
    links = [[c.value for c in row] for row in wb["Links to fix"].iter_rows()]
    assert links[1][1] == "Broken" and links[1][8] == "https://strange.example/"
    scans = [[c.value for c in row] for row in wb["Scans"].iter_rows()]
    assert scans[0][1] == "Started by" and scans[1][1] == STAFF
    assert client_as(STAFF).get(API + "/export.xlsx").status_code == 200


def test_the_download_of_an_empty_record_is_a_workbook_with_headings(monkeypatch):
    import io
    from openpyxl import load_workbook
    from tracker import paid_domains_export
    wb = load_workbook(io.BytesIO(paid_domains_export.workbook(MemStore(), [A], GROUPS)))
    assert wb.sheetnames == ["Findings", "Domain database", "In use today", "Links to fix", "Scans"] and all(ws.max_row == 1 for ws in wb)
