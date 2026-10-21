"""Cross-Channel & Attribution Analysis: the downloads."""
import csv
import io
import zipfile

import pytest

from tests import test_paid_xchannel as t
from tracker import paid_xchannel_export as ex


@pytest.fixture(autouse=True)
def fresh(monkeypatch):
    t.paid_gads.forget()
    monkeypatch.setattr(t.paid_segments, "_meta_target", lambda act, store: dict(t.META_ACCT, id=act))
    yield
    t.paid_gads.forget()


def full(**kw):
    kw = dict(t.settings_fns(), actual_conv="100", actual_value="9000", preset="90", **kw)
    return t.run(**kw)


def test_the_sheets_are_the_ones_the_report_has():
    tabs = ex.tables(full())
    assert list(tabs) == ["Summary", "Platforms", "Campaign types", "Cost change", "Weekly", "Weekly links", "Attribution", "Reconciliation", "Findings", "Notes"]


def test_a_report_without_a_split_or_reconciliation_leaves_those_sheets_out():
    tabs = ex.tables(t.run(account=t.GID))
    assert "Cost change" not in tabs and "Reconciliation" not in tabs and "Weekly links" not in tabs
    assert {"Summary", "Platforms", "Campaign types", "Weekly", "Findings", "Notes"} <= set(tabs)


def test_the_platform_rows_carry_the_figures_the_page_shows_and_a_blended_row():
    h, rows = ex.tables(t.run())["Platforms"]
    assert [r[0] for r in rows] == ["Google Ads", "Meta Ads", "Blended (the platforms' claims added)"]
    g = dict(zip(h, rows[0]))
    assert g["Spend"] == 1800.0 and g["Conversions"] == 90.0 and g["Cost per conversion"] == 20.0 and g["Conversion measure"] == "Conversions" and g["Share of conversions %"] == pytest.approx(66.7, abs=0.1)
    b = dict(zip(h, rows[2]))
    assert b["Spend"] == 3300.0 and b["Share of spend %"] == 100.0


def test_one_platform_has_no_blended_row():
    _h, rows = ex.tables(t.run(account=t.GID))["Platforms"]
    assert len(rows) == 1


def test_the_cost_change_sheet_adds_up():
    def meta(ci, d):
        return [(2000, 100, 15.0 if d < "2026-09-18" else 30.0, 1.5, 0.0, 2, 0.0, 0), (8000, 40, 20.0, 0, 0.0, 0, 0.0, 0)][ci]
    _h, rows = ex.tables(t.run(shape_fn=t.shape_fn(meta=meta)))["Cost change"]
    vals = {r[0]: r[1] for r in rows}
    assert vals["Change"] == pytest.approx(vals["Of which the mix of conversions between platforms"] + vals["Of which Google Ads' own cost per conversion"] + vals["Of which Meta Ads' own cost per conversion"], abs=0.02)


def test_the_attribution_sheet_lists_the_actions_the_windows_and_the_notes():
    _h, rows = ex.tables(full())["Attribution"]
    kinds = [(r[0], r[1]) for r in rows]
    assert ("Google Ads", "Counted conversion action") in kinds and ("Meta Ads", "Attribution window") in kinds and any(r[0] == "Both" and r[1] == "High" for r in rows)
    meta = next(r for r in rows if r[1] == "Attribution window")
    assert meta[2] == "7 day click, 1 day view" and "100.0 of spend in 1 ad set" in meta[3]


def test_the_reconciliation_sheet_names_what_was_scaled():
    _h, rows = ex.tables(full())["Reconciliation"]
    labels = [r[0] for r in rows]
    assert "Conversions your system recorded" in labels and "Google Ads cost per conversion scaled to what was recorded" in labels and "Return on the revenue recorded" in labels


def test_the_weekly_links_sheet_has_the_four_tests():
    _h, rows = ex.tables(full())["Weekly links"]
    assert len(rows) == 4 and {r[2] for r in rows} == {"in the same week", "a week later"}


def test_the_findings_sheet_has_a_level_a_detail_and_what_to_do():
    h, rows = ex.tables(full())["Findings"]
    assert h == ["Level", "Finding", "Detail", "Figures", "What to do", "How to do it", "What it is worth"] and rows[0][0] == "High" and any(r[4] for r in rows)


def test_a_finding_with_a_plan_carries_its_steps_and_worth_into_the_sheet_and_one_without_leaves_them_empty():
    h, rows = ex.tables(full())["Findings"]
    by = {r[1]: r for r in rows}
    gap = next(r for r in rows if "costs" in r[1] and "per conversion" in r[1])
    assert gap[5].startswith("1. In Ads Manager open the campaign's ad sets") and "\n2. In Google Ads open Campaigns" in gap[5] and "upper bound" in gap[6]
    view = next(r for r in rows if "saw the ad" in r[1])
    assert "Compare attribution settings" in view[5] and view[6].startswith("There is no figure for how many of Meta's conversions were view-through")          # a note rides in the worth column
    idle = next(r for r in rows if "recorded no conversions" in r[1])
    assert idle[5] == "" and idle[6] == ""


def test_the_excel_file_has_a_sheet_per_table_with_numbers_kept_as_numbers():
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(ex.xlsx(full())))
    assert wb.sheetnames == list(ex.tables(full()))
    ws = wb["Platforms"]
    assert ws["A2"].value == "Google Ads" and ws["C2"].value == 5400.0 and ws.freeze_panes == "A2"


def test_a_cell_that_starts_like_a_formula_is_text_in_excel_and_quoted_in_csv():
    from openpyxl import load_workbook
    report = full()
    report["findings"][0]["title"] = "=HYPERLINK(\"http://x\",\"y\")"
    wb = load_workbook(io.BytesIO(ex.xlsx(report)))
    c = wb["Findings"]["B2"]
    assert c.data_type == "s" and c.value.startswith("=")
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(report)))
    row = list(csv.reader(io.StringIO(z.read("findings.csv").decode())))[1]
    assert row[1].startswith("'=")


def test_the_zip_holds_one_csv_per_table():
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(full())))
    assert sorted(z.namelist()) == sorted(["summary.csv", "platforms.csv", "campaign-types.csv", "cost-change.csv", "weekly.csv", "weekly-links.csv", "attribution.csv", "reconciliation.csv", "findings.csv", "notes.csv"])
    rows = list(csv.reader(io.StringIO(z.read("platforms.csv").decode())))
    assert rows[0][0] == "Platform" and rows[1][0] == "Google Ads"


def test_the_pdf_is_a_pdf_with_the_main_sections():
    body = ex.pdf(full())
    assert body[:5] == b"%PDF-" and len(body) > 3000
    assert ex.pdf(t.run(account=t.GID))[:5] == b"%PDF-"


def test_the_pdf_survives_a_very_long_finding_and_markup_in_text():
    report = full()
    report["findings"][0]["text"] = "<b>bold</b> & <script>x</script> " + "long " * 2000
    assert ex.pdf(report)[:5] == b"%PDF-"


def test_the_file_name_is_made_from_the_account():
    r = full()
    assert ex.file_stem(r) == "Acme-Dental-Acme-Dental-Meta-cross-channel-and-attribution"
    assert ex.file_stem({"account": {"name": "", "id": ""}}) == "account-cross-channel-and-attribution"
    assert "/" not in ex.file_stem({"account": {"name": "../../etc/passwd"}})


def test_a_high_finding_is_shaded_in_excel():
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(ex.xlsx(full())))["Findings"]
    assert ws["A2"].value == "High" and ws["A2"].fill.fgColor.rgb.endswith("FFD9D6")


def test_returns_are_kept_to_three_decimals():
    rows = {r[0]: r[1] for r in ex.tables(full())["Reconciliation"][1]}
    assert rows["Google Ads return scaled to what was recorded"] == 1.667
