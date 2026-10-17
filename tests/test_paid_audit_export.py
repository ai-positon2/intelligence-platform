"""Account Audit: the downloads (tracker/paid_audit_export.py): every file is drawn from the same tables, opens, and cannot run a formula."""
import csv
import io
import zipfile

import openpyxl
import pytest

from tests import audit_fakes as A
from tests import audit_meta_fakes as F
from tracker import paid_audit as audit
from tracker import paid_audit_export as ex
from tracker import paid_audit_google_read as gread

NOW = audit.dt.datetime(2026, 10, 17, 12, 0, tzinfo=audit.dt.timezone.utc)


def make_report(**kw):
    google = A.a_read(answers={gread.CAMPAIGN_METRICS_GAQL: [A.campaign_metrics("10", 5000, 400, 1000.0, 0.0)], gread.TERMS_GAQL: [A.term("=cmd|' /C calc'!A0", 600.0, 60, 0.0)]})
    google["pages"] = {"checked": 1, "of": 1, "results": [{"url": "https://acme.test/x", "status": 404, "error": "", "ms": 120, "blocked": False, "finalUrl": "", "redirected": False}]}
    google["readAt"] = NOW.isoformat()
    meta = F.a_read()
    accounts = [{"platform": "google", "id": "1112223334", "name": "Acme Dental", "currency": "USD"}, {"platform": "meta", "id": F.ACT, "name": "Acme Meta", "currency": "USD"}]
    return audit.assemble(google, meta, accounts, now=NOW)


@pytest.fixture(scope="module")
def report():
    return make_report()


def test_the_tables_come_in_the_order_they_are_shown(report):
    assert list(ex.tables(report)) == ["Summary", "Fix first", "All checks", "Could not check", "Notes"]


def test_the_summary_says_who_what_and_how_well(report):
    header, rows = ex.tables(report)["Summary"]
    d = dict((r[0], r[1]) for r in rows)
    assert header == ["Item", "Detail"]
    assert d["Google Ads account"] == "Acme Dental (1112223334)" and d["Meta Ads account"] == "Acme Meta (%s)" % F.ACT
    assert d["Audit made on"] == "2026-10-17" and d["Read on"] == "2026-10-17 12:00"
    assert d["Overall score"] == report["summary"]["score"] and d["Overall grade"] == report["summary"]["grade"]
    assert d["Checks run"] == len(ex.tables(report)["All checks"][1]) and d["Failed critical checks"] == report["summary"]["critical"]
    assert any(k.startswith("Google Ads: Conversion tracking") for k in d) and any(k.startswith("Meta Ads: Pixel and tracking") for k in d)
    assert "%d%%" % report["summary"]["unknownShare"] == d["Share of the checks that could not be judged"]


def test_every_row_is_as_wide_as_its_header_and_numbers_stay_numbers(report):
    for name, (header, rows) in ex.tables(report).items():
        assert rows, name
        for r in rows:
            assert len(r) == len(header), (name, r)
    order = [r[0] for r in ex.tables(report)["Fix first"][1]]
    assert order == list(range(1, len(order) + 1))


def test_fix_first_lists_the_worst_findings_with_what_to_do(report):
    header, rows = ex.tables(report)["Fix first"]
    first = dict(zip(header, rows[0]))
    assert first["Check"].startswith("G-CNV-02") and first["Severity"] == "Critical" and first["Result"] == "Fail" and first["Platform"] == "Google Ads"
    assert "Check the tag" in first["What to do"] and "USD" in first["What was found"]


def test_all_checks_has_every_finding_with_why_and_fix(report):
    header, rows = ex.tables(report)["All checks"]
    assert len(rows) == report["summary"]["tally"]["total"]
    row = next(dict(zip(header, r)) for r in rows if r[2].startswith("G-ADS-09"))
    assert row["Result"] == "Fail" and "https://acme.test/x (HTTP 404)" in row["Affected"] and row["Why it matters"] and row["What to do"]
    assert {r[0] for r in rows} == {"Google Ads", "Meta Ads"}


def test_affected_things_are_joined_and_say_how_many_more():
    f = {"items": ["a", "b"], "count": 2}
    assert ex._affected(f) == "a; b" and ex._affected({"items": [], "count": 0}) == ""
    assert ex._affected({"items": ["a", "b"], "count": 5}) == "a; b (and 3 more)"
    many = {"items": ["a", "b", "c", "d", "e", "f", "g"], "count": 12}
    assert ex._affected(many, 5) == "a; b; c; d; e (and 7 more)" and ex._affected(many) == "a; b; c; d; e; f; g (and 5 more)" and ex._affected(many, 20) == ex._affected(many)


def test_what_could_not_be_checked_is_listed_with_why_not(report):
    header, rows = ex.tables(report)["Could not check"]
    assert header == ["Platform", "Check", "Why not"]
    assert any(r[1].startswith("G-CNV-11") and "does not show the GA4 link" in r[2] for r in rows)
    assert all(r[2].startswith("Could not be checked") for r in rows)


def test_notes_carry_the_caveats_the_apis_force_and_the_parts_not_read(report):
    rows = ex.tables(report)["Notes"][1]
    assert any("account-level negative" in r[1] for r in rows) and any("Conversions API" in r[1] for r in rows)
    bad = make_report()
    bad["warnings"] = ["The keywords could not be read (Google Ads answered 500)."]
    bad["stale"] = [{"platform": "google", "savedAt": "2026-10-16T08:30:00+00:00"}]
    rows = ex.tables(bad)["Notes"][1]
    assert ["Part that could not be read", "The keywords could not be read (Google Ads answered 500)."] in rows
    assert ["Saved read", "Google Ads was read on 2026-10-16 08:30 because the platform would not answer now."] in rows


def test_a_report_with_nothing_in_it_still_makes_every_file():
    empty = audit.assemble(None, None, [], now=NOW)
    assert list(ex.tables(empty)) == ["Summary"]
    assert ex.xlsx(empty)[:2] == b"PK" and ex.csv_zip(empty)[:2] == b"PK" and ex.pdf(empty)[:5] == b"%PDF-"
    assert dict(ex.tables(empty)["Summary"][1])["Overall score"] == "Nothing could be judged"


def test_the_workbook_has_a_sheet_per_table_with_results_coloured_and_text_as_text(report):
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(report)))
    assert wb.sheetnames == list(ex.tables(report))
    ws = wb["All checks"]
    header = [c.value for c in ws[1]]
    result_col = header.index("Result") + 1
    fills = {ws.cell(row=r, column=result_col).value: ws.cell(row=r, column=result_col).fill.fgColor.rgb for r in range(2, ws.max_row + 1)}
    assert fills["Fail"].endswith("FFD9D6") and fills["Pass"].endswith("DDF3E2") and fills["Could not check"].endswith("E4E8F0")
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref
    for row in ws.iter_rows(min_row=2):
        for c in row:
            assert not (isinstance(c.value, str) and c.value.startswith("=")) or c.data_type == "s"
    assert wb["Fix first"].cell(row=2, column=1).value == 1


def test_a_value_a_spreadsheet_would_run_never_runs(report):
    bad = make_report()
    first = bad["sections"][0]["findings"][0]
    first["found"] = "=HYPERLINK(\"http://evil.test\",\"x\")"
    first["items"] = ["+cmd|' /C calc'!A0"]
    first["count"] = 1
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(bad)))
    cell = next(c for row in wb["All checks"].iter_rows(min_row=2) for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value)
    assert cell.data_type == "s"
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(bad)))
    rows = list(csv.reader(io.StringIO(z.read("all-checks.csv").decode())))
    cells = [c for r in rows for c in r]
    assert "'=HYPERLINK(\"http://evil.test\",\"x\")" in cells and "'+cmd|' /C calc'!A0" in cells
    assert not any(c.startswith(("=", "+", "@")) for c in cells)


def test_the_csv_zip_has_one_file_per_table(report):
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(report)))
    assert z.namelist() == ["summary.csv", "fix-first.csv", "all-checks.csv", "could-not-check.csv", "notes.csv"]
    rows = list(csv.reader(io.StringIO(z.read("fix-first.csv").decode())))
    assert rows[0] == ex.FIX_HEADER and rows[1][0] == "1" and len(rows) == len(ex.tables(report)["Fix first"][1]) + 1


def test_the_pdf_is_a_document_with_the_summary_the_fixes_and_every_section(report):
    pdf = ex.pdf(report)
    assert pdf[:5] == b"%PDF-" and len(pdf) > 5000
    secs = ex.pdf_sections(report)
    assert [s[0] for s in secs][0] == "Google Ads: Account and settings" and [s[0] for s in secs][-1] == "Meta Ads: Ads"
    label, score, header, rows = next(s for s in secs if s[0].endswith("Ads and landing pages"))
    assert header == ["Result", "Severity", "Check", "What was found", "What to do"]
    page = next(r for r in rows if r[2].startswith("G-ADS-09"))
    assert page[0] == "Fail" and "Affected: https://acme.test/x (HTTP 404)" in page[3] and page[4]
    long_ones = ex.pdf_sections(make_long())
    finding = next(r for lbl, _s, _h, rows_ in long_ones for r in rows_ if r[2].startswith("G-KWD-07"))
    assert finding[3].endswith("(and 15 more)") and finding[3].count(";") == 4                     # the PDF names five, the workbook all the read kept
    passed = next(r for r in rows if r[0] == "Pass")
    assert passed[4] == "" and "Affected" not in passed[3]                        # passing checks are one line, with nothing to do


def make_long():
    r = make_report()
    f = next(f for s in r["sections"] for f in s["findings"] if f["id"] == "G-KWD-07")
    f["result"], f["items"], f["count"] = "fail", ["term %d" % i for i in range(8)], 20
    return r


def test_the_fix_first_page_in_the_pdf_names_five_affected_things_at_most():
    r = make_long()
    r["fixFirst"] = [next(f for s in r["sections"] for f in s["findings"] if f["id"] == "G-KWD-07")]
    assert ex.pdf(r)[:5] == b"%PDF-"


def test_the_file_stem_names_the_accounts_safely(report):
    assert ex.file_stem(report) == "Acme-Dental-Acme-Meta-account-audit"
    assert ex.file_stem({"accounts": []}) == "account-account-audit"
    assert ex.file_stem({"accounts": [{"name": "../../etc/passwd  &x", "id": "1"}]}) == "etc-passwd-x-account-audit"


# ── Cases the mutation round showed were not pinned ─────────────────────────

def test_the_file_stem_falls_back_to_the_id_and_is_cut_at_sixty_characters():
    assert ex.file_stem({"accounts": [{"id": "1112223334"}]}) == "1112223334-account-audit" and ex.file_stem({"accounts": [{"name": "", "id": "act_9"}]}) == "act-9-account-audit"
    long_name = "x" * 100
    assert ex.file_stem({"accounts": [{"name": long_name, "id": "1"}]}) == "x" * 60 + "-account-audit"
    assert ex.file_stem({"accounts": [{"name": "A" * 59 + "-" + "B" * 5, "id": "1"}]}) == "A" * 59 + "-account-audit"          # a dash left at the cut is trimmed


def test_one_more_is_said_when_exactly_one_more_was_kept_back():
    assert ex._affected({"items": ["a", "b"], "count": 3}) == "a; b (and 1 more)"


def test_a_platform_or_section_that_could_not_be_judged_says_so(report):
    mixed = make_report()
    mixed["summary"]["platforms"]["meta"]["score"] = None
    mixed["summary"]["platforms"]["meta"]["grade"] = None
    mixed["sections"][0]["score"] = None
    mixed["sections"][0]["grade"] = None
    d = dict((r[0], r[1]) for r in ex.tables(mixed)["Summary"][1])
    assert d["Meta Ads score"] == "Nothing could be judged" and d["Google Ads: Account and settings"] == "Nothing could be judged"
    assert ex.pdf_sections(mixed)[0][1] == "nothing could be judged" and ex.pdf(mixed)[:5] == b"%PDF-"


def test_the_workbook_columns_are_as_wide_as_their_text_up_to_sixty():
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(make_report())))
    ws = wb["All checks"]
    header = [c.value for c in ws[1]]
    assert ws.column_dimensions[openpyxl.utils.get_column_letter(header.index("Severity") + 1)].width == 10          # a short column has a floor
    assert ws.column_dimensions[openpyxl.utils.get_column_letter(header.index("What to do") + 1)].width == 60       # a long one a ceiling
    assert ws.column_dimensions[openpyxl.utils.get_column_letter(header.index("Result") + 1)].width >= 10


def test_a_passing_check_names_no_affected_things_in_the_pdf_even_when_the_read_kept_some():
    r = make_report()
    f = next(f for s in r["sections"] for f in s["findings"] if f["result"] == "pass")
    f["items"], f["count"] = ["a campaign name"], 1
    rows = [row for _l, _s, _h, rows_ in ex.pdf_sections(r) for row in rows_ if row[2].startswith(f["id"])]
    assert rows and rows[0][0] == "Pass" and "Affected" not in rows[0][3] and "a campaign name" not in rows[0][3]


def test_the_summary_says_why_a_grade_is_held_and_says_nothing_when_it_is_not(report):
    held = make_report()
    held["summary"]["capReasons"] = ["a critical check failed", "40% of the checks could not be read"]
    d = dict((r[0], r[1]) for r in ex.tables(held)["Summary"][1])
    assert d["Grade held at C or below because"] == "a critical check failed; 40% of the checks could not be read"
    clear = make_report()
    clear["summary"]["capReasons"] = []
    assert "Grade held at C or below because" not in dict((r[0], r[1]) for r in ex.tables(clear)["Summary"][1])
