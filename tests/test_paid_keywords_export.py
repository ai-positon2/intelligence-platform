"""Keyword Research: the downloads (tracker/paid_keywords_export.py): every file is drawn from the same tables, opens, and cannot run a formula."""
import csv
import io
import json
import zipfile

import openpyxl
import pytest

from tests import keyword_fakes as K, strategy_fakes as F
from tracker import paid_keywords as pk, paid_keywords_export as ex, paid_keywords_job as job, paid_strategy as ps


@pytest.fixture(scope="module")
def rec():
    store = F.FakeStrategyStore()
    inp = K.clean()
    r = store.create(by="ana@position2.com", inputs=inp, account_name="Acme Dental", steps=pk.steps_for(inp))
    store.begin(r["id"])
    job.run_build(r["id"], store=store, deps=K.deps())
    return json.loads(json.dumps(store.get(r["id"])))


def test_the_tables_come_in_the_order_they_are_shown(rec):
    assert list(ex.tables(rec)) == ["Overview", "Themes", "Keywords", "Add to account", "Already in the account", "Converting search terms", "Competitor keywords", "Negatives to add",
                                    "Seasonality", "Plan", "Questions for the client", "Ads Editor import", "Checks"]


def test_every_row_is_as_wide_as_its_header_and_text_is_text_and_numbers_are_numbers(rec):
    for name, (header, rows) in ex.tables(rec).items():
        assert rows, name
        for r in rows:
            assert len(r) == len(header), (name, r)
    head, rows = ex.tables(rec)["Keywords"]
    i = head.index("Monthly searches")
    assert all(isinstance(r[i], int) or r[i] == "" for r in rows) and any(isinstance(r[i], int) for r in rows)


def test_the_overview_gives_the_client_the_counts_and_the_short_read(rec):
    rows = dict(ex.tables(rec)["Overview"][1][:17])
    c = rec["plan"]["coverage"]
    assert rows["Client"] == "Acme Dental" and rows["Keywords judged"] == c["judged"] and rows["Judged a fit"] == c["yes"] and rows["Fit and new to the account"] == c["new"]
    assert rows["Already covered by the account, of the keywords judged a fit"] == "%d%%" % c["pct"] and rows["Smallest monthly searches looked at"] == 10
    assert rows["Already covered by the account, of the searches for ideas that share a word with what the client sells"] == "%d%%" % c["poolPct"] and c["poolPct"] != c["pct"] and rows["Built on"] == rec["plan"]["generatedAt"][:10]
    assert [r[1] for r in ex.tables(rec)["Overview"][1] if r[0] == "In short"] == K.PLAN["summary"]


def test_the_keywords_sheet_has_every_keyword_judged_a_fit_first_and_says_what_the_account_has(rec):
    head, rows = ex.tables(rec)["Keywords"]
    assert len(rows) == len(rec["plan"]["keywords"]) and head[0] == "Keyword" and head[-1] == "In the Ads Editor file"
    fits = [r[1] for r in rows]
    assert fits == sorted(fits, key=lambda f: {"yes": 0, "maybe": 1, "no": 2, "unjudged": 3}[f])
    by = {r[0]: r for r in rows}
    assert by["dental implants austin"][head.index("In the account")] == "Already a keyword (exact match in P2_Search_Implants_20250101)"
    assert by["cost of dental implants"][head.index("In the account")] == "Matched by \"dental implants\" (phrase) in P2_Search_Implants_20250101"
    assert by["emergency dentist near me"][head.index("In the account")] == "New to the account" and by["emergency dentist near me"][-1] == "Yes" and by["invisalign cost"][-1] == "No"
    assert by["dentist dallas"][head.index("In the account")] == "New to the account. Blocked in some campaigns by \"dallas\""
    assert by["emergency dentist near me"][head.index("Low bid (USD)")] == 7.0 and by["invisalign cost"][head.index("How it moves through the year")].startswith("Peaks in May")


def test_add_to_account_is_the_included_keywords_most_important_first_with_their_campaign_and_ad_group(rec):
    head, rows = ex.tables(rec)["Add to account"]
    assert len(rows) == rec["plan"]["coverage"]["include"] == 4 and head[:6] == ["Priority", "Keyword", "Match", "Theme", "Campaign", "Ad group"]
    order = {"high": 0, "medium": 1, "low": 2}
    assert [order[r[0]] for r in rows] == sorted(order[r[0]] for r in rows)
    r = next(x for x in rows if x[1] == "emergency dentist near me")
    assert r[2] == "phrase" and r[3] == "Emergency Dentist" and r[4] == "P2_Search_EmergencyDentist" and r[5] == "Emergency Dentist | Local" and r[6] == 9900


def test_the_already_covered_sheet_says_what_covers_each_keyword(rec):
    head, rows = ex.tables(rec)["Already in the account"]
    by = {r[0]: r for r in rows}
    assert by["cost of dental implants"][2:6] == ["Matched by a broader keyword", "dental implants", "Phrase", "P2_Search_Implants_20250101"] and by["dental implants austin"][2] == "Already a keyword"
    assert "emergency dentist near me" not in by and all(isinstance(r[1], int) for r in rows)


def test_converting_search_terms_carry_their_clicks_cost_and_conversions(rec):
    head, rows = ex.tables(rec)["Converting search terms"]
    by = {r[0]: r for r in rows}
    assert by["all on 4 implants austin tx"][1:5] == [4.0, 44, 380.0, "P2_Search_Implants_20250101"] and [r[1] for r in rows] == sorted([r[1] for r in rows], reverse=True)


def test_competitor_keywords_are_listed_with_the_decision_that_belongs_to_the_client(rec):
    head, rows = ex.tables(rec)["Competitor keywords"]
    assert [r[0] for r in rows] == ["smile austin dentistry", "smile austin veneers"] and rows[0][1] == "smileaustin.com" and "client's decision" in rows[0][4]


def test_the_seasonality_sheet_has_a_column_for_each_month(rec):
    head, rows = ex.tables(rec)["Seasonality"]
    assert head[0] == "Theme" and head[1] == "Jan 2026" and head[12] == "Dec 2026" and head[-1] == "How it moves" and all(len(r) == 14 for r in rows)


def test_the_checks_sheet_names_what_was_left_out_and_why(rec):
    rows = ex.tables(rec)["Checks"][1]
    kinds = [r[0] for r in rows]
    assert "Negative not suggested" in kinds and "Ideas below the smallest search volume" in kinds and "Ideas blocked in every campaign by the account's negatives" in kinds and "Assumption" in kinds
    assert any(r[0] == "Negative not suggested" and r[1].startswith("dental implants (phrase): It would block the keyword") for r in rows)
    assert dict((r[0], r[1]) for r in rows)["Ideas merged with a close variant"] == "1" and "Keywords blocked by a negative in only some campaigns" in kinds


def test_the_editor_file_is_the_keywords_to_add_paused(rec):
    body = ex.csv_editor(rec)
    rows = list(csv.reader(io.StringIO(body)))
    assert rows[0] == ["Campaign", "Ad group", "Keyword", "Criterion Type", "Status"] and len(rows) == 5 and all(r[4] == "Paused" for r in rows[1:])
    assert {r[2] for r in rows[1:]} == {"emergency dentist near me", "implant dentist near me", "all on 4 implants austin tx", "acme dental austin"}
    assert not any(r[2] in ("smile austin dentistry", "dentist dallas") for r in rows[1:]) and ex.tables(rec)["Ads Editor import"][1] == rec["plan"]["editor"]


def test_there_is_no_editor_file_when_there_is_nothing_to_add(rec):
    empty = json.loads(json.dumps(rec))
    empty["plan"]["editor"] = []
    assert ex.csv_editor(empty) is None and "Ads Editor import" not in ex.tables(empty) and ex.csv_editor({}) is None


def test_the_excel_file_opens_with_a_sheet_for_each_table_and_numbers_as_numbers(rec):
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(rec)))
    assert wb.sheetnames == [n[:31] for n in ex.tables(rec)] and wb["Keywords"].freeze_panes == "A2" and wb["Keywords"].auto_filter.ref
    ws = wb["Keywords"]
    col = [c.value for c in ws[1]].index("Monthly searches") + 1
    assert any(isinstance(ws.cell(row=r, column=col).value, int) for r in range(2, ws.max_row + 1))


def test_the_zip_has_a_csv_for_every_table_and_they_all_parse(rec):
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(rec)))
    assert sorted(z.namelist()) == sorted(n.lower().replace(" ", "-") + ".csv" for n in ex.tables(rec))
    for n in z.namelist():
        assert len(list(csv.reader(io.StringIO(z.read(n).decode("utf-8"))))) >= 2, n


def test_the_pdf_is_a_document_without_the_full_keyword_list(rec):
    pdf = ex.pdf(rec)
    assert pdf[:5] == b"%PDF-" and len(pdf) > 3000
    assert "Keywords" in ex.PDF_SKIP and "Ads Editor import" in ex.PDF_SKIP


def test_a_long_table_in_the_pdf_is_cut_and_says_so(rec, monkeypatch):
    big = json.loads(json.dumps(rec))
    row = dict(next(r for r in big["plan"]["keywords"] if r["include"]))
    big["plan"]["keywords"] = [dict(row, kw="added keyword %d" % i) for i in range(300)]
    monkeypatch.setattr(ex, "PDF_ROW_CAP", {"Add to account": 5})
    assert ex.pdf(big)[:5] == b"%PDF-"


def test_nothing_a_spreadsheet_would_run_as_a_formula_is_written_as_one(rec):
    bad = json.loads(json.dumps(rec))
    bad["name"] = "=1+1"
    bad["inputs"]["name"] = "=HYPERLINK(\"http://x\")"
    bad["plan"]["plan"]["summary"] = ["+cmd|' /C calc'!A0", "@SUM(1)"]
    bad["plan"]["editor"] = [["=Campaign", "-Group", "@kw", "Exact", "Paused"]]
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(bad)))
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value[:1] in "=+-@":
                    assert c.data_type == "s", (ws.title, c.coordinate, c.value)
    assert ex.csv_editor(bad).splitlines()[1].startswith("'=Campaign,'-Group,'@kw,")


def test_the_file_name_is_safe_and_names_the_agent(rec):
    assert ex.file_stem(rec) == "Acme-Dental-keyword-research" and ex.file_stem({"name": "../../etc/passwd  — x"}) == "etc-passwd-x-keyword-research" and ex.file_stem({}) == "client-keyword-research"


def test_a_result_with_nothing_in_it_still_makes_every_file():
    empty = {"name": "Nothing", "inputs": {"name": "Nothing"}, "plan": {"generatedAt": "2026-10-17T00:00:00+00:00"}, "research": {}}
    assert list(ex.tables(empty)) == ["Overview"] and ex.xlsx(empty)[:2] == b"PK" and ex.pdf(empty)[:5] == b"%PDF-" and ex.csv_zip(empty)[:2] == b"PK"


# ── what the mutation tests of this file found missing ─────────────────────

def test_the_pdf_holds_every_table_but_the_full_keyword_list_and_the_editor_file(rec):
    shown = [name for name, _h, _r, _n in ex.pdf_tables(rec)]
    assert "Keywords" not in shown and "Ads Editor import" not in shown and shown == [n for n in ex.tables(rec) if n not in ("Keywords", "Ads Editor import")]


def test_the_long_tables_of_the_pdf_are_cut_to_their_cap_and_say_how_many(rec):
    big = json.loads(json.dumps(rec))
    row = dict(next(r for r in big["plan"]["keywords"] if r["include"]))
    big["plan"]["keywords"] = [dict(row, kw="added keyword %d" % i) for i in range(300)]
    tables = {name: (rows, note) for name, _h, rows, note in ex.pdf_tables(big)}
    assert len(tables["Add to account"][0]) == ex.PDF_ROW_CAP["Add to account"] == 120 and tables["Add to account"][1] == "120 of 300 rows shown. All of them are in the Excel file."
    assert tables["Themes"][1] is None and len(tables["Themes"][0]) == len(big["plan"]["themes"])
    assert ex.PDF_ROW_CAP["Already in the account"] == 40
    assert ex.pdf(big)[:5] == b"%PDF-"


def test_a_count_of_zero_is_a_count_and_is_not_left_out_of_the_overview(rec):
    zero = json.loads(json.dumps(rec))
    zero["plan"]["coverage"].update(new=0, pct=0, yes=0)
    zero["plan"]["coverage"]["poolPct"] = 0
    rows = dict(ex.tables(zero)["Overview"][1][:17])
    assert rows["Fit and new to the account"] == 0 and rows["Already covered by the account, of the keywords judged a fit"] == "0%" and rows["Judged a fit"] == 0
    assert rows["Already covered by the account, of the searches for ideas that share a word with what the client sells"] == "0%"
    none = json.loads(json.dumps(rec))
    none["plan"]["coverage"].update(pct=None, poolPct=None)
    names = dict(ex.tables(none)["Overview"][1])
    assert "Already covered by the account, of the keywords judged a fit" not in names and "Already covered by the account, of the searches for ideas that share a word with what the client sells" not in names


def test_only_keywords_judged_a_fit_are_listed_as_already_in_the_account(rec):
    edited = json.loads(json.dumps(rec))
    covered = next(r for r in edited["plan"]["keywords"] if r["status"] == "covered")
    covered["fit"] = "maybe"
    names = [r[0] for r in ex.tables(edited)["Already in the account"][1]]
    assert covered["kw"] not in names and names


# ── what the independent review found ───────────────────────────────────────

def test_the_editor_column_says_why_a_keyword_that_fits_is_not_in_the_file(rec):
    fit = dict(kw="a", fit="yes", intent="local", status="new", include=False, editorOk=True)
    assert ex.editor_text(dict(fit, include=True)) == "Yes" and ex.editor_text(fit) == "No"
    assert ex.editor_text(fit, unread=True) == "No, the account's keywords could not be read" and ex.editor_text(dict(fit, fit="no"), unread=True) == "No"
    assert ex.editor_text(dict(fit, blockedInTarget=True)) == "No, the campaign it would go in already blocks it with a negative keyword"
    assert ex.editor_text(dict(fit, editorOk=False)) == "No, it holds a symbol Google Ads does not accept in a keyword"
    assert ex.editor_text(dict(fit, editorOk=False, status="covered", conv={"conv": 1})) == "No, it holds a symbol Google Ads does not accept in a keyword"
    assert ex.editor_text(dict(fit, editorOk=False, intent="competitor")) == "No" and ex.editor_text(dict(fit, editorOk=False, status="covered")) == "No"


def test_the_checks_sheet_carries_where_searches_were_counted_what_was_withheld_and_every_unbacked_number(rec):
    edited = json.loads(json.dumps(rec))
    chk = edited["plan"]["checks"]
    chk.update(geography="Searches are counted in Austin,Texas,United States.", editorWithheld="The account's keywords could not be read.", accountWarnings=["The account has more keywords than were read."],
               caveats=["Account-level negatives are not available."], numbers=[{"where": "In short", "text": "That brings 450 leads", "number": "450"}])
    edited["plan"]["coverage"]["blockedInTarget"] = 2
    rows = ex.tables(edited)["Checks"][1]
    kinds = [r[0] for r in rows]
    assert kinds[0] == "Where searches were counted" and rows[0][1] == "Searches are counted in Austin,Texas,United States." and "Google Ads Editor file not made" in kinds
    assert ["Account", "The account has more keywords than were read."] in rows and ["Not available from Google Ads", "Account-level negatives are not available."] in rows
    assert ["A number in the write-up that is not in the research: check it", "In short: 450 (That brings 450 leads)"] in rows
    assert ["Keywords left out of the Editor file because the campaign they would go in blocks them", "2"] in rows


def test_the_keywords_sheet_marks_a_keyword_blocked_in_its_own_campaign(rec):
    edited = json.loads(json.dumps(rec))
    row = next(r for r in edited["plan"]["keywords"] if r["fit"] == "yes" and r["status"] == "new")
    row.update(include=False, blockedInTarget=True)
    head, rows = ex.tables(edited)["Keywords"]
    assert {r[0]: r[-1] for r in rows}[row["kw"]] == "No, the campaign it would go in already blocks it with a negative keyword"


# ── Meta ────────────────────────────────────────────────────────────────────

def _built(platform):
    from tests import test_paid_keywords_meta_job as MJ
    store, sid = MJ.make(platform)
    MJ.run(store, sid)
    return json.loads(json.dumps(store.get(sid)))


@pytest.fixture(scope="module")
def meta_rec():
    return _built("meta")


def test_a_meta_result_has_its_own_sheets_and_none_of_googles(meta_rec):
    assert list(ex.tables(meta_rec)) == ["Overview", "Meta audiences", "Meta interests", "Meta in use", "Meta ad sets", "Meta plan", "Meta questions", "Meta checks"]
    for name, (header, rows) in ex.tables(meta_rec).items():
        assert rows, name
        assert all(len(r) == len(header) for r in rows), name
    ov = dict(ex.tables(meta_rec)["Overview"][1])
    assert ov["Meta ad account"] == "Acme Meta" and "Google Ads account" not in ov and "Keywords judged" not in ov and ov["Meta interests judged"] == meta_rec["plan"]["meta"]["coverage"]["judged"]
    assert ov["Ad set spend on interests judged not a fit (USD)"] == 600.0 and any(r[0] == "Meta: what the numbers say" for r in ex.tables(meta_rec)["Overview"][1])


def test_the_meta_sheets_say_what_each_interest_cost_and_how_to_act(meta_rec):
    head, rows = ex.tables(meta_rec)["Meta interests"]
    by = {r[0]: dict(zip(head, r)) for r in rows}
    assert by["Insulin pump"]["Fit"] == "yes" and by["Insulin pump"]["Audience size"] == "1.0M to 1.2M" and by["Insulin pump"]["Spent (USD)"] == 600.0 and by["Insulin pump"]["Cost per result (USD)"] == 30.0
    assert by["Diabetes"]["Fit"] == "no" and by["Insulin"]["Aimed at by ad sets"] == "" and rows[0][2] == "yes" and rows[-1][2] in ("no", "unjudged")
    head, plan = ex.tables(meta_rec)["Meta plan"]
    kinds = [r[0] for r in plan]
    assert "Quick win" in kinds and "Lookalike source" in kinds and any(k.startswith("How to: Start the audiences") for k in kinds)
    assert "1. In Ads Manager" in [r for r in plan if r[0].startswith("How to: Start")][0][2]


def test_a_build_for_both_has_the_google_sheets_and_then_the_meta_ones():
    rec = _built("both")
    names = list(ex.tables(rec))
    assert names.index("Keywords") < names.index("Meta audiences") < names.index("Meta interests") and names.index("Meta checks") < names.index("Checks")
    ov = dict(ex.tables(rec)["Overview"][1])
    assert ov["Google Ads account"] and ov["Meta ad account"] == "Acme Meta" and "Keywords judged" in ov and "Meta interests judged" in ov


def test_the_meta_files_open_and_an_interest_that_starts_like_a_formula_is_text(meta_rec):
    rec = json.loads(json.dumps(meta_rec))
    rec["plan"]["meta"]["interests"][0]["name"] = "=HYPERLINK(\"http://evil\")"
    wb = openpyxl.load_workbook(io.BytesIO(ex.xlsx(rec)))
    assert "Meta interests" in wb.sheetnames and "Meta audiences" in wb.sheetnames
    cells = [c for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value]
    assert cells and all(c.data_type == "s" for c in cells)                                       # the writer marks it as text, so a spreadsheet never runs it
    assert ex.pdf(meta_rec)[:4] == b"%PDF"
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(rec)))
    csv_cells = [c for row in csv.reader(io.StringIO(z.read("meta-interests.csv").decode("utf-8"))) for c in row if "HYPERLINK" in c]
    assert csv_cells and all(c.startswith("'=") for c in csv_cells)                              # in a CSV it gets a leading quote mark instead
    z = zipfile.ZipFile(io.BytesIO(ex.csv_zip(meta_rec)))
    assert "meta-audiences.csv" in z.namelist() and "google" not in " ".join(z.namelist()) and ex.csv_editor(meta_rec) is None


def test_the_pdf_cuts_the_long_meta_lists_and_says_so(meta_rec):
    rec = json.loads(json.dumps(meta_rec))
    base = rec["plan"]["meta"]["interests"][0]
    rec["plan"]["meta"]["interests"] = [dict(base, id=str(i), name="Interest %d" % i) for i in range(90)]
    shown = {name: (rows, cut) for name, _h, rows, cut in ex.pdf_tables(rec)}
    assert len(shown["Meta interests"][0]) == 60 and shown["Meta interests"][1] == "60 of 90 rows shown. All of them are in the Excel file."


def test_inside_a_fit_an_interest_no_ad_set_aims_at_comes_before_one_in_use_and_the_note_for_a_cut_list_and_advantage_are_written():
    meta = {"currency": "USD", "interests": [
        {"id": "1", "name": "Big in use", "kind": "interest", "fit": "yes", "intent": "core", "theme": "", "size": 9_000_000, "sizeText": "9.0M", "inUse": {"adsets": 1, "cost": 10.0, "conv": 1.0, "cpa": 10.0}, "from": "", "path": [], "note": ""},
        {"id": "2", "name": "Small new", "kind": "interest", "fit": "yes", "intent": "core", "theme": "", "size": 90_000, "sizeText": "90K", "inUse": None, "from": "", "path": [], "note": ""}],
        "adsets": [{"name": "A", "campaign": "C", "status": "ACTIVE", "interests": [], "behaviors": [], "audiences": [], "excluded": [], "advantage": True, "cost": 1.0, "conv": 0.0, "cpa": None},
                   {"name": "B", "campaign": "C", "status": "ACTIVE", "interests": [], "behaviors": [], "audiences": [], "excluded": [], "advantage": False, "cost": 1.0, "conv": 0.0, "cpa": None}],
        "checks": {"cut": True, "found": 150, "advantageAdsets": 1}}
    t = ex.meta_tables(meta)
    assert [r[0] for r in t["Meta interests"][1]] == ["Small new", "Big in use"]
    assert [r[7] for r in t["Meta ad sets"][1]] == ["Yes", "No"]
    checks = dict(t["Meta checks"][1])
    assert checks["More interests were found than were judged"] == "150 found, the first 2 judged (the ones in use first)." and "suggestions to Meta" in checks["Ad sets on Advantage+ audience"]


def test_a_platform_that_could_not_be_written_is_named_in_the_overview(meta_rec):
    rec = json.loads(json.dumps(meta_rec))
    rec["plan"]["errors"] = {"Google": "Keyword Planner gave no ideas"}
    assert ["Google part not written", "Keyword Planner gave no ideas"] in ex.tables(rec)["Overview"][1]
    assert "Meta audiences" in ex.PDF_PAGE_BREAK
