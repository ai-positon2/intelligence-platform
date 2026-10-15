"""Paid Marketing Agents: QA Agent, the campaign book the team adds.

Pinned here: the Campaign Book layout (Field | Value | Other things to Check) is read into rows,
and anything else on the sheets is kept as text for the AI; placeholders like "Depends on the
campaign book" are not mistaken for an answer; labels match through the workbook's typos; a file
that is too big, is not a workbook, is corrupt or is an old .xls is refused in words; a Google
Sheet is fetched only from docs.google.com, only as an export of that sheet, and a private one is
said to be private rather than parsed. No network: the fetcher is a stand-in.
"""
import base64
import io
import os
import sys

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from tracker import paid_qa_book as book  # noqa: E402


def workbook(rows, title="Sheet1", extra=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    for name, rr in (extra or {}).items():
        w2 = wb.create_sheet(name)
        for r in rr:
            w2.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


TEMPLATE = [["Field ", "Value", "Other things to Check"],
            ["Campaign Name", "Should follow the format - P2_Campaign Type(Search or Pmax)_Launchdate"],
            ["Bid Strategy Type", "Maximize conversions"],
            ["Devices", "Desktop and mobile, no tablets", "Check the devices"],
            ["Ad Shceduling", "Mon-Fri 8am to 6pm"],
            ["Genders", "Depends on the campaign book", "But you can still check from the domain"],
            ["Household Income", "Depends on the client"],
            ["Age", "Check the age"],
            [None, None, None],
            ["Language targeting", "English"]]


# ── Reading a workbook ──────────────────────────────────────────────────────

def test_the_template_layout_becomes_field_value_and_other_rows():
    b = book.parse_xlsx(workbook(TEMPLATE), "Campaign Book.xlsx")
    fields = {r["field"]: r for r in b["rows"]}
    assert fields["Devices"] == {"field": "Devices", "value": "Desktop and mobile, no tablets", "other": "Check the devices"}
    assert fields["Language targeting"]["value"] == "English"
    assert "Field" not in fields and len(b["rows"]) == 8                      # the header row is not data; blank rows are skipped
    assert b["filename"] == "Campaign Book.xlsx" and b["source"] == "upload" and b["sheets"] == ["Sheet1"]


def test_everything_on_every_sheet_is_kept_as_text_for_the_ai():
    b = book.parse_xlsx(workbook(TEMPLATE, extra={"Notes": [["Launch", "2026-10-14"], ["Client", "Example Co"]]}))
    assert "Sheet1 | Devices | Desktop and mobile, no tablets | Check the devices" in b["text"]
    assert "Notes | Launch | 2026-10-14" in b["text"] and b["sheets"] == ["Sheet1", "Notes"]


def test_a_book_that_does_not_use_the_template_layout_is_still_read():
    b = book.parse_xlsx(workbook([["Campaign", "Brand Search", "Non-brand"], ["Daily budget", 50, 120], ["Bid strategy", "tCPA", "Max conversions"]]))
    assert {"field": "Daily budget", "value": "50", "other": "120"} in b["rows"]
    assert "Sheet1 | Daily budget | 50 | 120" in b["text"]


def test_numbers_and_floats_are_read_as_plain_text():
    b = book.parse_xlsx(workbook([["Field", "Value"], ["Budget", 50.0], ["Target CPA", 42.5]]))
    assert {r["field"]: r["value"] for r in b["rows"]} == {"Budget": "50", "Target CPA": "42.5"}


def test_long_cells_are_cut_and_the_text_dump_is_capped():
    b = book.parse_xlsx(workbook([["Field", "Value"]] + [["Row %d" % i, "x" * 700] for i in range(60)]))
    assert all(len(r["value"]) <= book.MAX_CELL for r in b["rows"]) and len(b["text"]) <= book.MAX_TEXT


def test_an_empty_workbook_is_refused():
    with pytest.raises(book.BookError, match="no text"):
        book.parse_xlsx(workbook([[None, None]]))


def test_a_file_that_is_not_a_workbook_a_corrupt_one_and_a_huge_one_are_refused():
    with pytest.raises(book.BookError, match="not an .xlsx"):
        book.parse_xlsx(b"just some text, not a zip")
    with pytest.raises(book.BookError, match="could not be read"):
        book.parse_xlsx(b"PK\x03\x04 this is not really a workbook")
    with pytest.raises(book.BookError, match="over 4 MB"):
        book.parse_xlsx(b"PK" + b"0" * (book.MAX_BYTES + 1))


def test_a_workbook_that_unzips_to_something_enormous_is_refused(monkeypatch):
    monkeypatch.setattr(book, "MAX_UNZIPPED", 1000)
    with pytest.raises(book.BookError, match="too large"):
        book.parse_xlsx(workbook(TEMPLATE))


# ── CSV and browser uploads ─────────────────────────────────────────────────

def test_a_csv_is_read_the_same_way():
    b = book.parse_csv(b"Field,Value,Other things to Check\r\nDevices,Mobile only,\r\nGenders,Women,From the domain\r\n", "book.csv")
    assert [(r["field"], r["value"]) for r in b["rows"]] == [("Devices", "Mobile only"), ("Genders", "Women")]


def test_an_upload_from_the_browser_is_base64_with_or_without_the_data_url_prefix():
    raw = workbook(TEMPLATE)
    plain = book.from_upload("Book.xlsx", base64.b64encode(raw).decode())
    prefixed = book.from_upload("Book.xlsx", "data:application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;base64," + base64.b64encode(raw).decode())
    assert plain["rows"] == prefixed["rows"] and plain["filename"] == "Book.xlsx"
    assert book.from_upload("b.csv", base64.b64encode(b"Field,Value\nDevices,Mobile").decode())["rows"][0]["value"] == "Mobile"


def test_bad_uploads_are_refused_in_words():
    for name, data, says in (("b.xlsx", "", "empty"), ("b.xlsx", None, "empty"), ("old.xls", base64.b64encode(b"x").decode(), "old .xls"),
                             ("b.xlsx", base64.b64encode(b"not a zip").decode(), "not an .xlsx")):
        with pytest.raises(book.BookError, match=says):
            book.from_upload(name, data)


def test_a_file_name_cannot_carry_markup_or_a_path():
    b = book.from_upload("../../<script>alert(1)</script> Book.xlsx", base64.b64encode(workbook(TEMPLATE)).decode())
    assert "<" not in b["filename"] and "/" not in b["filename"] and len(b["filename"]) <= 120


# ── Finding a value ─────────────────────────────────────────────────────────

def test_a_value_is_found_through_the_workbooks_typos_and_case():
    b = book.parse_xlsx(workbook(TEMPLATE))
    assert book.lookup(b, "bid strategy type", "bid strategy") == "Maximize conversions"
    assert book.lookup(b, "ad scheduling") == "Mon-Fri 8am to 6pm"                     # the sheet says "Shceduling"
    assert book.lookup(b, "DEVICES") == "Desktop and mobile, no tablets"
    assert book.lookup(b, "language") == "English"                                      # a label that starts with the name


def test_a_placeholder_is_not_an_answer():
    b = book.parse_xlsx(workbook(TEMPLATE))
    assert book.lookup(b, "genders") is None and book.lookup(b, "household income") is None and book.lookup(b, "age") is None
    assert book.lookup(b, "locations") is None and book.lookup(None, "devices") is None and book.lookup({}, "devices") is None
    assert not book.filled("Depends on the campaign book") and not book.filled("  ") and book.filled("Women only")


# ── A Google Sheet link ─────────────────────────────────────────────────────

SHEET = "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/edit?gid=0#gid=0"


def test_only_a_google_sheets_link_is_accepted_and_it_becomes_an_export_of_that_sheet():
    assert book.sheet_export_url(SHEET) == "https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/export?format=xlsx"
    for bad in ("", None, "https://example.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789", "http://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789",
                "https://docs.google.com.evil.example/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789", "https://docs.google.com/document/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789",
                "https://docs.google.com/spreadsheets/d/short", "javascript:alert(1)", "file:///etc/passwd"):
        with pytest.raises(book.BookError, match="not a Google Sheets link"):
            book.sheet_export_url(bad)


class Raw:
    def __init__(self, body):
        self.body = body

    def read(self, n=None, decode_content=True):
        return self.body[:n] if n else self.body


class Resp:
    def __init__(self, status, body):
        self.status_code, self.raw, self.closed = status, Raw(body), False

    def close(self):
        self.closed = True


def test_a_shared_sheet_is_fetched_as_an_export_and_read():
    seen = []

    def fetcher(url, timeout=20, headers=None):
        seen.append(url)
        return Resp(200, workbook(TEMPLATE))
    b = book.from_google_sheet(SHEET, fetcher)
    assert seen == ["https://docs.google.com/spreadsheets/d/1AbCdEfGhIjKlMnOpQrStUvWxYz0123456789/export?format=xlsx"]
    assert b["source"] == "google_sheet" and b["filename"] == "Google Sheet" and len(b["rows"]) == 8


def test_a_private_sheet_is_said_to_be_private_not_parsed():
    for status, body in ((200, b"<html>Sign in to continue</html>"), (401, b""), (403, b"nope"), (404, b"gone")):
        with pytest.raises(book.BookError, match="private"):
            book.from_google_sheet(SHEET, lambda url, timeout=20, headers=None, s=status, b=body: Resp(s, b))


def test_a_fetch_that_fails_is_reported_not_raised():
    def boom(url, timeout=20, headers=None):
        raise TimeoutError("slow")
    with pytest.raises(book.BookError, match="could not be fetched"):
        book.from_google_sheet(SHEET, boom)


def test_the_real_fetcher_refuses_private_addresses_so_a_link_cannot_reach_our_network(monkeypatch):
    # the link is checked to be docs.google.com first; even if that host resolved to a private address the guard refuses it
    from tracker import event_intel_http
    monkeypatch.setattr(event_intel_http.socket, "getaddrinfo", lambda *a, **k: [(2, 1, 6, "", ("10.0.0.5", 443))])
    with pytest.raises(book.BookError, match="could not be fetched"):
        book.from_google_sheet(SHEET)


# ── What the page is told ───────────────────────────────────────────────────

def test_the_page_is_told_the_file_and_how_much_was_read_but_not_the_text():
    b = book.parse_xlsx(workbook(TEMPLATE), "Campaign Book.xlsx")
    b.update(uploadedBy="a@position2.com", uploadedAt="2026-10-15T09:00:00+00:00")
    pub = book.public(b)
    assert pub == {"filename": "Campaign Book.xlsx", "source": "upload", "fields": 8, "sheets": ["Sheet1"],
                   "uploadedBy": "a@position2.com", "uploadedAt": "2026-10-15T09:00:00+00:00"}
    assert book.public(None) is None and book.public({}) is None
